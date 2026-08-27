#!/usr/bin/env python3
"""Check whether source/target rows preserve numbers and URLs."""

from __future__ import annotations

import re
import unicodedata
from collections import Counter
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from pathlib import Path

from openpyxl.utils import column_index_from_string

from tools.excel_output import (
    PROBLEM_BASE_HEADERS,
    build_prefixed_output_path,
    find_last_value_row,
    load_workbook_for_editing,
    validate_distinct_source_target_columns,
    write_output_table,
)
from tools.tag_placeholder_checker.check_tags_and_placeholders import (
    CANONICAL_TOKEN_TYPES,
    extract_token_details,
)


NUMBER_RULE = "number"
URL_RULE = "url"
SUPPORTED_RULES = (NUMBER_RULE, URL_RULE)
NUMBER_PROBLEM_SHEET_NAME = "数字不一致"
URL_PROBLEM_SHEET_NAME = "URL不一致"

URL_PATTERN = re.compile(
    r"(?i)(?<![\w@])(?:(?:https?|ftp|file)://|mailto:|www\.)"
    r"[^\s<>\"'“”‘’]+"
)
NUMBER_PATTERN = re.compile(
    r"(?<!\d)[+\-]?(?:"
    r"\d+(?:[.,:/\-'’]\d+)+"
    r"|\d{1,3}(?:[ '\u00a0\u202f’]\d{3})+(?:[.,]\d+)?"
    r"|\d+(?:[.,]\d+)?"
    r")(?:[%‰])?(?!\d)"
)
ENCODED_CHARACTER_PATTERN = re.compile(
    r"&#(?:x[0-9a-f]+|\d+);|\\u[0-9a-f]{4,8}|\\x[0-9a-f]{2}",
    re.IGNORECASE,
)
TRAILING_URL_PUNCTUATION = ".,;:!?，。；：！？、…"
TRAILING_URL_PAIRS = {
    ")": "(",
    "]": "[",
    "}": "{",
    "）": "（",
    "】": "【",
    "》": "《",
    "」": "「",
    "』": "『",
}
DASH_TRANSLATION = str.maketrans({"−": "-", "–": "-", "—": "-", "﹣": "-"})


@dataclass(frozen=True)
class CheckSummary:
    output_path: Path
    worksheet_title: str
    source_column: str
    target_column: str
    start_row: int
    total_rows_checked: int
    selected_rules: tuple[str, ...]
    number_problem_rows: int
    url_problem_rows: int


def normalize_column(column_name: str) -> str:
    normalized = column_name.strip().upper()
    column_index_from_string(normalized)
    return normalized


def cell_text(value: object) -> str:
    return "" if value is None else str(value)


def normalize_rules(rules: Iterable[str] | None) -> tuple[str, ...]:
    raw_rules = tuple(SUPPORTED_RULES if rules is None else rules)
    invalid_rules = [rule for rule in raw_rules if rule not in SUPPORTED_RULES]
    if invalid_rules:
        raise ValueError(f"不支持的内容保真规则: {'、'.join(invalid_rules)}")
    normalized = tuple(rule for rule in SUPPORTED_RULES if rule in raw_rules)
    if not normalized:
        raise ValueError("请至少选择一种内容保真检查规则。")
    return normalized


def build_default_output_path(input_file: str | Path) -> Path:
    return build_prefixed_output_path(input_file, "content_fidelity_check_")


def _trim_url(candidate: str) -> str:
    trimmed = candidate.rstrip(TRAILING_URL_PUNCTUATION)
    while trimmed:
        closing = trimmed[-1]
        opening = TRAILING_URL_PAIRS.get(closing)
        if opening is None or trimmed.count(closing) <= trimmed.count(opening):
            break
        trimmed = trimmed[:-1].rstrip(TRAILING_URL_PUNCTUATION)
    return trimmed


def iter_url_spans(text: str) -> Iterator[tuple[int, int, str]]:
    for match in URL_PATTERN.finditer(text):
        token = _trim_url(match.group(0))
        if token:
            yield match.start(), match.start() + len(token), token


def extract_urls(text: object) -> tuple[str, ...]:
    raw_text = cell_text(text)
    return tuple(token for _, _, token in iter_url_spans(raw_text))


def _mask_ranges(text: str, ranges: Iterable[tuple[int, int]]) -> str:
    characters = list(text)
    for start, end in ranges:
        for index in range(max(0, start), min(len(characters), end)):
            characters[index] = " "
    return "".join(characters)


def extract_numbers(text: object) -> tuple[str, ...]:
    normalized_text = unicodedata.normalize("NFKC", cell_text(text)).translate(
        DASH_TRANSLATION
    )
    masked_ranges = [
        (start, end) for start, end, _ in iter_url_spans(normalized_text)
    ]
    masked_ranges.extend(
        (match.start(), match.end())
        for match in ENCODED_CHARACTER_PATTERN.finditer(normalized_text)
    )
    masked_ranges.extend(
        (token.start, token.end)
        for token in extract_token_details(
            normalized_text,
            token_types=CANONICAL_TOKEN_TYPES,
        )
    )
    searchable_text = _mask_ranges(normalized_text, masked_ranges)
    return tuple(match.group(0) for match in NUMBER_PATTERN.finditer(searchable_text))


def _format_values(counter: Counter[str]) -> str:
    return " | ".join(
        f"{value} × {count}" if count > 1 else value
        for value, count in counter.items()
    )


def _problem_entry(
    *,
    row_index: int,
    source_text: str,
    target_text: str,
    label: str,
    source_values: tuple[str, ...],
    target_values: tuple[str, ...],
) -> tuple[int, str, str, str, str, str, str, str]:
    source_counter = Counter(source_values)
    target_counter = Counter(target_values)
    return (
        row_index,
        source_text,
        target_text,
        f"{label}不一致",
        _format_values(source_counter),
        _format_values(target_counter),
        _format_values(source_counter - target_counter),
        _format_values(target_counter - source_counter),
    )


def process_excel(
    input_file: str | Path,
    source_column: str,
    target_column: str,
    sheet: str | None = None,
    start_row: int = 2,
    rules: Iterable[str] | None = None,
    output_file: str | Path | None = None,
) -> CheckSummary:
    if start_row < 1:
        raise ValueError("开始行必须大于等于 1。")

    input_path = Path(input_file).expanduser().resolve()
    if not input_path.exists():
        raise FileNotFoundError(f"输入文件不存在: {input_path}")
    source_column = normalize_column(source_column)
    target_column = normalize_column(target_column)
    validate_distinct_source_target_columns(source_column, target_column)
    rules = normalize_rules(rules)
    output_path = (
        Path(output_file).expanduser().resolve()
        if output_file
        else build_default_output_path(input_path)
    )

    workbook = load_workbook_for_editing(input_path)
    try:
        summary = process_workbook(
            workbook=workbook,
            output_path=output_path,
            source_column=source_column,
            target_column=target_column,
            sheet=sheet,
            start_row=start_row,
            rules=rules,
        )
        workbook.save(output_path)
        return summary
    finally:
        workbook.close()


def process_workbook(
    *,
    workbook,
    output_path: Path,
    source_column: str,
    target_column: str,
    sheet: str | None = None,
    start_row: int = 2,
    rules: Iterable[str] | None = None,
    format_output: bool = True,
) -> CheckSummary:
    """Run the selected checks against an already-open workbook without saving it."""
    if start_row < 1:
        raise ValueError("开始行必须大于等于 1。")

    source_column = normalize_column(source_column)
    target_column = normalize_column(target_column)
    validate_distinct_source_target_columns(source_column, target_column)
    selected_rules = normalize_rules(rules)
    worksheet = workbook[sheet] if sheet else workbook.active
    last_row = find_last_value_row(
        worksheet,
        (source_column, target_column),
        start_row=start_row,
    )
    number_entries: list[tuple[object, ...]] = []
    url_entries: list[tuple[object, ...]] = []

    for row_index in range(start_row, last_row + 1):
        source_text = cell_text(worksheet[f"{source_column}{row_index}"].value)
        target_text = cell_text(worksheet[f"{target_column}{row_index}"].value)
        if NUMBER_RULE in selected_rules:
            source_numbers = extract_numbers(source_text)
            target_numbers = extract_numbers(target_text)
            if Counter(source_numbers) != Counter(target_numbers):
                number_entries.append(
                    _problem_entry(
                        row_index=row_index,
                        source_text=source_text,
                        target_text=target_text,
                        label="数字",
                        source_values=source_numbers,
                        target_values=target_numbers,
                    )
                )
        if URL_RULE in selected_rules:
            source_urls = extract_urls(source_text)
            target_urls = extract_urls(target_text)
            if Counter(source_urls) != Counter(target_urls):
                url_entries.append(
                    _problem_entry(
                        row_index=row_index,
                        source_text=source_text,
                        target_text=target_text,
                        label="URL",
                        source_values=source_urls,
                        target_values=target_urls,
                    )
                )

    detail_headers = ("Source 内容", "Target 内容", "Target 缺少", "Target 多出")
    if NUMBER_RULE in selected_rules:
        write_output_table(
            workbook,
            current_sheet_name=worksheet.title,
            sheet_name=NUMBER_PROBLEM_SHEET_NAME,
            headers=PROBLEM_BASE_HEADERS + detail_headers,
            rows=number_entries,
            row_link_target_column=target_column,
            format_output=format_output,
        )
    if URL_RULE in selected_rules:
        write_output_table(
            workbook,
            current_sheet_name=worksheet.title,
            sheet_name=URL_PROBLEM_SHEET_NAME,
            headers=PROBLEM_BASE_HEADERS + detail_headers,
            rows=url_entries,
            row_link_target_column=target_column,
            format_output=format_output,
        )

    return CheckSummary(
        output_path=output_path,
        worksheet_title=worksheet.title,
        source_column=source_column,
        target_column=target_column,
        start_row=start_row,
        total_rows_checked=max(0, last_row - start_row + 1),
        selected_rules=selected_rules,
        number_problem_rows=len(number_entries),
        url_problem_rows=len(url_entries),
    )
