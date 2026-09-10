from __future__ import annotations

import math
from pathlib import Path
import re

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import column_index_from_string

from tools.excel_file_ops import BatchSummary, OperationResult, excel_files, optional_output_directory, create_output_directory
from tools.operation_logs import emit_log, log_result, log_summary
from tools.content_sync.master_to_target import _save_output, _sheet


def count_units(value: object, mode: str) -> int:
    text = "" if value is None else str(value).strip()
    if mode == "chinese_chars":
        return len(re.findall(r"[\u4e00-\u9fa5]", text))
    if mode == "english_words":
        return len(re.findall(r"\b[a-zA-Z]+(?:['\-][a-zA-Z]+)*\b", text))
    raise ValueError("计数模式必须是 chinese_chars 或 english_words。")


def translation_empty(value: object) -> bool:
    # Statistics intentionally retains momoTools' 'nan' convention; sync does not.
    return value is None or str(value).strip().casefold() in {"", "nan"}


def untranslated_stats(target_dir: str | Path, *, output_dir: str | Path | None = None,
                       source_column: str = "B", target_column: str = "C",
                       sheet: str | None = None, header_rows: int = 1,
                       mode: str = "chinese_chars", progress_callback=None, log_callback=None) -> BatchSummary:
    count_units("", mode)
    emit_log(log_callback, f"开始未翻译统计：{target_dir}（{'英文词数' if mode == 'english_words' else '中文字符数'}）")
    source_col, target_col = (column_index_from_string(col.strip().upper()) for col in (source_column, target_column))
    if source_col == target_col or not all(1 <= col <= 16384 for col in (source_col, target_col)):
        raise ValueError("Source 和译文列必须在 A–XFD 范围内且不同。")
    if isinstance(header_rows, bool) or not isinstance(header_rows, int) or not 0 <= header_rows < 1048576:
        raise ValueError("表头行数必须是 0–1048575 的整数。")
    folder = Path(target_dir).expanduser().resolve()
    output, inplace = optional_output_directory(output_dir, default=folder, inputs=(folder,))
    report_file = output / "未翻译统计.xlsx"
    files, skipped = excel_files(folder, extensions={".xlsx", ".xlsm"}, exclude=() if inplace else (output,))
    # The tool's reserved result name must never become its next source workbook.
    files = [path for path in files if path.name.casefold() != "未翻译统计.xlsx"]
    if not files:
        raise ValueError("目录中没有支持的 .xlsx/.xlsm 文件。")
    summary = BatchSummary("untranslated-stats", output, details={"mode": mode, "source_column": source_column,
                                                                 "target_column": target_column, "sheet": sheet,
                                                                 "header_rows": header_rows})
    totals = [0, 0, 0, 0]
    for index, path in enumerate(files, 1):
        workbook = None
        try:
            workbook = load_workbook(path, read_only=True, data_only=True)
            worksheet = _sheet(workbook, sheet)
            if worksheet.max_column is not None and max(source_col, target_col) > worksheet.max_column:
                raise ValueError("工作表缺少配置列。")
            counts = [0, 0, 0, 0]  # untranslated units/rows, all units/rows
            start, end = min(source_col, target_col), max(source_col, target_col)
            for row in worksheet.iter_rows(min_row=header_rows + 1, min_col=start, max_col=end, values_only=True):
                source, target = row[source_col - start], row[target_col - start]
                if source is None or not str(source).strip() or isinstance(source, float) and math.isnan(source):
                    continue
                units = count_units(source, mode)
                counts[2] += units
                counts[3] += 1
                if translation_empty(target):
                    counts[0] += units
                    counts[1] += 1
            totals = [a + b for a, b in zip(totals, counts)]
            summary.results.append(OperationResult(str(path.relative_to(folder)), "read", details={"counts": counts}))
        except Exception as exc:
            summary.results.append(OperationResult(str(path.relative_to(folder)), "failed", message=str(exc)))
        finally:
            if workbook is not None:
                workbook.close()
        log_result(log_callback, summary.results[-1])
        if progress_callback:
            progress_callback(index, len(files))
    summary.results.extend(OperationResult(str(path.relative_to(folder)), "skipped", message="不支持的格式") for path in skipped)
    for result in summary.results:
        if result.status == "skipped":
            log_result(log_callback, result)
    if not inplace:
        create_output_directory(output)
    report = Workbook()
    try:
        worksheet = report.active
        worksheet.title = "未翻译统计"
        unit = "词数" if mode == "english_words" else "字数"
        worksheet.append(["文件名", f"未翻译{unit}", "未翻译行数", f"总{unit}", "总行数"])
        for result in summary.results:
            if result.status == "read":
                worksheet.append([result.source, *result.details["counts"]])
                worksheet.cell(worksheet.max_row, 1).data_type = "s"
        worksheet.append(["总计", *totals])
        for cell in worksheet[worksheet.max_row]:
            cell.font = Font(bold=True)
        worksheet.column_dimensions["A"].width = 45
        for col in "BCDE":
            worksheet.column_dimensions[col].width = 16
        worksheet.freeze_panes = "A2"
        failures = [item for item in summary.results if item.status in {"failed", "skipped"}]
        if failures:
            errors = report.create_sheet("未统计文件")
            errors.append(["文件名", "状态", "原因"])
            for result in failures:
                errors.append([result.source, result.status, result.message])
                for cell in errors[errors.max_row]:
                    cell.data_type = "s"
            errors.column_dimensions["A"].width = 45
            errors.column_dimensions["C"].width = 80
        _save_output(report, report_file, report_file, changed=True)
    finally:
        report.close()
    summary.details.update({"totals": totals, "excel_report": str(report_file)})
    emit_log(log_callback, f"统计表: {report_file}")
    log_summary(log_callback, summary, f"未翻译量 {totals[0]}/{totals[2]}，未翻译行 {totals[1]}/{totals[3]}")
    return summary
