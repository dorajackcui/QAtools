"""Match Master rows to target workbooks without a DataFrame conversion layer."""

from __future__ import annotations

from collections.abc import Callable
from contextlib import ExitStack
from dataclasses import asdict, dataclass, field
import os
from pathlib import Path
import shutil
import tempfile

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

from tools.excel_output import iter_value_cells, load_workbook_for_editing
from tools.excel_file_ops import create_output_directory, optional_output_directory, BACKUP_DIRECTORY
from tools.operation_logs import emit_log, log_result, log_summary
from .parallel import ordered_results, validate_workers


SUPPORTED_EXTENSIONS = {".xlsx", ".xlsm"}


def content_text(value: object) -> str:
    """An empty cell never becomes a visible 'None' or 'nan' marker."""
    return "" if value is None else str(value)


def is_blank(value: object) -> bool:
    return value is None or isinstance(value, str) and not value.strip()


def row_identity(key: object, match: object) -> tuple[str, str] | None:
    parts = (content_text(key).strip(), content_text(match).strip())
    return parts if all(parts) else None


@dataclass(frozen=True)
class ColumnMapping:
    key: str
    match: str
    content: str

    def indexes(self, count: int) -> tuple[int, int, int]:
        indexes = tuple(column_index_from_string(col.strip().upper()) for col in
                        (self.key, self.match, self.content))
        key, match, content = indexes
        if not all(1 <= col <= 16384 for col in indexes) or content + count - 1 > 16384:
            raise ValueError("列配置超出 Excel 的 A–XFD 范围。")
        if key == match or any(content <= col < content + count for col in (key, match)):
            raise ValueError("Key、原文及内容列范围不能重叠。")
        return key, match, content


@dataclass
class FileResult:
    source: str
    output: str | None = None
    status: str = "failed"
    sheet: str | None = None
    matched_rows: int = 0
    updated_cells: int = 0
    unchanged_cells: int = 0
    blank_source_cells: int = 0
    occupied_cells: int = 0
    invalid_rows: int = 0
    unmatched_rows: list[int] = field(default_factory=list)
    error: str | None = None
    postprocess_error: str | None = None


@dataclass
class SyncSummary:
    output_dir: Path
    master_row_count: int
    duplicate_keys: list[dict[str, object]]
    files: list[FileResult]
    warnings: list[str] = field(default_factory=list)

    @property
    def succeeded_files(self) -> int:
        return sum(item.status in {"updated", "unchanged"} and not item.postprocess_error for item in self.files)

    @property
    def failed_files(self) -> int:
        return sum(item.status == "failed" or bool(item.postprocess_error) for item in self.files)

    @property
    def skipped_files(self) -> int:
        return sum(item.status == "skipped" for item in self.files)

    @property
    def updated_cells(self) -> int:
        return sum(item.updated_cells for item in self.files if item.status == "updated")


def format_summary(summary: SyncSummary) -> str:
    return "\n".join((
        f"输出目录: {summary.output_dir}",
        f"成功文件: {summary.succeeded_files}；失败: {summary.failed_files}；跳过: {summary.skipped_files}",
        f"实际更新单元格: {summary.updated_cells}",
        f"Master 重复身份: {len(summary.duplicate_keys)}（后行优先）",
        f"未匹配行: {sum(len(item.unmatched_rows) for item in summary.files)}",
        "逐文件详情请查看运行日志。",
    ) + tuple(summary.warnings))


def build_default_output_dir(target_dir: str | Path) -> Path:
    folder = Path(target_dir).expanduser().absolute()
    return folder.with_name(f"{folder.name}_synced")


def _sheet(workbook, name: str | None):
    if name and name not in workbook.sheetnames:
        raise ValueError(f"找不到工作表: {name}")
    worksheet = workbook[name] if name else workbook.active
    if worksheet is None or not hasattr(worksheet, "iter_rows"):
        raise ValueError("请选择数据工作表。")
    return worksheet


def _is_previous_output(folder: Path) -> bool:
    """Exclude earlier generated trees when the user reruns a parent folder."""
    from tools.excel_file_ops import generated_directory
    return generated_directory(folder)


def _input_files(folder: Path, output: Path | None, master: Path):
    files, skipped = [], []
    def scan_failed(error: OSError) -> None:
        raise error

    for root, directories, names in os.walk(folder, followlinks=False, onerror=scan_failed):
        parent = Path(root)
        directories[:] = [name for name in directories
                          if name.casefold() != BACKUP_DIRECTORY and (parent / name).resolve() != output
                          and (parent / name).resolve().is_relative_to(folder)
                          and not _is_previous_output(parent / name)]
        for name in names:
            path = parent / name
            if name.startswith("~$"):
                continue
            resolved = path.resolve()
            if resolved == master or (output is not None and resolved.is_relative_to(output)) or not resolved.is_relative_to(folder):
                continue
            if path.suffix.lower() in SUPPORTED_EXTENSIONS:
                files.append(path)
            elif path.suffix.lower() in {".xls", ".xlsb"}:
                skipped.append(FileResult(str(path.relative_to(folder)), status="skipped", error="不支持的格式"))
    files.sort(key=lambda p: (str(p).casefold(), str(p)))
    skipped.sort(key=lambda item: (item.source.casefold(), item.source))
    return files, skipped


def _read_master(path: Path, mapping: ColumnMapping, count: int, sheet: str | None,
                 header_rows: int, *, skip_blank_content: bool = False, log_callback=None):
    key_col, match_col, content_col = mapping.indexes(count)
    end_col = max(key_col, match_col, content_col + count - 1)
    start_col = min(key_col, match_col, content_col)
    records: dict[tuple[str, str], tuple[int, tuple[str, ...]]] = {}
    duplicates: list[dict[str, object]] = []
    # Read formulas as well as values so missing caches cannot silently clear targets.
    with ExitStack() as stack:
        # Own the stream as well: an interrupted read-only iterator may retain a
        # ZIP member after Workbook.close(), otherwise locking the file on Windows.
        stream = stack.enter_context(path.open("rb"))
        workbook = load_workbook(stream, read_only=True, data_only=False)
        stack.callback(workbook.close)
        worksheet = _sheet(workbook, sheet)
        sheet_name = worksheet.title
        if worksheet.max_column is not None and max(key_col, match_col) > worksheet.max_column:
            raise ValueError("来源工作表缺少配置的 Key 或原文列。")
        cached_rows = None
        cached_row_number = 0
        selected = {key_col, match_col, *range(content_col, content_col + count)}
        for row_number, cells in enumerate(worksheet.iter_rows(
            min_row=header_rows + 1, min_col=start_col, max_col=end_col,
        ), header_rows + 1):
            values = {col: cells[col - start_col].value for col in selected}
            if any(cells[col - start_col].data_type != "f" and not content_text(values[col]).strip()
                   for col in (key_col, match_col)):
                continue
            formula_cols = [col for col in selected if cells[col - start_col].data_type == "f"]
            if formula_cols:
                if cached_rows is None:
                    cached_stream = stack.enter_context(path.open("rb"))
                    cached_book = load_workbook(cached_stream, read_only=True, data_only=True)
                    stack.callback(cached_book.close)
                    cached_rows = iter(cached_book[sheet_name].iter_rows(
                        min_row=row_number, min_col=start_col, max_col=end_col, values_only=True,
                    ))
                    cached_row_number = row_number - 1
                cached_values = ()
                while cached_row_number < row_number:
                    cached_values = next(cached_rows, ())
                    cached_row_number += 1
                for col in formula_cols:
                    value = cached_values[col - start_col] if col - start_col < len(cached_values) else None
                    if value is None:
                        address = f"{sheet_name}!{get_column_letter(col)}{row_number}"
                        raise ValueError(f"来源公式缺少缓存值: {address}；请在 Excel 中计算并保存，或转换为值。")
                    values[col] = value
            identity = row_identity(values[key_col], values[match_col])
            if identity is None:
                continue
            content = tuple(content_text(values[col]) for col in range(content_col, content_col + count))
            if skip_blank_content and all(is_blank(value) for value in content):
                continue
            if identity in records:
                old_row, old_content = records[identity]
                duplicates.append({
                    "key": identity[0], "match": identity[1], "previous_row": old_row,
                    "selected_row": row_number, "previous_content": old_content,
                    "selected_content": content, "conflicting": old_content != content,
                })
            records[identity] = row_number, content
    return records, duplicates, sheet_name


def _save_output(workbook, source: Path, destination: Path, *, changed: bool) -> None:
    if not changed and source.resolve() == destination.resolve():
        return
    destination.parent.mkdir(parents=True, exist_ok=True)
    handle, temporary = tempfile.mkstemp(dir=destination.parent, suffix=destination.suffix)
    os.close(handle)
    temporary_path = Path(temporary)
    try:
        if changed:
            workbook.save(temporary_path)
        else:
            shutil.copy2(source, temporary_path)
        os.replace(temporary_path, destination)
    finally:
        temporary_path.unlink(missing_ok=True)


def _sync_file(source: Path, destination: Path, records, mapping: ColumnMapping,
               count: int, sheet: str | None, header_rows: int,
               fill_blank_only: bool, allow_blank_write: bool, result: FileResult,
               identities: set | None = None) -> None:
    key_col, match_col, content_col = mapping.indexes(count)
    workbook = load_workbook_for_editing(source)
    try:
        worksheet = _sheet(workbook, sheet)
        result.sheet = worksheet.title
        if max(key_col, match_col) > worksheet.max_column:
            raise ValueError("目标工作表缺少配置的 Key 或原文列。")
        # Sparse row selection avoids traversing millions of formatted empty rows.
        rows = sorted({cell.row for cell in iter_value_cells(worksheet)
                       if cell.column in {key_col, match_col} and cell.row > header_rows})
        for row in rows:
            key_cell, match_cell = worksheet.cell(row, key_col), worksheet.cell(row, match_col)
            if key_cell.data_type == "f" or match_cell.data_type == "f":
                raise ValueError(f"目标第 {row} 行的 Key/原文包含公式，请先转换为值。")
            identity = row_identity(key_cell.value, match_cell.value)
            if identity is None:
                result.invalid_rows += 1
                continue
            if identities is not None:
                identities.add(identity)
            if identity not in records:
                result.unmatched_rows.append(row)
                continue
            result.matched_rows += 1
            for offset, value in enumerate(records[identity][1]):
                if not allow_blank_write and is_blank(value):
                    result.blank_source_cells += 1
                    continue
                cell = worksheet.cell(row, content_col + offset)
                if fill_blank_only and not is_blank(cell.value):
                    result.occupied_cells += 1
                    continue
                if (cell.value == value and cell.data_type not in {"f", "e"}
                        or cell.value is None and value == ""):
                    result.unchanged_cells += 1
                    continue
                # Transferred content is text, including strings beginning with '='
                # or matching an Excel error token. Do not reinterpret it as a formula.
                cell.value = value
                cell.data_type = "s"
                result.updated_cells += 1
        _save_output(workbook, source, destination, changed=result.updated_cells > 0)
        result.output = str(destination)
        result.status = "updated" if result.updated_cells else "unchanged"
    finally:
        workbook.close()


def sync_master_to_targets(
    master_file: str | Path, target_dir: str | Path, *,
    output_dir: str | Path | None = None,
    master_columns: ColumnMapping = ColumnMapping("B", "C", "D"),
    target_columns: ColumnMapping = ColumnMapping("A", "B", "C"),
    column_count: int = 1, master_sheet: str | None = None, target_sheet: str | None = None,
    master_header_rows: int = 1, target_header_rows: int = 1,
    fill_blank_only: bool = False, allow_blank_write: bool = False,
    compatibility_resave: bool = False,
    workers: int = 2,
    progress_callback: Callable[[int, int], None] | None = None,
    log_callback=None,
) -> SyncSummary:
    """Update targets in place by default; an explicit output writes a new tree."""
    validate_workers(workers)
    emit_log(log_callback, f"开始同步：{master_file} → {target_dir}")
    if not isinstance(column_count, int) or isinstance(column_count, bool) or not 1 <= column_count <= 16384:
        raise ValueError("更新列数必须是 1–16384 的整数。")
    for header_rows in (master_header_rows, target_header_rows):
        if not isinstance(header_rows, int) or isinstance(header_rows, bool) or not 0 <= header_rows < 1048576:
            raise ValueError("表头行数必须是 0–1048575 的整数。")
    master_columns.indexes(column_count)
    target_columns.indexes(column_count)
    master = Path(master_file).expanduser().resolve()
    folder = Path(target_dir).expanduser().resolve()
    if not master.is_file() or master.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("Master 必须是存在的 .xlsx/.xlsm 文件。")
    if not folder.is_dir():
        raise ValueError(f"小表目录不存在: {folder}")
    output, inplace = optional_output_directory(output_dir, default=folder, inputs=(folder, master))
    emit_log(log_callback, "保存方式：原位更新小表" if inplace else f"保存方式：另存到 {output}")
    files, skipped = _input_files(folder, None if inplace else output, master)
    if not files:
        raise ValueError("目录中没有可同步的 .xlsx/.xlsm 小表。")
    records, duplicates, master_sheet_name = _read_master(
        master, master_columns, column_count, master_sheet, master_header_rows, log_callback=log_callback,
    )
    if not records:
        raise ValueError("Master 中没有有效的 Key + 原文。")
    emit_log(log_callback, f"Master {master_sheet_name}：{len(records)} 个有效身份；待处理 {len(files)} 个小表")
    for duplicate in duplicates:
        emit_log(log_callback, "Master 重复身份（后行优先）", duplicate)
    for result in skipped:
        log_result(log_callback, result)
    summary = SyncSummary(output, len(records), duplicates, skipped)
    # Linked inputs may observe another target's writes. COM remains on its owner
    # thread, including the preceding sync, to retain the original save sequence.
    if compatibility_resave or any(path.is_symlink() for path in files):
        workers = 1
    workers = min(workers, len(files))

    def process(source):
        relative = source.relative_to(folder)
        result = FileResult(str(relative))
        try:
            _sync_file(source, output / relative, records, target_columns, column_count,
                       target_sheet, target_header_rows, fill_blank_only, allow_blank_write, result)
        except Exception as exc:
            result.error = str(exc)
            result.updated_cells = 0
        return result

    session = None
    if compatibility_resave:
        emit_log(log_callback, "正在启动 Excel 兼容性重存会话…")
        from tools.excel_com import edit_copy, excel_session
        session = excel_session()
        application = session.__enter__()  # Preflight before reserving or writing outputs.
    try:
        if not inplace:
            create_output_directory(output)
        for index, (result, error) in enumerate(ordered_results(process, files, workers), 1):
            if error is not None:
                result = FileResult(str(files[index - 1].relative_to(folder)), error=str(error))
            summary.files.append(result)
            if compatibility_resave and result.status == "updated":
                try:
                    destination = Path(result.output)
                    edit_copy(application, destination, destination)
                except Exception as exc:
                    result.postprocess_error = f"Excel 重存失败，保留已同步内容: {exc}"
            log_result(log_callback, result)
            if progress_callback:
                progress_callback(index, len(files))
    finally:
        if session is not None:
            try:
                session.__exit__(None, None, None)
            except Exception as exc:
                summary.warnings.append(f"Excel 会话未正常结束: {exc}")
    log_summary(log_callback, summary, f"更新 {summary.updated_cells} 格")
    return summary
