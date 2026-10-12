"""Group filename-labelled rows and export complete per-file contents."""
from __future__ import annotations

from copy import copy
from dataclasses import dataclass
import os
from pathlib import Path
import tempfile

from tools.excel_file_ops import create_output_directory
from .grouping import (
    FileGroup, FileTexts, GroupingExport, ScanIssue, _digest, _empty, _fingerprint,
    _groups, _header_columns,
)


@dataclass(frozen=True)
class MergedGroupingResult:
    input_file: Path
    filename_column: str
    source_header: str
    target_header: str
    files: tuple[FileTexts, ...]
    groups: tuple[FileGroup, ...]
    issues: tuple[ScanIssue, ...]
    fingerprint: tuple[int, ...]
    digest: str
    row_count: int

    @property
    def can_export(self) -> bool:
        return bool(self.files) and not self.issues


def _filename(value) -> str:
    return "" if _empty(value) else Path(str(value).strip().replace("\\", "/")).as_posix()


def _value(row, column):
    # Dimensionless tables omit trailing cells; all other columns stay intact.
    return row[column].value if column < len(row) else None


def _filename_index(column):
    from openpyxl.utils import column_index_from_string
    column = column.strip().upper()
    try:
        index = column_index_from_string(column)
    except ValueError as exc:
        raise ValueError("文件名列必须是 A–XFD 范围内的列字母。") from exc
    if not 1 <= index <= 16384:
        raise ValueError("文件名列必须是 A–XFD 范围内的列字母。")
    return column, index - 1


def scan_merged_table(input_file: str | Path, *, filename_column: str = "A",
                      source_header: str = "MsgStr", target_header: str = "fr",
                      progress_callback=None) -> MergedGroupingResult:
    from openpyxl import load_workbook

    if not str(input_file or "").strip():
        raise ValueError("请选择合并表。")
    path = Path(input_file).expanduser().resolve()
    if not path.is_file() or path.suffix.casefold() not in {".xlsx", ".xlsm"}:
        raise ValueError("请选择存在的 .xlsx/.xlsm 合并表。")
    source_header, target_header = source_header.strip(), target_header.strip()
    if not source_header or not target_header or source_header == target_header:
        raise ValueError("原文和译文表头不能为空，且必须不同。")
    filename_column, filename_index = _filename_index(filename_column)
    before, digest = _fingerprint(path), _digest(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    records, issues, matched, row_count = {}, [], False, 0
    try:
        for sheet_index, sheet in enumerate(workbook.worksheets, 1):
            rows = sheet.iter_rows()
            headers = next(rows, ())
            try:
                columns = _header_columns(headers, source_header, target_header, sheet.title)
                if columns is None:
                    continue
                matched = True
                if filename_index >= len(headers) or filename_index in columns:
                    raise ValueError(f"工作表 {sheet.title} 的文件名列必须存在，且与原文、译文列不同。")
                source_index, target_index = columns
                for row_number, row in enumerate(rows, 2):
                    if all(_empty(cell.value) for cell in row):
                        continue
                    name = _filename(_value(row, filename_index))
                    if not name:
                        raise ValueError(f"工作表 {sheet.title} 第 {row_number} 行有内容但缺少文件名。")
                    formula = next((cell for cell in row if cell.data_type == "f"), None)
                    if formula is not None:
                        raise ValueError(f"工作表 {sheet.title} 的 {formula.coordinate} 含公式；分组导出会重排行，请先将公式转换为值。")
                    record = records.setdefault(name.casefold(), [Path(name), set(), 0])
                    row_count += 1
                    source, target = _value(row, source_index), _value(row, target_index)
                    if not _empty(source) and _empty(target):
                        record[1].add(str(source))
                        record[2] += 1
            except ValueError as exc:
                issues.append(ScanIssue(Path(sheet.title), str(exc)))
            if progress_callback:
                progress_callback(sheet_index, len(workbook.worksheets))
    finally:
        workbook.close()
    if not matched:
        issues.append(ScanIssue(Path(path.name), f"未找到首行同时包含 {source_header} 和 {target_header} 的工作表。"))
    if _fingerprint(path) != before:
        raise ValueError("合并表在扫描期间发生变化，请重新扫描。")
    files = tuple(FileTexts(record[0], frozenset(record[1]), record[2]) for record in sorted(
        records.values(), key=lambda record: (record[0].as_posix().casefold(), record[0].as_posix())))
    return MergedGroupingResult(path, filename_column, source_header, target_header, files,
                                _groups(files), tuple(issues), before, digest, row_count)


def _check_input(result):
    try:
        if _fingerprint(result.input_file) != result.fingerprint or _digest(result.input_file) != result.digest:
            raise ValueError("合并表已变化，请重新扫描。")
    except OSError as exc:
        raise ValueError(f"合并表不可读，请重新扫描：{exc}") from exc


def _copy_row(row, sheet, styles):
    from openpyxl.cell import WriteOnlyCell

    copied = []
    for cell in row:
        has_style = getattr(cell, "has_style", False)
        if cell.value is None and not has_style:
            copied.append(None)
            continue
        new_cell = WriteOnlyCell(sheet, value=cell.value)
        new_cell.data_type = cell.data_type  # Preserve literal text beginning with '='.
        if has_style:
            key = tuple(cell.style_array)
            if key not in styles:
                styles[key] = (copy(cell.font), copy(cell.fill), copy(cell.border),
                               copy(cell.alignment), cell.number_format, copy(cell.protection))
            (new_cell.font, new_cell.fill, new_cell.border, new_cell.alignment,
             new_cell.number_format, new_cell.protection) = styles[key]
        copied.append(new_cell)
    return copied


def _publish_book(book, result, output, destination):
    def check_destination():
        if output.resolve() != output or not destination.resolve().is_relative_to(output):
            raise ValueError("输出路径已变化，请重新导出。")
        if destination.exists() or destination.is_symlink():
            raise FileExistsError(f"输出文件已存在，未覆盖：{destination.name}")

    check_destination()
    descriptor, name = tempfile.mkstemp(prefix=".qatools-", suffix=".tmp", dir=output)
    os.close(descriptor)
    temporary = Path(name)
    try:
        book.save(temporary)
        _check_input(result)
        check_destination()
        if os.name == "nt":
            os.rename(temporary, destination)
        else:
            os.link(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def export_merged_groups(result: MergedGroupingResult, output_dir: str | Path,
                         *, progress_callback=None) -> GroupingExport:
    from openpyxl import Workbook, load_workbook

    if not result.can_export:
        raise ValueError("扫描不完整或没有可导出的文件，请修正问题后重新扫描。")
    if not str(output_dir or "").strip():
        raise ValueError("请选择新输出目录。")
    output = Path(output_dir).expanduser().resolve()
    if output.exists() or output.is_symlink() or result.input_file.is_relative_to(output):
        raise ValueError("输出目录必须是尚不存在的新目录，且不能覆盖输入。")
    _check_input(result)
    _, filename_index = _filename_index(result.filename_column)
    owners = {file.as_posix().casefold(): group.name for group in result.groups for file in group.files}
    books, sheets, styles, copied = {}, {}, {}, 0
    workbook = load_workbook(result.input_file, read_only=True, data_only=False)
    try:
        # Reserve before writing; each workbook is published only after a complete save.
        create_output_directory(output)
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows()
            headers = next(rows, ())
            if _header_columns(headers, result.source_header, result.target_header, sheet.title) is None:
                continue
            for row in rows:
                if all(_empty(cell.value) for cell in row):
                    continue
                key = _filename(_value(row, filename_index)).casefold()
                group = owners.get(key)
                if group is None:
                    raise ValueError("文件名或合并表内容已变化，请重新扫描。")
                if group not in books:
                    books[group] = Workbook(write_only=True)
                sheet_key = (group, sheet.title)
                if sheet_key not in sheets:
                    target = books[group].create_sheet(sheet.title)
                    target.freeze_panes = "A2"
                    sheets[sheet_key] = target
                    target.append(_copy_row(headers, target, styles))
                target = sheets[sheet_key]
                target.append(_copy_row(row, target, styles))
        _check_input(result)
        for group in result.groups:
            book = books[group.name]
            destination = output / f"{group.name.replace(' ', '')}.xlsx"
            _publish_book(book, result, output, destination)
            copied += 1
            if progress_callback:
                progress_callback(copied, len(result.groups))
        _check_input(result)
    except Exception as exc:
        raise RuntimeError(f"导出中止：{exc}\n已导出 {copied} 个工作簿，结果保留在：{output}") from exc
    finally:
        workbook.close()
        for book in books.values():
            for sheet in book.worksheets:
                if not sheet.closed:
                    sheet.close()
            book.close()
    return GroupingExport(output, copied)
