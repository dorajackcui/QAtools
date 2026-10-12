"""Read-only workbook scans, connected file groups and guarded exports."""
from __future__ import annotations

from dataclasses import dataclass
import hashlib
import math
import os
from pathlib import Path
import re
import shutil
import tempfile

from tools.excel_file_ops import create_output_directory, excel_files


@dataclass(frozen=True)
class ScannedFile:
    path: Path  # Relative to the scan root.
    sources: frozenset[str]
    pending_rows: int
    fingerprint: tuple[int, ...]
    digest: str


@dataclass(frozen=True)
class FileGroup:
    name: str
    files: tuple[Path, ...]
    unique_chars: int


@dataclass(frozen=True)
class ScanIssue:
    path: Path
    message: str


@dataclass(frozen=True)
class GroupingResult:
    source_dir: Path
    source_header: str
    target_header: str
    files: tuple[ScannedFile, ...]
    groups: tuple[FileGroup, ...]
    issues: tuple[ScanIssue, ...]

    @property
    def can_export(self) -> bool:
        return bool(self.files) and not self.issues


@dataclass(frozen=True)
class GroupingExport:
    output_dir: Path
    file_count: int


def _fingerprint(path: Path) -> tuple[int, ...]:
    info = path.stat()
    return (info.st_dev, info.st_ino, info.st_size, info.st_mtime_ns, info.st_ctime_ns)


def _digest(path: Path) -> str:
    with path.open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def _empty(value: object) -> bool:
    return (value is None or isinstance(value, str) and not value.strip()
            or isinstance(value, float) and math.isnan(value))


def _read_file(path: Path, root: Path, source_header: str, target_header: str) -> ScannedFile:
    from openpyxl import load_workbook

    before, digest = _fingerprint(path), _digest(path)
    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        sources, pending, matched = set(), 0, False
        for sheet in workbook.worksheets:
            rows = sheet.iter_rows()
            headers = next(rows, ())
            source_cols = [index for index, cell in enumerate(headers)
                           if str(cell.value or "").strip() == source_header]
            target_cols = [index for index, cell in enumerate(headers)
                           if str(cell.value or "").strip() == target_header]
            if not source_cols and not target_cols:
                continue
            if len(source_cols) != 1 or len(target_cols) != 1:
                raise ValueError(f"工作表 {sheet.title} 的首行缺少或重复表头：{source_header} / {target_header}")
            matched = True
            source_col, target_col = source_cols[0], target_cols[0]
            for row in rows:
                if max(source_col, target_col) >= len(row):
                    continue
                source, target = row[source_col], row[target_col]
                if _empty(source.value) or not _empty(target.value):
                    continue
                if source.data_type == "f":
                    raise ValueError(f"工作表 {sheet.title} 第 {source.row} 行的待翻译原文是公式，请先转换为值。")
                sources.add(str(source.value))
                pending += 1
        if not matched:
            raise ValueError(f"未找到首行同时包含 {source_header} 和 {target_header} 的工作表。")
    finally:
        workbook.close()
    if _fingerprint(path) != before:
        raise ValueError("文件在扫描期间发生变化，请重新扫描。")
    return ScannedFile(path.relative_to(root), frozenset(sources), pending, before, digest)


def _groups(files: tuple[ScannedFile, ...]) -> tuple[FileGroup, ...]:
    # An inverted index joins file owners directly; no quadratic file-pair graph.
    parents = list(range(len(files)))

    def find(index):
        while parents[index] != index:
            parents[index] = parents[parents[index]]
            index = parents[index]
        return index

    owners = {}
    for index, file in enumerate(files):
        for source in file.sources:
            previous = owners.setdefault(source, index)
            parents[find(index)] = find(previous)
    components = {}
    for index in range(len(files)):
        components.setdefault(find(index), []).append(files[index])
    linked = sorted((component for component in components.values() if len(component) > 1),
                    key=lambda items: (-len(items), items[0].path.as_posix().casefold(), items[0].path.as_posix()))
    independent = [component[0] for component in components.values() if len(component) == 1]

    def make_group(name, items):
        sources = set().union(*(item.sources for item in items))
        chars = sum(len(re.findall(r"[\u4e00-\u9fa5]", source)) for source in sources)
        return FileGroup(name, tuple(item.path for item in items), chars)

    groups = [make_group(f"文件群 {index}", items) for index, items in enumerate(linked, 1)]
    if independent:
        groups.append(make_group("其他", independent))
    return tuple(groups)


def scan_file_groups(source_dir: str | Path, *, source_header: str = "MsgStr", target_header: str = "fr",
                     progress_callback=None) -> GroupingResult:
    if not str(source_dir or "").strip():
        raise ValueError("请选择来源目录。")
    source_header, target_header = source_header.strip(), target_header.strip()
    if not source_header or not target_header or source_header == target_header:
        raise ValueError("原文和译文表头不能为空，且必须不同。")
    root = Path(source_dir).expanduser().resolve()
    paths, skipped = excel_files(root, extensions={".xlsx", ".xlsm"}, skip_links=True)
    if not paths and not skipped:
        raise ValueError("目录中没有支持的 .xlsx/.xlsm 文件。")
    files, issues = [], [ScanIssue(path.relative_to(root), "不支持 .xls/.xlsb，请先转换为 .xlsx/.xlsm。")
                        for path in skipped]
    for index, path in enumerate(paths, 1):
        try:
            files.append(_read_file(path, root, source_header, target_header))
        except Exception as exc:
            issues.append(ScanIssue(path.relative_to(root), str(exc)))
        if progress_callback:
            progress_callback(index, len(paths))
    scanned = tuple(files)
    return GroupingResult(root, source_header, target_header, scanned, _groups(scanned), tuple(issues))


def _check_file(result: GroupingResult, file: ScannedFile) -> Path:
    path = result.source_dir / file.path
    if (not path.resolve().is_relative_to(result.source_dir)
            or _fingerprint(path) != file.fingerprint or _digest(path) != file.digest):
        raise ValueError(f"来源文件已变化，请重新扫描：{file.path}")
    return path


def _check_scan(result: GroupingResult) -> None:
    try:
        paths, skipped = excel_files(result.source_dir, extensions={".xlsx", ".xlsm"}, skip_links=True)
        if skipped or tuple(path.relative_to(result.source_dir) for path in paths) != tuple(file.path for file in result.files):
            raise ValueError("来源文件列表已变化，请重新扫描。")
        for file in result.files:
            _check_file(result, file)
    except OSError as exc:
        raise ValueError(f"来源文件不可读，请重新扫描：{exc}") from exc


def _copy_file(result: GroupingResult, file: ScannedFile, output: Path, relative: Path) -> None:
    source = _check_file(result, file)
    destination = output / relative

    def check_destination():
        if output.resolve() != output or not destination.resolve().is_relative_to(output):
            raise ValueError("输出路径已变化，请重新导出。")
        if destination.exists() or destination.is_symlink():
            raise FileExistsError(f"输出文件已存在，未覆盖：{relative}")

    check_destination()
    destination.parent.mkdir(parents=True, exist_ok=True)
    check_destination()
    descriptor, name = tempfile.mkstemp(prefix=".qatools-", suffix=".tmp", dir=destination.parent)
    os.close(descriptor)
    temporary = Path(name)
    try:
        shutil.copyfile(source, temporary)
        info = source.stat()
        os.utime(temporary, ns=(info.st_atime_ns, info.st_mtime_ns))
        _check_file(result, file)
        if _digest(temporary) != file.digest:
            raise ValueError(f"复制内容校验失败：{file.path}")
        check_destination()
        if os.name == "nt":
            os.rename(temporary, destination)
        else:
            os.link(temporary, destination)
    finally:
        temporary.unlink(missing_ok=True)


def export_file_groups(result: GroupingResult, output_dir: str | Path, *, progress_callback=None) -> GroupingExport:
    if not result.can_export:
        raise ValueError("扫描不完整或没有可导出的文件，请修正问题后重新扫描。")
    if not str(output_dir or "").strip():
        raise ValueError("请选择新输出目录。")
    output = Path(output_dir).expanduser().resolve()
    source = result.source_dir
    if source == output or source.is_relative_to(output) or output.is_relative_to(source):
        raise ValueError("来源目录和输出目录不能相同或互相包含。")
    if output.exists() or output.is_symlink():
        raise ValueError(f"输出目录已存在，请指定新目录：{output}")
    _check_scan(result)  # Check all files before creating any output.
    entries = {file.path: file for file in result.files}
    destinations = [(entries[path], Path(group.name.replace(" ", "")) / path)
                    for group in result.groups for path in group.files]
    # Protect Windows exports even when sources were prepared on a case-sensitive filesystem.
    keys = [relative.as_posix().casefold() for _, relative in destinations]
    if len(set(keys)) != len(keys):
        raise ValueError("分组中有忽略大小写后的输出路径冲突。")
    key_set = set(keys)
    if any(parent.as_posix().casefold() in key_set for _, path in destinations for parent in path.parents):
        raise ValueError("分组中有文件与目录的输出路径冲突。")
    create_output_directory(output)
    copied = 0
    try:
        for file, relative in destinations:
            _copy_file(result, file, output, relative)
            copied += 1
            if progress_callback:
                progress_callback(copied, len(destinations))
        _check_scan(result)
    except Exception as exc:
        raise RuntimeError(f"导出中止：{exc}\n已导出 {copied} 个文件，结果保留在：{output}") from exc
    return GroupingExport(output, copied)
