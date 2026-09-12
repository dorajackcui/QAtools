"""Exercise Qt buttons, real background workers and the resulting workbooks."""

from __future__ import annotations

import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from toolshub_gui import ToolshubApp


class GuiFunctionalTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        for module in ("tools.header_aliases", "tools.workflow.gui_options"):
            self.enterContext(patch(f"{module}.default_header_aliases_path", return_value=self.root / "aliases.json"))
        self.enterContext(patch("tools.tb_projects.default_tb_projects_path", return_value=self.root / "projects.json"))
        self.errors = self.enterContext(patch.object(QMessageBox, "critical"))
        self.warnings = self.enterContext(patch.object(QMessageBox, "warning"))
        self.enterContext(patch.object(QMessageBox, "information"))
        self.callback_errors = self.enterContext(patch("sys.excepthook"))
        self.window = ToolshubApp(show_window=False)
        self.addCleanup(self.dispose_window)

    def dispose_window(self):
        self.window.close()
        self.window.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)

    def workbook(self, relative, rows):
        path = self.root / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        try:
            book.active.title = "Data"
            for row in rows:
                book.active.append(row)
            book.save(path)
        finally:
            book.close()
        return path

    def read(self, path, sheet="Data"):
        book = load_workbook(path)
        try:
            return list(book[sheet].values)
        finally:
            book.close()

    def edit(self, path, sheet, cell, value):
        book = load_workbook(path)
        try:
            book[sheet][cell] = value
            book.save(path)
        finally:
            book.close()

    def run_button(self, page, button, *, expect_error=False):
        self.assertTrue(button.isEnabled())
        button.click()
        deadline = time.monotonic() + 60
        while page.has_running_tasks() and time.monotonic() < deadline:
            QTest.qWait(10)
        self.assertFalse(page.has_running_tasks(), "GUI worker did not complete")
        self.app.processEvents()
        self.callback_errors.assert_not_called()
        self.assertTrue(button.isEnabled(), "Action was not restored after completion")
        if expect_error:
            self.errors.assert_called()
            self.errors.reset_mock()
        else:
            self.assertEqual(self.errors.call_count, 0, str(self.errors.call_args_list))
            self.assertEqual(self.warnings.call_count, 0, str(self.warnings.call_args_list))

    def test_qa_all_checks_and_revision_roundtrip(self):
        source = self.workbook("qa.xlsx", [
            ["key", "source", "target"],
            ["a", "<b>Hello {name}</b> 123", "你好 456"],
            ["b", "<b>Hello {name}</b> 123", "Autre"],
        ])
        original = source.read_bytes()
        page = self.window.tool_frames["workflow"]
        page.load_input_file(str(source))
        page.set_all_tasks(True)
        page.tag_check_order.setChecked(True)
        self.run_button(page, page.run_button)
        report = Path(page.last_workflow_output_path)
        problems = self.read(report, "问题处理")
        self.assertGreater(len(problems), 1)
        row = problems[1][0]
        self.edit(report, "问题处理", "D2", "Corrected translation")
        output = self.root / "revised.xlsx"
        with patch("PySide6.QtWidgets.QFileDialog.getOpenFileName", return_value=(str(report), "")), \
             patch("PySide6.QtWidgets.QFileDialog.getSaveFileName", return_value=(str(output), "")):
            self.run_button(page, page.revision_button)
        self.assertEqual(self.read(output)[row - 1][2], "Corrected translation")
        self.assertEqual(source.read_bytes(), original)

    def test_phraseloom_export_and_restore(self):
        source = self.workbook("strings-source.xlsx", [
            ["key", "source", "target"], ["a", "Hello Alice", None],
            ["b", "Hello Alice", None], ["c", "Hello Bob", "Bonjour Bob"],
        ])
        page = self.window.tool_frames["phraseloom"]
        with patch("phraseloom.qt_page._choose_excel", return_value=str(source)):
            page.input_picker.choose_button.click()
        self.run_button(page, page.export_button)
        strings = self.root / "strings-source_strings.xlsx"
        self.assertEqual(len(self.read(strings, "strings")), 2)
        headers = self.read(strings, "strings")[0]
        target_column = get_column_letter(headers.index("target") + 1)
        self.edit(strings, "strings", f"{target_column}2", "Bonjour Alice")
        with patch("phraseloom.qt_page._choose_excel", return_value=str(strings)):
            self.run_button(page, page.restore_button)
        restored = self.read(self.root / "strings-source_translated.xlsx")
        self.assertEqual([row[2] for row in restored[1:]], ["Bonjour Alice", "Bonjour Alice", "Bonjour Bob"])

    def test_batch_split_and_restore(self):
        rows = [["source", "target"], ["First", "Premier"], ["Second", "Deuxième"]]
        source = self.workbook("batch.xlsx", rows)
        page = self.window.tool_frames["excel_batcher"]
        with patch("tools.excel_batcher.qt_page._choose_excel", return_value=str(source)):
            page.split_input.choose_button.click()
        page.batch_size.setValue(1)
        self.run_button(page, page.split_button)
        directory = Path(page.split_output.path())
        batches = sorted(directory.glob("batch_batch_*.xlsx"))
        self.assertEqual(len(batches), 2)
        self.edit(batches[1], "Data", "B2", "Updated")
        page.tabs.setCurrentIndex(1)
        page.restore_dir.set_path(str(directory))
        page.restore_output.set_path(str(self.root / "restored.xlsx"))
        self.run_button(page, page.restore_button)
        rows[2][1] = "Updated"
        self.assertEqual(self.read(self.root / "restored.xlsx"), [tuple(row) for row in rows])

    def test_nbsp_restores_result_column(self):
        source = self.workbook("french.xlsx", [["source", "target"], ["Hello!", "Bonjour !"]])
        page = self.window.tool_frames["french_nbsp"]
        page.load_input_file(str(source))
        page.result_column.setText("C")
        self.run_button(page, page.run_button)
        output = next(path for path in self.root.glob("*.xlsx") if path != source)
        self.assertEqual(self.read(output)[1][1:], ("Bonjour !", "Bonjour\u00a0!"))

    def test_merger_headers_and_source_files(self):
        self.workbook("merge/a.xlsx", [["source", "target"], ["Hello", "Bonjour"]])
        self.workbook("merge/b.xlsx", [["source", "target"], ["Bye", "Salut"]])
        page = self.window.tool_frames["excel_merger"]
        page.input_dir.set_path(str(self.root / "merge"))
        page.keep_headers.setChecked(True)
        self.run_button(page, page.run_button)
        output = next(self.root.glob("merge_merged_active_sheet_*.xlsx"))
        rows = self.read(output, "MergedData")
        self.assertEqual([row[1:] for row in rows], [("source", "target"), ("Hello", "Bonjour"), ("source", "target"), ("Bye", "Salut")])

    def test_xbench_converts_report(self):
        source = self.workbook("xbench.xlsx", [
            [None, None, "Source", "Target", "Comments", "Metadata"],
            ["Key Term Mismatch (Hello / Bonjour)"],
            [None, None, "Hello!", "Salut!", None, "key1\nfile.xlsx"],
        ])
        page = self.window.tool_frames["xbench_report"]
        with patch("tools.xbench_report_transformer.qt_page._choose_excel", return_value=str(source)):
            page.input_picker.choose_button.click()
        self.run_button(page, page.run_button)
        output = next(path for path in self.root.glob("*.xlsx") if path != source)
        rows = self.read(output, "Xbench QA整理")
        self.assertEqual(rows[1][:4], ("file.xlsx", "key1", "Hello!", "Salut!"))

    def sync_roundtrip(self, *, resave=False):
        master = self.workbook("master.xlsx", [["key", "source", "target"], ["a", "Hello", "Bonjour"]])
        self.workbook("targets/a.xlsx", [["key", "source", "target"], ["a", "Hello", None]])
        sync = self.window.tool_frames["content_sync"]
        for index in (0, 1):
            sync.tabs.setCurrentIndex(index)
            page = sync.tabs.currentWidget()
            for widget, column in zip(page.master_columns, ("A", "B", "C")):
                widget.setText(column)
            page.master_picker.set_path(str(master))
            page.target_picker.set_path(str(self.root / ("targets" if index == 0 else "synced")))
            output = self.root / ("synced" if index == 0 else "returned")
            page.output_picker.set_path(str(output))
            page.compatibility_resave.setChecked(resave and index == 0)
            if index == 1:
                self.edit(self.root / "synced/a.xlsx", "Data", "C2", "Updated")
            self.run_button(page, page.run_button)
            result = output / ("a.xlsx" if index == 0 else "master.xlsx")
            self.assertEqual(self.read(result)[1][2], "Bonjour" if index == 0 else "Updated")
            page.logs_button.click()
            self.app.processEvents()
            self.assertTrue(page.logs.text.toPlainText())
            page.logs.close()
        self.assertEqual(self.read(master)[1][2], "Bonjour")

    def test_content_sync_both_directions_and_logs(self):
        self.sync_roundtrip()

    def test_deep_replace_and_untranslated_statistics(self):
        source = self.workbook("sources/a.xlsx", [["key", "source", "target"], ["a", "Hello world", None]])
        self.workbook("targets/a.xlsx", [["key", "source", "target"], ["a", "Old", "Old"]])
        page = self.window.tool_frames["deep_replace"]
        page.source_picker.set_path(str(source.parent))
        page.input_picker.set_path(str(self.root / "targets"))
        page.output_picker.set_path(str(self.root / "replaced"))
        self.run_button(page, page.run_button)
        self.assertEqual((self.root / "replaced/a.xlsx").read_bytes(), source.read_bytes())
        stats = self.window.tool_frames["untranslated_stats"]
        stats.input_picker.set_path(str(self.root / "replaced"))
        stats.output_picker.set_path(str(self.root / "statistics"))
        stats.mode.setCurrentIndex(1)
        self.run_button(stats, stats.run_button)
        rows = self.read(self.root / "statistics/未翻译统计.xlsx", "未翻译统计")
        self.assertEqual(rows[1][1:], (2, 1, 2, 1))

    def test_failed_worker_can_be_retried_and_settings_are_saved(self):
        page = self.window.tool_frames["phraseloom"]
        page.input_picker.set_path(str(self.root / "missing.xlsx"))
        self.run_button(page, page.export_button, expect_error=True)
        source = self.workbook("retry.xlsx", [["source", "target"], ["Hello", None]])
        page.input_picker.set_path(str(source))
        self.run_button(page, page.export_button)
        self.assertTrue((self.root / "retry_strings.xlsx").is_file())
        settings = self.window.tool_frames["settings"]
        settings.source_aliases.setPlainText("Original")
        settings.target_aliases.setPlainText("Translation")
        settings.save_button.click()
        settings.reload_aliases()
        self.assertEqual(settings.source_aliases.toPlainText(), "Original")
        self.assertEqual(settings.target_aliases.toPlainText(), "Translation")

    @unittest.skipUnless(os.environ.get("QATOOLS_TEST_EXCEL_COM") == "1", "opt-in real Microsoft Excel GUI test")
    def test_native_excel_column_actions_and_resave(self):
        source = self.workbook("native/a.xlsx", [["key", "source", "target"], ["a", "Hello", "Bonjour"]])
        original = source.read_bytes()
        page = self.window.tool_frames["column_tools"]
        page.input_picker.set_path(str(source.parent))
        for index, action in enumerate(("clear", "insert", "delete")):
            with self.subTest(action=action):
                page.action.setCurrentIndex(index)
                page.output_picker.set_path(str(self.root / action))
                self.run_button(page, page.run_button)
                rows = self.read(self.root / action / "a.xlsx")
                if action == "clear":
                    self.assertIsNone(rows[1][2])
                elif action == "insert":
                    self.assertEqual(rows[0][2], "Translation")
                    self.assertEqual(rows[1][3], "Bonjour")
                else:
                    self.assertEqual(rows[1], ("a", "Hello"))
        page = self.window.tool_frames["compatibility"]
        page.input_picker.set_path(str(source.parent))
        page.output_picker.set_path(str(self.root / "resaved"))
        self.run_button(page, page.run_button)
        self.assertEqual(self.read(self.root / "resaved/a.xlsx")[1][2], "Bonjour")
        self.assertEqual(source.read_bytes(), original)

    @unittest.skipUnless(os.environ.get("QATOOLS_TEST_EXCEL_COM") == "1", "opt-in real Microsoft Excel GUI test")
    def test_content_sync_with_native_excel_resave(self):
        self.sync_roundtrip(resave=True)


if __name__ == "__main__":
    unittest.main()
