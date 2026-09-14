from __future__ import annotations

import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools.excel_output import PROBLEM_BASE_HEADERS
from tools.workflow.review_sheet import (
    WORKFLOW_METADATA_SHEET_NAME,
    WORKFLOW_REVIEW_HEADERS,
    WORKFLOW_REVIEW_SHEET_NAME,
    read_review_metadata,
    write_review_sheet,
)
from tools.workflow.revision_applier import (
    apply_workflow_revisions,
    build_default_revised_output_path,
)


class WorkflowMetadataTests(unittest.TestCase):
    def setUp(self) -> None:
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.report_path = self.root / "report.xlsx"
        self.workbook = Workbook()
        self.addCleanup(self.workbook.close)
        data = self.workbook.active
        data.title = "Data"
        data.append(["source", "target"])
        for number in range(2, 5):
            data.append([f"Source {number}", f"Target {number}"])

    def write_report(self, *, empty: bool = False) -> None:
        issues = self.workbook.create_sheet("Issues")
        issues.append(PROBLEM_BASE_HEADERS)
        if not empty:
            for number in range(2, 5):
                issues.append([number, f"Source {number}", f"Target {number}", "缺少 Tag"])
        try:
            write_review_sheet(
                self.workbook, current_sheet_name="Data", input_file=self.root / "original.xlsx",
                source_column="A", target_column="B", start_row=2,
                problem_sheets=[("Tag 检查", "Issues")],
                generated_sheet_names=[WORKFLOW_REVIEW_SHEET_NAME], remove_term_helper=False,
            )
        finally:
            del self.workbook["Issues"]
        self.workbook.save(self.report_path)

    def test_metadata_survives_review_row_and_column_cleanup_and_roundtrip(self) -> None:
        self.write_report()
        report = load_workbook(self.report_path)
        try:
            review = report[WORKFLOW_REVIEW_SHEET_NAME]
            self.assertEqual(tuple(next(review.values)), WORKFLOW_REVIEW_HEADERS)
            self.assertEqual(review.max_column, 6)
            self.assertEqual(review.max_row, 4)
            self.assertFalse(review.protection.sheet)
            metadata_sheet = report[WORKFLOW_METADATA_SHEET_NAME]
            self.assertEqual(metadata_sheet.sheet_state, "veryHidden")
            self.assertTrue(metadata_sheet.protection.sheet)
            metadata = read_review_metadata(report)
            self.assertEqual(metadata["schema_version"], "3")
            review.delete_cols(7, 2)
            review.delete_rows(2)
            # Reverse the remaining issues, as with a user sorting the review table.
            rows = list(review.iter_rows(min_row=2, max_row=3, values_only=True))
            for index, values in enumerate(reversed(rows), start=2):
                for column, value in enumerate(values, start=1):
                    review.cell(index, column).value = value
            review["D2"] = "Corrected fourth row"
            self.assertEqual(read_review_metadata(report), metadata)
            report.save(self.report_path)
        finally:
            report.close()

        result = apply_workflow_revisions(self.report_path)
        self.assertEqual(result.output_path, self.root / "revised_original.xlsx")
        self.assertEqual(result.revised_count, 1)
        revised = load_workbook(result.output_path)
        try:
            self.assertEqual(revised.sheetnames, ["Data"])
            self.assertEqual(revised["Data"]["B4"].value, "Corrected fourth row")
            self.assertEqual(revised["Data"]["B2"].value, "Target 2")
            self.assertEqual(revised["Data"]["B3"].value, "Target 3")
        finally:
            revised.close()

    def test_empty_review_has_only_six_headers_and_no_metadata_rows(self) -> None:
        self.write_report(empty=True)
        report = load_workbook(self.report_path)
        try:
            review = report[WORKFLOW_REVIEW_SHEET_NAME]
            self.assertEqual(list(review.values), [WORKFLOW_REVIEW_HEADERS])
            self.assertEqual(read_review_metadata(report)["data_sheet_name"], "Data")
        finally:
            report.close()
        self.assertEqual(apply_workflow_revisions(self.report_path).revised_count, 0)

    def test_missing_or_damaged_metadata_does_not_write_a_revision(self) -> None:
        for damage in ("missing", "version", "required_field"):
            with self.subTest(damage=damage):
                self.write_report()
                sheet = self.workbook[WORKFLOW_METADATA_SHEET_NAME]
                if damage == "missing":
                    del self.workbook[WORKFLOW_METADATA_SHEET_NAME]
                else:
                    key = "schema_version" if damage == "version" else "target_column"
                    for row in sheet:
                        if row[0].value == key:
                            row[1].value = "unsupported" if damage == "version" else None
                self.workbook.save(self.report_path)
                before = self.report_path.read_bytes()
                output = self.root / "revision.xlsx"
                with self.assertRaisesRegex(ValueError, "重新运行一键质量检查"):
                    apply_workflow_revisions(self.report_path, output)
                self.assertFalse(output.exists())
                self.assertEqual(self.report_path.read_bytes(), before)

    def test_does_not_overwrite_a_business_sheet_with_the_reserved_name(self) -> None:
        sheet = self.workbook.create_sheet(WORKFLOW_METADATA_SHEET_NAME.upper())
        sheet["A1"] = "Business data"
        with self.assertRaisesRegex(ValueError, "重名"):
            self.write_report()
        self.assertEqual(sheet["A1"].value, "Business data")
        self.assertIn(sheet.title, self.workbook.sheetnames)
        self.assertFalse(self.report_path.exists())

    def test_regenerating_report_replaces_its_own_metadata(self) -> None:
        self.write_report()
        self.write_report(empty=True)
        self.assertEqual(self.workbook.sheetnames.count(WORKFLOW_METADATA_SHEET_NAME), 1)
        self.assertEqual(read_review_metadata(self.workbook)["target_column"], "B")

    def test_previously_generated_schema_two_reports_still_apply(self) -> None:
        review = self.workbook.create_sheet(WORKFLOW_REVIEW_SHEET_NAME)
        review.append(WORKFLOW_REVIEW_HEADERS)
        review.append([2, "Source 2", "Target 2", "Revised", "Issue", "Tag 检查"])
        metadata = {
            "schema_version": "2", "input_file": str(self.root / "original.xlsx"),
            "data_sheet_name": "Data", "source_column": "A", "target_column": "B",
            "start_row": 2, "generated_sheet_names": WORKFLOW_REVIEW_SHEET_NAME,
            "remove_term_helper": "0",
        }
        for row, (key, value) in enumerate(metadata.items(), start=1):
            review.cell(row, 7, key)
            review.cell(row, 8, value)
        self.workbook.save(self.report_path)
        result = apply_workflow_revisions(self.report_path)
        self.assertEqual(result.output_path.name, "revised_original.xlsx")
        self.assertEqual(result.revised_count, 1)

    def test_default_output_naming_closes_workbook_when_metadata_is_invalid(self) -> None:
        self.write_report()
        with (
            patch("tools.workflow.revision_applier.load_workbook") as load,
            patch("tools.workflow.revision_applier.read_review_metadata", side_effect=ValueError),
        ):
            load.return_value.sheetnames = [WORKFLOW_REVIEW_SHEET_NAME]
            self.assertEqual(
                build_default_revised_output_path(self.report_path).name, "revised_report.xlsx",
            )
            load.return_value.close.assert_called_once_with()


if __name__ == "__main__":
    unittest.main()
