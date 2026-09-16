"""Semantic and allocation regressions for the QA performance paths."""
from __future__ import annotations

import importlib
from pathlib import Path
import random
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook

from tools.excel_output import existing_cell_value, value_row_numbers
from tools.substring_consistency_checker import check_substring_consistency as substring
from tools.term_pair_checker import extract_terms_from_excel as terms


CHECK_MODULES = (
    "source_consistency_checker.check_source_consistency",
    "target_consistency_checker.check_target_consistency",
    "substring_consistency_checker.check_substring_consistency",
    "term_pair_checker.extract_terms_from_excel",
    "tag_placeholder_checker.check_tags_and_placeholders",
    "line_break_checker.check_line_breaks",
    "content_fidelity_checker.check_content_fidelity",
    "chinese_target_checker.check_chinese_target",
    "target_text_checker.check_target_text",
)


class SparseQaTests(unittest.TestCase):
    def test_sparse_row_selection_keeps_values_and_order_without_creating_cells(self):
        book = Workbook()
        self.addCleanup(book.close)
        ws = book.active
        for row, column, value in ((90, 2, ""), (7, 1, False), (3, 2, 0),
                                   (5, 1, "=1+1"), (2, 1, "   "), (1, 1, "header")):
            ws.cell(row, column, value)
        ws["Z200"] = "ignored"
        ws["A1000"].number_format = "@"
        original = set(ws._cells)
        self.assertEqual(value_row_numbers(ws, ("a", "B"), start_row=2), [2, 3, 5, 7, 90])
        self.assertIsNone(existing_cell_value(ws, 700, 2))
        self.assertEqual(set(ws._cells), original)

    def test_all_processors_preserve_sparse_data_and_span_statistics(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            path.touch()
            for name in CHECK_MODULES:
                with self.subTest(check=name):
                    book = Workbook()
                    try:
                        ws = book.active
                        ws.title = "Data"
                        ws.append(["source", "target"])
                        # Insert out of row order, with either side absent.
                        ws["B100001"] = "译文.. ("
                        ws["A2"] = "[Alpha] 12\n{name}"
                        ws["A80000"].number_format = "@"
                        ws["Z200000"] = "unrelated"
                        before = {key: (cell.value, cell.number_format) for key, cell in ws._cells.items()}
                        options = dict(workbook=book, output_path=path, source_column="A",
                                       target_column="B", sheet="Data", format_output=False)
                        if name.startswith("term_pair"):
                            options["input_file"] = path
                        summary = importlib.import_module("tools." + name).process_workbook(**options)
                        self.assertEqual(before, {key: (cell.value, cell.number_format)
                                                  for key, cell in ws._cells.items()})
                        count = getattr(summary, "total_rows_checked", getattr(summary, "processed_count", None))
                        if count is not None:
                            self.assertEqual(count, 100000)
                        issue_rows = {cell.value for sheet in book if sheet.title != "Data"
                                      and sheet["A1"].value == "行号"
                                      for cell in sheet["A"][1:]}
                        if not name.startswith(("source_consistency", "target_consistency", "substring_consistency")):
                            self.assertTrue(issue_rows)
                            self.assertTrue(issue_rows <= {2, 100001})
                    finally:
                        book.close()


class ConsistencyReportAllocationTests(unittest.TestCase):
    def test_long_group_is_shared_and_omitting_detail_keeps_review_identical(self):
        from tools.workflow.review_sheet import collect_review_rows
        for name, label in ((CHECK_MODULES[0], "同 Source 不同 Target"),
                            (CHECK_MODULES[1], "同 Target 不同 Source")):
            module = importlib.import_module("tools." + name)
            with self.subTest(check=name):
                book = Workbook()
                try:
                    ws = book.active
                    ws.title = "Data"
                    ws.append(["source", "target"])
                    for index in range(8000):
                        values = ("shared", f"variant {index % 2}")
                        ws.append(values if name == CHECK_MODULES[0] else values[::-1])
                    options = dict(workbook=book, output_path=Path("unused.xlsx"),
                                   source_column="A", target_column="B", sheet="Data", format_output=False)
                    original = module.process_workbook(**options)
                    sheet = book[module.PROBLEM_SHEET_NAME]
                    expected = "、".join(map(str, range(2, 8002)))[:32767]
                    self.assertEqual(sheet["F2"].value, expected)
                    self.assertIs(sheet["F2"].value, sheet["F8001"].value)
                    review = collect_review_rows(book, [(label, module.PROBLEM_SHEET_NAME)])
                    reduced = module.process_workbook(**options, include_grouped_rows=False)
                    self.assertEqual(original, reduced)
                    self.assertEqual(review, collect_review_rows(book, [(label, module.PROBLEM_SHEET_NAME)]))
                finally:
                    book.close()


class TermTransactionTests(unittest.TestCase):
    def test_row_conflict_rolls_back_new_terms_and_empty_target_upgrades(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            book = Workbook()
            try:
                ws = book.active
                ws.title = "Data"
                for row in [
                    ("source", "target"), ("[Base]", "[Correct]"), ("[Empty]", ""),
                    ("[Empty] [New] [Base]", "[Filled] [First] [Wrong]"),
                    ("[New]", "[Second]"), ("[Empty]", "[Final]"),
                    ("[Twin] [Twin]", "[One] [Two]"), ("[Twin]", "[Three]"),
                ]:
                    ws.append(row)
                path.touch()
                learned = {}
                original_build = terms.build_term_mapping_entries

                def capture(values):
                    values = list(values)
                    learned.update((item.source_plain_text, item.target_plain_text) for item in values)
                    return original_build(values)

                with patch.object(terms, "build_term_mapping_entries", side_effect=capture):
                    terms.process_workbook(workbook=book, input_file=path, output_path=path,
                                           source_column="A", target_column="B", sheet="Data")
                self.assertEqual(learned, {"Base": "Correct", "Empty": "Final", "New": "Second", "Twin": "Three"})
                issues = {row[0]: row[3] for row in list(book["问题列"].values)[1:]}
                self.assertIn("Wrong", issues[4])
                self.assertIn("Two", issues[7])
                self.assertNotIn(5, issues)
                self.assertNotIn(6, issues)
                self.assertNotIn(8, issues)
            finally:
                book.close()


class DenseSubstringTests(unittest.TestCase):
    def test_fallback_has_exact_original_match_order_and_boundaries(self):
        rng = random.Random(512)
        samples = [("中" * 80, ["中" * n for n in range(1, 70)]),
                   ("𠀀a 𠀀a𠀁 a𠀀 SSＡ", ["𠀀", "𠀀a", "a𠀀", "𠀁", "SS", "SSＡ"]),
                   ("aaaaaa aa aaaa _aa aa_ 中文aa！", ["a", "aa", "aaa", "aaaa", "中文", "文a", "a！"]),
                   ("abc ab abc αabß ab１ ab! _ ab", ["ab", "abc", "b", "b１", "αab", "abß", "_", " ab"])]
        alphabet = "ab12_ 中文ßα１![]\n"
        for _ in range(100):
            text = "".join(rng.choices(alphabet, k=100))
            patterns = {text[start:start + rng.randrange(1, 9)] for start in rng.sample(range(90), 30)}
            patterns.update(("absent", "中中中"))
            samples.append((text, patterns))
        for text, patterns in samples:
            index = substring._build_index(patterns)
            expected = list(dict.fromkeys(pattern for pattern, _, _ in substring._matches(text, index)))
            with self.subTest(text=text):
                with patch.object(substring, "MIN_POSITION_BUDGET", 0), patch.object(substring, "POSITIONS_PER_CHARACTER", 0):
                    self.assertEqual(substring._unique_patterns(text, index, {}), expected)

    def test_dense_full_reports_preserve_counts_descriptions_and_order(self):
        rows = [("中" * n, "译" * n) for n in range(3, 48)]
        rows += [("中" * 65, "译" * 12), ("中" * 65, "译" * 12), ("中" * 70, "Wrong")]

        def run(budget):
            book = Workbook()
            try:
                book.active.append(["source", "target"])
                for row in rows:
                    book.active.append(row)
                with patch.object(substring, "MIN_POSITION_BUDGET", budget):
                    summary = substring.process_workbook(workbook=book, output_path=Path("unused.xlsx"),
                                                         source_column="A", target_column="B")
                return summary, list(book[substring.PROBLEM_SHEET_NAME].values)
            finally:
                book.close()

        self.assertEqual(run(1_000_000), run(1))

    def test_dense_position_enumeration_is_bounded_without_losing_patterns(self):
        patterns = ["中" * n for n in range(1, 200)]
        index = substring._build_index(patterns)

        class CountingIndex:
            count = 0

            def __len__(self):
                return len(index)

            def keys(self):
                return index.keys()

            def iter(self, text):
                for item in index.iter(text):
                    self.count += 1
                    yield item

        counted = CountingIndex()
        found = substring._unique_patterns("中" * 200, counted, {})
        self.assertEqual(found, patterns)
        self.assertLessEqual(counted.count, 400)


if __name__ == "__main__":
    unittest.main()
