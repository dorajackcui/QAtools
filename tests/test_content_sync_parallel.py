"""Serial/parallel equivalence, deliberate out-of-order completion and isolation."""
from contextlib import contextmanager
from dataclasses import asdict
from pathlib import Path
import tempfile
import threading
import unittest
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.styles import PatternFill

from tools.content_sync import master_to_target as forward, target_to_master as reverse
from tools.content_sync.parallel import ordered_results


class ParallelSyncTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.master = self.root / "master.xlsx"
        self.targets = self.root / "targets"
        self.targets.mkdir()

    def write(self, path, rows):
        book = Workbook()
        try:
            for row in rows:
                book.active.append(row)
            # Literal formula/error-like source text remains text.
            for row in book.active:
                for cell in row:
                    if isinstance(cell.value, str):
                        cell.data_type = "s"
            book.active["C2"].fill = PatternFill("solid", fgColor="FF0011")
            book.create_sheet("Other")["A1"] = "=1+2"
            book.save(path)
        finally:
            book.close()

    def normalized(self, value, output):
        if isinstance(value, dict):
            return {k: self.normalized(v, output) for k, v in value.items()}
        if isinstance(value, (tuple, list)):
            return [self.normalized(v, output) for v in value]
        if isinstance(value, (str, Path)):
            return str(value).replace(str(output), "OUTPUT")
        return value

    def equivalent_outputs(self, first, second):
        names = sorted(p.relative_to(first) for p in first.rglob("*.xlsx"))
        self.assertEqual(names, sorted(p.relative_to(second) for p in second.rglob("*.xlsx")))
        for name in names:
            with ZipFile(first / name) as a, ZipFile(second / name) as b:
                self.assertEqual(a.namelist(), b.namelist())
                for entry in a.namelist():
                    if entry != "docProps/core.xml":  # Save timestamps are not content.
                        self.assertEqual(a.read(entry), b.read(entry), (name, entry))

    def compare(self, operation, *, case="", **kwargs):
        before = {p: p.read_bytes() for p in self.root.rglob("*.xlsx")}
        results, logs = [], []
        owner = threading.get_ident()
        for workers in (1, 2):
            output = self.root / f"{case}out-{workers}"
            lines, progress, callback_threads = [], [], []

            def log(line):
                callback_threads.append(threading.get_ident())
                lines.append(line)

            def update(done, total):
                callback_threads.append(threading.get_ident())
                progress.append((done, total))

            result = operation(self.master, self.targets, output_dir=output, workers=workers,
                               log_callback=log, progress_callback=update, **kwargs)
            self.assertEqual(set(callback_threads), {owner})
            count = len(list(self.targets.glob("*.xlsx")))
            self.assertEqual(progress, [(i, count) for i in range(1, count + 1)])
            results.append(self.normalized(asdict(result), output))
            logs.append(self.normalized(lines, output))
        self.assertEqual(results[0], results[1])
        self.assertEqual(logs[0], logs[1])
        self.equivalent_outputs(self.root / f"{case}out-1", self.root / f"{case}out-2")
        for path, data in before.items():
            self.assertEqual(path.read_bytes(), data)
        return results[1]

    def test_forward_values_multicolumn_counts_logs_and_save_failure(self):
        self.write(self.master, [["id", "k", "s", "t", "t2"],
            [1, "k", "S", "  nan\n ", None], [2, "n", "N", 0, False],
            [3, "k", "S", "=literal", "#N/A"]])
        rows = [["k", "s", "t", "t2"], ["k", "S", "old", "keep"],
                ["n", "N", None, None], ["unknown", "S", "keep"], [None, "S", "keep"]]
        for name in ("a", "b", "failed"):
            self.write(self.targets / f"{name}.xlsx", rows)
        (self.targets / "broken.xlsx").write_bytes(b"bad")
        original = forward._save_output

        def save(book, source, destination, **kwargs):
            if source.stem == "failed":
                raise PermissionError("locked")
            return original(book, source, destination, **kwargs)

        with patch.object(forward, "_save_output", save):
            result = self.compare(forward.sync_master_to_targets, column_count=2)
        self.assertEqual([r["updated_cells"] for r in result["files"]], [4, 4, 0, 0])
        self.assertEqual(sorted(p.name for p in (self.root / "out-2").iterdir()), ["a.xlsx", "b.xlsx"])

    def test_reverse_out_of_order_completion_keeps_conflicts_and_blank_policy(self):
        original = reverse._read_master
        for fill in (False, True):
            for allow in (False, True):
                with self.subTest(fill=fill, allow=allow):
                    self.write(self.master, [["id", "k", "s", "t"],
                        [1, "k", "S", "old"], [2, "n", "N", None], [3, "k", "S", None]])
                    self.write(self.targets / "a.xlsx", [["k", "s", "t"],
                        ["k", "S", "first"], ["n", "N", "None"]])
                    self.write(self.targets / "z.xlsx", [["k", "s", "t"],
                        ["k", "S", "  last\n "], ["k", "S", None], ["new", "S", "nan"]])
                    later_done = threading.Event()
                    owner = threading.get_ident()

                    def read(path, *args, **kwargs):
                        if threading.get_ident() != owner and path.name == "a.xlsx":
                            self.assertTrue(later_done.wait(5), "second file was not processed concurrently")
                        data = original(path, *args, **kwargs)
                        if threading.get_ident() != owner and path.name == "z.xlsx":
                            later_done.set()
                        return data

                    with patch.object(reverse, "_read_master", read):
                        self.compare(reverse.sync_targets_to_master, case=f"{fill}-{allow}-",
                                     fill_blank_only=fill, allow_blank_write=allow)
                    self.assertTrue(later_done.is_set())

    def test_reverse_rejects_entire_source_with_missing_formula_cache(self):
        self.write(self.master, [["id", "k", "s", "t"], [1, "k", "S", "old"]])
        self.write(self.targets / "a.xlsx", [["k", "s", "t"], ["k", "S", "valid"]])
        book = Workbook()
        try:
            for row in [["k", "s", "t"], ["k", "S", "partial"], ["n", "N", "=1+1"]]:
                book.active.append(row)
            book.save(self.targets / "z.xlsx")
        finally:
            book.close()
        result = self.compare(reverse.sync_targets_to_master, allow_blank_write=True)
        self.assertEqual(result["details"]["updated_cells"], 1)
        self.assertEqual(result["results"][1]["status"], "failed")
        self.assertEqual(result["details"]["conflicts"], [])

    def test_com_and_linked_inputs_use_coordinator_thread(self):
        self.write(self.master, [["id", "k", "s", "t"], [1, "k", "S", "new"]])
        for name in ("a", "b"):
            self.write(self.targets / f"{name}.xlsx", [["k", "s", "t"], ["k", "S", "old"]])
        owner = threading.get_ident()
        calls = []
        original = forward._sync_file

        def sync(*args, **kwargs):
            calls.append(("sync", threading.get_ident()))
            return original(*args, **kwargs)

        @contextmanager
        def session():
            calls.append(("open", threading.get_ident()))
            yield object()
            calls.append(("close", threading.get_ident()))

        def resave(*args):
            calls.append(("resave", threading.get_ident()))

        with patch.object(forward, "_sync_file", sync), \
             patch("tools.excel_com.excel_session", session), patch("tools.excel_com.edit_copy", resave):
            result = forward.sync_master_to_targets(self.master, self.targets, workers=2, compatibility_resave=True)
        self.assertEqual(result.updated_cells, 2)
        self.assertEqual([kind for kind, _ in calls], ["open", "sync", "resave", "sync", "resave", "close"])
        self.assertEqual({thread for _, thread in calls}, {owner})
        calls.clear()
        with patch.object(Path, "is_symlink", return_value=True), patch.object(forward, "_sync_file", sync):
            forward.sync_master_to_targets(self.master, self.targets, workers=2)
        self.assertEqual(calls, [("sync", owner), ("sync", owner)])

    def test_invalid_workers_fail_before_any_io(self):
        for operation in (forward.sync_master_to_targets, reverse.sync_targets_to_master):
            for workers in (0, 5, True, 1.5, "2"):
                with self.subTest(operation=operation.__name__, workers=workers), self.assertRaisesRegex(ValueError, "并行任务数"):
                    operation(self.master, self.targets, workers=workers)
        self.assertEqual(list(self.targets.iterdir()), [])

    def test_failed_source_releases_file_even_while_traceback_is_retained(self):
        book = Workbook()
        try:
            book.active.append(["k", "s", "t"])
            book.active.append(["k", "S", "=1+1"])
            book.save(self.master)
        finally:
            book.close()
        errors = []
        try:
            forward._read_master(self.master, forward.ColumnMapping("A", "B", "C"), 1, None, 1)
        except ValueError as exc:
            errors.append(exc)
        self.assertEqual(len(errors), 1)
        moved = self.master.with_name("moved.xlsx")
        self.master.replace(moved)
        moved.replace(self.master)


class OrderedResultsTests(unittest.TestCase):
    def test_bounded_submission_ordered_errors_and_executor_cleanup(self):
        second_done, release = threading.Event(), threading.Event()
        started, threads = [], set()

        def task(index):
            started.append(index)
            threads.add(threading.current_thread())
            if index == 0:
                if not second_done.wait(5):
                    raise RuntimeError("second task never started")
                raise ValueError("bad first file")
            if index == 1:
                second_done.set()
            if index >= 2:
                release.wait(5)
            return index

        iterator = ordered_results(task, range(100), 2)
        try:
            value, error = next(iterator)
            self.assertIsNone(value)
            self.assertEqual(error, "bad first file")
            self.assertTrue(set(started).issubset({0, 1, 2}))
            self.assertEqual(next(iterator), (1, None))
        finally:
            release.set()
            iterator.close()
        self.assertTrue(all(not thread.is_alive() for thread in threads))
