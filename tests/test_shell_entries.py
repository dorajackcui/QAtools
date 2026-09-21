from __future__ import annotations

import os
from pathlib import Path
import sys
import tempfile
import time
import unittest
import uuid
from unittest.mock import Mock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook, load_workbook
from PySide6.QtCore import QCoreApplication, QEvent, QEventLoop, QProcess, QTimer
from PySide6.QtWidgets import QApplication

from toolshub_gui import ToolshubApp, _initial_request, build_argument_parser, main
from tools.qt_navigation import ToolNavigationServer, send_path_request
from tools.workflow.file_receiver import ToolFileRequest


class ShellEntryTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.folder = self.root / "中文 folder & data"
        self.folder.mkdir()
        self.file = self.folder / "法语 input.xlsx"
        self.write_workbook(self.file, [["Source", "Target"], ["Hello!", "Bonjour !"]])
        for target, filename in (
            ("tools.workflow.gui_options.default_header_aliases_path", "aliases.json"),
            ("tools.tb_projects.default_tb_projects_path", "projects.json"),
        ):
            mocked = patch(target, return_value=self.root / filename)
            mocked.start()
            self.addCleanup(mocked.stop)

    @staticmethod
    def write_workbook(path, rows):
        workbook = Workbook()
        try:
            for row in rows:
                workbook.active.append(row)
            workbook.save(path)
        finally:
            workbook.close()

    def window(self):
        window = ToolshubApp(show_window=False)
        window._bring_window_to_front = Mock()
        self.addCleanup(self.dispose, window)
        return window

    @staticmethod
    def dispose(window):
        window.close()
        window.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)

    def test_all_five_startup_actions_and_existing_instance_forwarding(self):
        for flag, action, path in (
            ("--qa-workflow", "qa_workflow", self.file),
            ("--nbsp-restore", "french_nbsp_restore", self.file),
            ("--compatibility-dir", "compatibility", self.folder),
            ("--merge-dir", "excel_merger", self.folder),
            ("--untranslated-dir", "untranslated_stats", self.folder),
        ):
            with self.subTest(flag=flag):
                request = _initial_request(build_argument_parser().parse_args([flag, str(path)]))
                self.assertEqual(request, ToolFileRequest(action, str(path)))
                with patch("toolshub_gui.ToolshubApp") as factory, patch("toolshub_gui.WorkflowFileReceiver"):
                    self.assertEqual(main([flag, str(path), "--smoke-test"]), 0)
                    factory.return_value.handle_file_request.assert_called_once_with(request)
                with (
                    patch("toolshub_gui._acquire_gui_instance_lock", return_value=None),
                    patch("toolshub_gui.send_path_request", return_value=True) as send,
                    patch("toolshub_gui.ToolshubApp") as factory,
                ):
                    self.assertEqual(main([flag, str(path)]), 0)
                    send.assert_called_once_with(action, str(path))
                    factory.assert_not_called()

    def test_invalid_paths_and_unreachable_instance_show_errors(self):
        with patch("toolshub_gui.show_error") as error:
            self.assertEqual(main(["--qa-workflow", str(self.root / "missing.xlsx")]), 2)
            self.assertEqual(main(["--merge-dir", str(self.file)]), 2)
            with patch("toolshub_gui._acquire_gui_instance_lock", return_value=None), patch(
                "toolshub_gui.send_path_request", return_value=False
            ):
                self.assertEqual(main(["--qa-workflow", str(self.file)]), 1)
            self.assertEqual(error.call_count, 3)
        with self.assertRaises(SystemExit):
            build_argument_parser().parse_args(["--merge-dir", str(self.folder), "--qa-workflow", str(self.file)])

    def test_directory_actions_load_pages_without_running_and_preserve_options(self):
        window = self.window()
        for action in ("compatibility", "excel_merger", "untranslated_stats"):
            with self.subTest(action=action):
                page = window.tool_frames[action]
                page.run_in_background = Mock()
                if action != "excel_merger":
                    page.output_picker.set_path("chosen-output")
                window.handle_file_request(ToolFileRequest(action, str(self.folder)))
                picker = page.input_dir if action == "excel_merger" else page.input_picker
                self.assertEqual(picker.path(), str(self.folder))
                self.assertEqual(window.current_tool_key, action)
                page.run_in_background.assert_not_called()
                if action != "excel_merger":
                    self.assertEqual(page.output_picker.path(), "chosen-output")
                else:
                    self.assertIn(self.folder.name, page.preview.text())

    def test_qa_loads_and_detects_columns_without_running(self):
        window = self.window()
        page = window.tool_frames["workflow"]
        page.run_in_background = Mock()
        window.handle_file_request(ToolFileRequest("qa_workflow", str(self.file)))
        self.assertEqual(window.current_tool_key, "workflow")
        self.assertEqual(page.input_picker.path(), str(self.file))
        self.assertEqual((page.source_column.text(), page.target_column.text()), ("A", "B"))
        page.run_in_background.assert_not_called()

    def test_busy_pages_reject_paths_before_changing_input(self):
        window = self.window()
        for action, key, path in (
            ("qa_workflow", "workflow", self.file),
            ("french_nbsp_restore", "french_nbsp", self.file),
            ("compatibility", "compatibility", self.folder),
            ("excel_merger", "excel_merger", self.folder),
            ("untranslated_stats", "untranslated_stats", self.folder),
        ):
            page = window.tool_frames[key]
            picker = page.input_dir if key == "excel_merger" else page.input_picker
            before = picker.path()
            with patch.object(page, "has_running_tasks", return_value=True), patch("toolshub_gui.show_error") as error:
                window._open_forwarded_path(action, str(path))
                error.assert_called_once()
                self.assertEqual(picker.path(), before)

    def test_nbsp_direct_action_writes_copy_and_keeps_source(self):
        window = self.window()
        page = window.tool_frames["french_nbsp"]
        page.start_row.setValue(10)
        page.result_column.setText("D")
        original = self.file.read_bytes()
        with patch("tools.french_nbsp_restorer.qt_page.show_info") as done:
            window.handle_file_request(ToolFileRequest("french_nbsp_restore", str(self.file)))
            deadline = time.monotonic() + 15
            while page.has_running_tasks() and time.monotonic() < deadline:
                self.app.processEvents()
                time.sleep(0.01)
            self.assertFalse(page.has_running_tasks())
            self.app.processEvents()
            done.assert_called_once()
        self.assertEqual(self.file.read_bytes(), original)
        result = self.folder / ("french_nbsp_restore_" + self.file.name)
        workbook = load_workbook(result)
        try:
            self.assertEqual(workbook.active["B2"].value, "Bonjour\u00a0!")
            self.assertIsNone(workbook.active["D2"].value)
        finally:
            workbook.close()

    def test_nbsp_missing_ambiguous_or_unreadable_headers_never_auto_run(self):
        window = self.window()
        page = window.tool_frames["french_nbsp"]
        page.run_restore = Mock()
        for headers in (["Source", "Unknown"], ["Target", "Target"]):
            self.write_workbook(self.file, [headers, ["Bonjour !", "Salut !"]])
            with patch("tools.french_nbsp_restorer.qt_page._show_error") as error:
                window.handle_file_request(ToolFileRequest("french_nbsp_restore", str(self.file)))
                error.assert_called_once()
                self.assertEqual(page.target_column.text(), "")
        self.file.write_bytes(b"invalid workbook")
        with patch("tools.french_nbsp_restorer.qt_page._show_error") as error:
            window.handle_file_request(ToolFileRequest("french_nbsp_restore", str(self.file)))
            error.assert_called_once()
        self.write_workbook(self.file, [["Source", "Target"]])
        with patch("tools.french_nbsp_restorer.qt_page.detect_source_target_columns", side_effect=OSError("locked")), patch(
            "tools.french_nbsp_restorer.qt_page._show_error"
        ):
            window.handle_file_request(ToolFileRequest("french_nbsp_restore", str(self.file)))
        page.run_restore.assert_not_called()

    def test_real_local_socket_acknowledges_unicode_path_after_delayed_start(self):
        name = f"qatools-shell-test-{uuid.uuid4().hex}"
        server = ToolNavigationServer(name=name)
        client = QProcess()
        loop = QEventLoop()
        received = []
        server.path_requested.connect(lambda *args: received.append(args))
        client.finished.connect(loop.quit)
        timer = QTimer()
        timer.setSingleShot(True)
        timer.timeout.connect(loop.quit)
        startup = QTimer()
        startup.setSingleShot(True)
        startup.timeout.connect(server.start)
        code = (
            "import sys\nfrom unittest.mock import patch\n"
            "from PySide6.QtCore import QCoreApplication\n"
            "from tools.qt_navigation import send_path_request\n"
            "app = QCoreApplication([])\n"
            "with patch('tools.qt_navigation.navigation_server_name', return_value=sys.argv[1]):\n"
            "    raise SystemExit(0 if send_path_request('qa_workflow', sys.argv[2]) else 1)\n"
        )
        try:
            client.setWorkingDirectory(str(Path(__file__).resolve().parents[1]))
            client.start(sys.executable, ["-c", code, name, str(self.file)])
            self.assertTrue(client.waitForStarted(5000))
            startup.start(1000)
            timer.start(10000)
            loop.exec()
            self.app.processEvents()
            self.assertEqual(client.state(), QProcess.ProcessState.NotRunning)
            self.assertEqual(client.exitCode(), 0, bytes(client.readAllStandardError()).decode(errors="replace"))
            self.assertEqual(received, [("qa_workflow", str(self.file))])
        finally:
            startup.stop()
            timer.stop()
            if client.state() != QProcess.ProcessState.NotRunning:
                client.kill()
                client.waitForFinished(5000)
            server.close()
            server.deleteLater()
            QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)

    def test_path_protocol_rejects_unknown_actions_and_invalid_paths(self):
        server = ToolNavigationServer()
        self.addCleanup(server.deleteLater)
        received = []
        server.path_requested.connect(lambda *args: received.append(args))
        for payload in (b'{"action":"unknown","path":"x"}\n', b'{"action":"qa_workflow","path":7}\n',
                        b'{"action":"qa_workflow","path":""}\n', b'{"action":[],"path":"x"}\n'):
            socket = Mock()
            socket.bytesAvailable.return_value = len(payload)
            socket.canReadLine.return_value = True
            socket.readLine.return_value = payload
            server._read(socket)
            socket.write.assert_not_called()
        self.app.processEvents()
        self.assertEqual(received, [])
        self.assertFalse(send_path_request("unknown", str(self.file)))
        self.assertFalse(send_path_request("qa_workflow", "x" * 5000))


if __name__ == "__main__":
    unittest.main()
