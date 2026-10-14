from __future__ import annotations

import base64
from pathlib import Path
import random
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils.cell import range_boundaries
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from phraseloom import workbook_schema as schema
from phraseloom.errors import WorkbookFormatError
from phraseloom.strings_package import load_strings_metadata
from phraseloom.strings_workflow import export_strings_workbook, restore_strings_workbook
from phraseloom.workbook_backup import CHUNK_CHARACTERS, read_original_file
from phraseloom.workbook_io import workbook_metadata


def write_sparse_source(path: Path, *, large_backup: bool = False) -> None:
    book = Workbook()
    try:
        sheet = book.active
        # Original names can now coincide with package sheet names.
        sheet.title = schema.STRINGS_SHEET
        sheet.append(["source", "target", "context"])
        sheet["A2"] = "Early message"
        sheet["C2"] = "Early context"
        sheet["A4"] = "Already translated"
        sheet["B4"] = "Existing translation"
        sheet["A100001"] = "Distant message"
        sheet["C100001"] = "Distant context"
        sheet["B100001"].fill = PatternFill("solid", fgColor="FFDDEE")
        sheet["B100001"].font = Font(bold=True, color="123456")
        sheet["XFD1048576"].number_format = "@"
        sheet.row_dimensions[100001].height = 35
        sheet.column_dimensions["C"].width = 45
        sheet.freeze_panes = "C2"
        sheet.merge_cells("D7:E8")
        sheet["D7"] = "Merged note"
        sheet["C2"].comment = Comment("Retain comment", "QA")
        sheet["C2"].hyperlink = "https://example.com/context"
        for index in range(10, 5010):
            sheet.cell(index, 2).number_format = "@"
        notes = book.create_sheet(schema.METADATA_SHEET)
        notes["A1"] = "Notes"
        notes["B3"] = False
        notes["C3"] = 0
        notes["D3"] = "=1+1"
        notes["XFD1048576"] = "Real distant content"
        if large_backup:
            rng = random.Random(42)
            for row in range(10, 40):
                notes.cell(row, 1, "".join(rng.choices("abcdefghijklmnopqrstuvwxyz0123456789", k=2000)))
        hidden = book.create_sheet(schema.ORIGINAL_FILE_SHEET)
        hidden["A1"] = "Hidden original content"
        hidden.sheet_state = "veryHidden"
        sheet.sheet_state = "hidden"
        book.active = 1
        book.save(path)
    finally:
        book.close()


class WorkbookBackupTests(unittest.TestCase):
    def test_backup_accepts_integral_numeric_xml_encodings(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            package = Path(directory) / "strings.xlsx"
            write_sparse_source(source)
            export_strings_workbook(source, package)
            with ZipFile(package) as archive:
                entries = [(item, archive.read(item)) for item in archive.infolist()]
            # Spreadsheet writers may encode an integer as 1.0 or 1e0.
            for suffix in (".0", "e0"):
                with self.subTest(suffix=suffix):
                    with ZipFile(package, "w", compression=ZIP_DEFLATED) as archive:
                        for item, data in entries:
                            if item.filename in ("xl/worksheets/sheet4.xml", "xl/worksheets/sheet5.xml"):
                                root = ET.fromstring(data)
                                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                                for cell in root.findall("s:sheetData/s:row/s:c", ns):
                                    value = cell.find("s:v", ns)
                                    if cell.get("t") == "n" and value is not None:
                                        value.text += suffix
                                data = ET.tostring(root)
                            archive.writestr(item, data)
                    book = load_workbook(package, read_only=True)
                    try:
                        with read_original_file(book, workbook_metadata(book)) as original:
                            self.assertEqual(original.read(), source.read_bytes())
                    finally:
                        book.close()

    def test_all_export_sheets_are_compact_and_roundtrip_preserves_original(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            package = Path(directory) / "strings.xlsx"
            restored = Path(directory) / "restored.xlsx"
            write_sparse_source(source, large_backup=True)
            original_bytes = source.read_bytes()
            with patch.object(ReadOnlyWorksheet, "iter_rows", side_effect=AssertionError("dense read")):
                export_strings_workbook(source, package, context_col="context")
            self.assertEqual(source.read_bytes(), original_bytes)
            book = load_workbook(package)
            try:
                self.assertEqual(book.sheetnames, [schema.STRINGS_SHEET, schema.COMPLETED_STRINGS_SHEET,
                    schema.STRINGS_MAP_SHEET, schema.METADATA_SHEET, schema.ORIGINAL_FILE_SHEET])
                self.assertEqual(book.active.title, schema.STRINGS_SHEET)
                metadata = workbook_metadata(book)
                self.assertEqual(metadata[schema.SCHEMA_VERSION_KEY], "2.0")
                self.assertGreater(metadata[schema.ORIGINAL_FILE_CHUNKS_KEY], 1)
                backup = book[schema.ORIGINAL_FILE_SHEET]
                self.assertEqual(backup.sheet_state, "hidden")
                self.assertEqual(backup.max_column, 2)
                self.assertTrue(all(len(row[1].value) <= CHUNK_CHARACTERS
                                    for row in backup.iter_rows(min_row=2)))
                with read_original_file(book, metadata) as original:
                    self.assertEqual(original.read(), original_bytes)
                for sheet in book:
                    self.assertLess(sheet.max_row * sheet.max_column, 200)
                self.assertEqual(book[schema.STRINGS_MAP_SHEET]["A3"].value, 100001)
                self.assertEqual(book[schema.COMPLETED_STRINGS_SHEET]["A2"].value, 4)
            finally:
                book.close()
            # Check every actual XML dimension, including hidden backup/metadata.
            with ZipFile(package) as archive:
                for name in archive.namelist():
                    if name.startswith("xl/worksheets/") and name.endswith(".xml"):
                        root = ET.fromstring(archive.read(name))
                        ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                        bounds = range_boundaries(root.find("s:dimension", ns).get("ref"))
                        self.assertLess((bounds[2] - bounds[0] + 1) * (bounds[3] - bounds[1] + 1), 200)

            # An editor can resave, reorder chunks or introduce ghost formatting.
            book = load_workbook(package)
            try:
                book[schema.STRINGS_SHEET]["D2"] = "Early translation"
                book[schema.STRINGS_SHEET]["D3"] = "Distant translation"
                backup = book[schema.ORIGINAL_FILE_SHEET]
                first = [backup.cell(2, col).value for col in (1, 2)]
                last = [backup.cell(backup.max_row, col).value for col in (1, 2)]
                for col in (1, 2):
                    backup.cell(2, col).value = last[col - 1]
                    backup.cell(backup.max_row, col).value = first[col - 1]
                backup.move_range(f"A{backup.max_row}:B{backup.max_row}", rows=100000)
                for sheet in book:
                    sheet["XFD1048576"].number_format = "@"
                book.save(package)
            finally:
                book.close()
            source.unlink()  # Restore must be entirely self-contained.
            with patch.object(ReadOnlyWorksheet, "iter_rows", side_effect=AssertionError("dense read")):
                stats = restore_strings_workbook(package, restored)
            self.assertEqual(stats["restored_row_count"], 2)
            self.assertEqual(stats["issue_count"], 0)
            book = load_workbook(restored)
            try:
                self.assertEqual(book.sheetnames, [schema.STRINGS_SHEET, schema.METADATA_SHEET, schema.ORIGINAL_FILE_SHEET])
                self.assertEqual([sheet.sheet_state for sheet in book], ["hidden", "visible", "veryHidden"])
                self.assertEqual(book.active.title, schema.METADATA_SHEET)
                sheet = book[schema.STRINGS_SHEET]
                self.assertEqual(sheet["B2"].value, "Early translation")
                self.assertEqual(sheet["B100001"].value, "Distant translation")
                self.assertEqual(sheet["B4"].value, "Existing translation")
                self.assertNotIn((3, 2), sheet._cells)
                self.assertEqual(sheet["XFD1048576"].number_format, "@")
                self.assertEqual(sheet["B100001"].fill.fgColor.rgb, "00FFDDEE")
                self.assertTrue(sheet["B100001"].font.bold)
                self.assertEqual(sheet.row_dimensions[100001].height, 35)
                self.assertEqual(sheet.column_dimensions["C"].width, 45)
                self.assertEqual(sheet.freeze_panes, "C2")
                self.assertIn("D7:E8", sheet.merged_cells)
                self.assertEqual(sheet["D7"].value, "Merged note")
                self.assertEqual(sheet["C2"].comment.text, "Retain comment")
                self.assertEqual(sheet["C2"].hyperlink.target, "https://example.com/context")
                notes = book[schema.METADATA_SHEET]
                self.assertIs(notes["B3"].value, False)
                self.assertEqual(notes["C3"].value, 0)
                self.assertEqual(notes["D3"].value, "=1+1")
                self.assertEqual(notes["XFD1048576"].value, "Real distant content")
                self.assertEqual(book[schema.ORIGINAL_FILE_SHEET]["A1"].value, "Hidden original content")
            finally:
                book.close()
            package.rename(package.with_name("package_closed.xlsx"))
            restored.rename(restored.with_name("restored_closed.xlsx"))

    def test_corrupt_backup_is_rejected_without_overwriting_destination(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            package = Path(directory) / "strings.xlsx"
            edited = Path(directory) / "edited.xlsx"
            restored = Path(directory) / "restored.xlsx"
            write_sparse_source(source, large_backup=True)
            export_strings_workbook(source, package)
            for damage in ("missing_sheet", "missing_chunk", "duplicate_chunk", "bad_base64",
                           "truncated_chunk", "changed_bytes", "invalid_index", "fractional_index",
                           "wrong_size", "wrong_hash"):
                with self.subTest(damage=damage):
                    book = load_workbook(package)
                    try:
                        backup = book[schema.ORIGINAL_FILE_SHEET]
                        if damage == "missing_sheet":
                            book.remove(backup)
                        elif damage == "missing_chunk":
                            backup.delete_rows(2)
                        elif damage == "duplicate_chunk":
                            backup["A3"] = backup["A2"].value
                        elif damage == "bad_base64":
                            backup["B2"] = "?" + backup["B2"].value[1:]
                        elif damage == "truncated_chunk":
                            backup["B2"] = backup["B2"].value[:-4]
                        elif damage == "changed_bytes":
                            data = bytearray(base64.b64decode(backup["B2"].value))
                            data[-1] ^= 1
                            backup["B2"] = base64.b64encode(data).decode("ascii")
                        elif damage == "invalid_index":
                            backup["A2"] = True
                        elif damage == "fractional_index":
                            backup["A2"] = 1.5
                        else:
                            key = schema.ORIGINAL_FILE_SIZE_KEY if damage == "wrong_size" else schema.ORIGINAL_FILE_SHA256_KEY
                            for row in book[schema.METADATA_SHEET]:
                                if row[0].value == key:
                                    row[1].value = 1 if damage == "wrong_size" else "0" * 64
                        book.save(edited)
                    finally:
                        book.close()
                    restored.write_bytes(b"existing destination")
                    with self.assertRaises(WorkbookFormatError):
                        restore_strings_workbook(edited, restored)
                    self.assertEqual(restored.read_bytes(), b"existing destination")

    def test_missing_legacy_and_future_versions_require_reexport(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            package = Path(directory) / "strings.xlsx"
            write_sparse_source(source)
            export_strings_workbook(source, package)
            for version in (None, "1.0", "1.1", "3.0"):
                with self.subTest(version=version):
                    book = load_workbook(package)
                    try:
                        for row in book[schema.METADATA_SHEET]:
                            if row[0].value == schema.SCHEMA_VERSION_KEY:
                                row[1].value = version
                        book.save(package)
                    finally:
                        book.close()
                    with self.assertRaisesRegex(WorkbookFormatError, "Export the original workbook again"):
                        load_strings_metadata(package)
