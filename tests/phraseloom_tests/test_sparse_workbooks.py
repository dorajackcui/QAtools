from __future__ import annotations

from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from phraseloom import workbook_schema as schema
from phraseloom.strings_package import read_strings_rows, read_mapping_rows
from phraseloom.strings_workflow import export_strings_workbook, restore_strings_workbook
from phraseloom.workbook_io import read_headers, read_source_rows


class SparseWorkbookTests(unittest.TestCase):
    def test_export_and_restore_keep_far_data_and_bounded_generated_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            package = Path(directory) / "strings.xlsx"
            restored = Path(directory) / "restored.xlsx"
            book = Workbook()
            try:
                sheet = book.active
                sheet.title = "Data"
                sheet.append(["source", "target", "context"])
                sheet["A2"] = "Early message"
                sheet["C2"] = "first context"
                sheet["A100001"] = "Distant message"
                sheet["C100001"] = "last context"
                sheet["XFD1048576"].number_format = "@"
                book.save(source)
            finally:
                book.close()
            with patch.object(ReadOnlyWorksheet, "iter_rows", side_effect=AssertionError("dense read")):
                self.assertEqual(read_headers(source), ["source", "target", "context"])
                rows = read_source_rows(source, "source", "target")
                self.assertEqual([row.row_number for row in rows], [2, 100001])
                self.assertEqual([len(row.original_values) for row in rows], [3, 3])
                stats = export_strings_workbook(source, package, context_col="context")
            self.assertEqual(stats["source_row_count"], 2)
            book = load_workbook(package)
            try:
                strings = book[schema.STRINGS_SHEET]
                mapping = book[schema.STRINGS_MAP_SHEET]
                self.assertEqual(strings.dimensions, "A1:G3")
                self.assertEqual(book[schema.COMPLETED_STRINGS_SHEET].dimensions, "A1:E1")
                self.assertEqual(mapping.dimensions, "A1:K3")
                self.assertEqual([mapping.cell(row, 1).value for row in (2, 3)], [2, 100001])
                self.assertEqual([strings.cell(row, 6).value for row in (2, 3)],
                                 ["first context", "last context"])
                strings["D2"] = "First translation"
                strings["D3"] = "Last translation"
                # Simulate editing a package: large physical gaps and formatting
                # must neither skip strings nor instantiate a full rectangle.
                strings.move_range("A3:G3", rows=100000)
                mapping.move_range("A3:K3", rows=100000)
                for sheet, reader in ((strings, read_strings_rows), (mapping, read_mapping_rows)):
                    sheet["XFD1048576"].number_format = "@"
                    before = set(sheet._cells)
                    with patch.object(sheet, "iter_rows", side_effect=AssertionError("dense read")):
                        self.assertEqual(len(reader(sheet)), 2)
                    self.assertEqual(set(sheet._cells), before)
                book.save(package)
            finally:
                book.close()
            result = restore_strings_workbook(package, restored)
            self.assertEqual(result["restored_row_count"], 2)
            self.assertEqual(result["issue_count"], 0)
            book = load_workbook(restored)
            try:
                sheet = book["Data"]
                self.assertEqual(sheet["B2"].value, "First translation")
                self.assertEqual(sheet["B100001"].value, "Last translation")
                self.assertNotIn((3, 2), sheet._cells)
                self.assertEqual(sheet["XFD1048576"].number_format, "@")
                self.assertLess(len(sheet._cells), 20)
            finally:
                book.close()

    def test_missing_target_is_created_after_values_not_after_ghost_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            package = Path(directory) / "strings.xlsx"
            restored = Path(directory) / "restored.xlsx"
            book = Workbook()
            try:
                book.active.append(["source"])
                book.active.append(["An untranslated message"])
                book.active["XFD1048576"].number_format = "@"
                book.save(source)
            finally:
                book.close()
            export_strings_workbook(source, package)
            book = load_workbook(package)
            try:
                book[schema.STRINGS_SHEET]["D2"] = "Translated"
                book.save(package)
            finally:
                book.close()
            restore_strings_workbook(package, restored)
            book = load_workbook(restored)
            try:
                self.assertEqual(book.active["B1"].value, "target")
                self.assertEqual(book.active["B2"].value, "Translated")
                self.assertEqual(book.active.max_column, 16384)
            finally:
                book.close()
