from __future__ import annotations

import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import Mock, patch
from zipfile import ZipFile

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook
from PySide6.QtCore import Qt
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QComboBox, QLineEdit

from tools.excel_preview import MAIN, ExcelPreview, read_excel_preview
from tools.qt_excel_preview import ColumnChoice, ExcelColumnSelector
from tools.qt_gui_common import AsyncPage, PathPicker


class ExcelPreviewTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.path = self.root / "input.xlsx"
        workbook = Workbook()
        try:
            worksheet = workbook.active
            worksheet.title = "First"
            worksheet.append(["Key", "Source", "Target"])
            for row in range(2, 20):
                worksheet.append([row, f"Source {row}", f"Target {row}"])
            other = workbook.create_sheet("Other")
            other["AA1"] = "标题"
            other["AA2"] = "正文"
            workbook.active = 1
            workbook.save(self.path)
        finally:
            workbook.close()

    def replace_parts(self, **parts):
        with ZipFile(self.path) as archive:
            contents = {name: archive.read(name) for name in archive.namelist()}
        contents.update(parts)
        with ZipFile(self.path, "w") as archive:
            for name, content in contents.items():
                archive.writestr(name, content)

    def test_active_explicit_and_first_sheet_keep_coordinates(self):
        active = read_excel_preview(str(self.path))
        self.assertEqual(active.sheet, "Other")
        self.assertEqual(active.columns, 27)
        self.assertEqual(active.cells[(0, 26)], "标题")
        first = read_excel_preview(str(self.path), first_sheet=True)
        self.assertEqual(first.sheet, "First")
        self.assertEqual(first.cells[(5, 1)], "Source 6")
        self.assertFalse(any(row > 5 for row, _ in first.cells))
        self.assertEqual(read_excel_preview(str(self.path), "First"), first)
        with self.assertRaisesRegex(ValueError, "工作表不存在"):
            read_excel_preview(str(self.path), "Missing")
        # All archive and partial XML stream handles are closed.
        self.path.rename(self.root / "renamed.xlsx")

    def test_prefix_read_skips_styles_dimensions_and_large_tail(self):
        xml = (f'<worksheet xmlns="{MAIN[1:-1]}"><dimension ref="A1:XFD1048576"/><sheetData>'
               '<row r="1"><c r="A1" t="inlineStr"><is><t>Source</t></is></c></row>'
               '<row r="6"><c r="C6" t="b"><v>0</v></c><c r="D6"><f>1+2</f><v>3</v></c></row>'
               '<row r="7"><c r="A7" t="inlineStr"><is><t>' + "x" * 3_000_000 +
               '</t></is></c></row></sheetData></worksheet>')
        self.replace_parts(**{"xl/worksheets/sheet1.xml": xml, "xl/styles.xml": b"unused invalid styles"})
        reads = {}
        original = ZipFile.open

        def track(archive, name, *args, **kwargs):
            stream = original(archive, name, *args, **kwargs)
            read = stream.read

            def counted(size=-1):
                chunk = read(size)
                reads[name] = reads.get(name, 0) + len(chunk)
                return chunk

            stream.read = counted
            return stream

        with patch.object(ZipFile, "open", track):
            result = read_excel_preview(str(self.path), "First")
        self.assertLess(reads["xl/worksheets/sheet1.xml"], 32768)
        self.assertNotIn("xl/styles.xml", reads)
        self.assertEqual(result.columns, 4)
        self.assertEqual(result.cells[(5, 2)], "FALSE")
        self.assertEqual(result.cells[(5, 3)], "=1+2")
        self.assertNotIn((6, 0), result.cells)

    def test_shared_strings_stop_when_needed_values_are_found(self):
        sheet = (f'<worksheet xmlns="{MAIN[1:-1]}"><sheetData><row r="1">'
                 '<c r="B1" t="s"><v>0</v></c></row></sheetData></worksheet>')
        strings = (f'<sst xmlns="{MAIN[1:-1]}"><si><r><t>原</t></r><r><t>文</t></r>'
                   '<rPh sb="0" eb="1"><t>ignored</t></rPh></si><si><t>' + "x" * 100000 + '</t></si></sst>')
        self.replace_parts(**{"xl/worksheets/sheet1.xml": sheet, "xl/sharedStrings.xml": strings})
        with patch("tools.excel_preview.STRINGS_BYTE_LIMIT", 16384):
            result = read_excel_preview(str(self.path), "First")
        self.assertEqual(result.cells[(0, 1)], "原文")
        self.assertEqual(result.note, "")

    def test_shared_headers_keep_full_identity_while_samples_are_clipped(self):
        header, sample = "Source " + "x" * 1500, "sample " + "y" * 1500
        sheet = (f'<worksheet xmlns="{MAIN[1:-1]}"><sheetData><row r="1">'
                 '<c r="A1" t="s"><v>0</v></c></row><row r="2">'
                 '<c r="A2" t="s"><v>1</v></c></row></sheetData></worksheet>')
        strings = f'<sst xmlns="{MAIN[1:-1]}"><si><t>{header}</t></si><si><t>{sample}</t></si></sst>'
        self.replace_parts(**{"xl/worksheets/sheet1.xml": sheet, "xl/sharedStrings.xml": strings})
        result = read_excel_preview(str(self.path), "First")
        self.assertEqual(result.headers[0], header)
        self.assertEqual(result.cells[(0, 0)], header[:1000])
        self.assertEqual(result.cells[(1, 0)], sample[:1000])

    def test_oversized_rows_and_shared_strings_have_bounded_fallback(self):
        sheet = (f'<worksheet xmlns="{MAIN[1:-1]}"><sheetData><row r="1">'
                 '<c r="A1" t="s"><v>1</v></c></row><row r="2"><c r="A2" t="inlineStr"><is><t>' +
                 "x" * 100000 + '</t></is></c></row></sheetData></worksheet>')
        strings = f'<sst xmlns="{MAIN[1:-1]}"><si><t>' + "x" * 100000 + '</t></si><si><t>late</t></si></sst>'
        self.replace_parts(**{"xl/worksheets/sheet1.xml": sheet, "xl/sharedStrings.xml": strings})
        with patch("tools.excel_preview.SHEET_BYTE_LIMIT", 16384), patch("tools.excel_preview.STRINGS_BYTE_LIMIT", 16384):
            result = read_excel_preview(str(self.path), "First")
        self.assertEqual(result.cells[(0, 0)], "[文本未展开]")
        self.assertIn("部分", result.note)
        self.assertIn("共享文本", result.note)

    def test_folder_chooses_one_eligible_file_and_excludes_backups_and_outputs(self):
        excluded = self.root / "excluded.xlsx"
        excluded.touch()
        (self.root / "~$locked.xlsx").touch()
        (self.root / "未翻译统计.xlsx").touch()
        backup = self.root / ".qatools-backups"
        backup.mkdir()
        (backup / "backup.xlsx").touch()
        with patch("tools.excel_preview.random.choice", side_effect=lambda files: files[-1]) as choose:
            result = read_excel_preview(str(self.root), folder=True, exclude=(str(excluded),),
                                        exclude_names=("未翻译统计.xlsx",))
        self.assertEqual(choose.call_args.args[0], [self.path])
        self.assertEqual(result.path, self.path)
        with self.assertRaisesRegex(ValueError, "没有可预览"):
            read_excel_preview(str(backup), folder=True, exclude=(str(backup),))

    def test_first_sheet_preview_uses_first_data_sheet_when_a_chart_is_first(self):
        from openpyxl.chart import BarChart, Reference
        workbook = Workbook()
        try:
            data = workbook.active
            data.title = "Data"
            data.append(["Source", "Target"])
            data.append([1, 2])
            chart = BarChart()
            chart.add_data(Reference(data, min_col=1, min_row=1, max_row=2), titles_from_data=True)
            chartsheet = workbook.create_chartsheet("Chart", index=0)
            chartsheet.add_chart(chart)
            workbook.active = 0
            workbook.save(self.path)
        finally:
            workbook.close()
        result = read_excel_preview(str(self.path), first_sheet=True)
        self.assertEqual(result.sheet, "Data")
        self.assertEqual(result.sheet_names, ("Data",))
        self.assertEqual(result.cells[(0, 0)], "Source")


class ExcelPreviewGuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.page = AsyncPage()
        self.addCleanup(self.page.close)
        self.picker = PathPicker("Input", parent=self.page)
        self.sheet = QComboBox(self.page)
        self.sheet.addItems(["One", "Two"])
        self.selector = ExcelColumnSelector(self.page, self.picker, self.sheet)

    def wait_for_preview(self, page):
        deadline = time.monotonic() + 5
        while page.has_running_tasks() and time.monotonic() < deadline:
            QTest.qWait(10)
        self.assertFalse(page.has_running_tasks())

    def test_restoring_paths_does_not_read_and_old_results_are_discarded(self):
        self.page.run_in_background = Mock()
        with patch("tools.excel_preview.read_excel_preview") as read:
            self.picker.set_path("first.xlsx")
            self.sheet.setCurrentText("Two")
            self.app.processEvents()
            read.assert_not_called()
            self.page.run_in_background.assert_not_called()
            self.selector.open_preview()
        pending = self.page.run_in_background.call_args.kwargs
        self.assertEqual(pending["args"], ("first.xlsx", "Two"))
        self.picker.set_path("second.xlsx")
        pending["on_success"](ExcelPreview(Path("first.xlsx"), "Two", ("Two",), {(0, 0): "old"}, 1))
        self.assertIsNone(self.selector.model.preview)
        self.assertFalse(self.selector.dialog.isVisible())
        self.selector.open_preview()
        pending = self.page.run_in_background.call_args.kwargs
        self.sheet.setCurrentText("One")
        pending["on_error"]("old error")
        self.assertNotIn("old error", self.selector.status.text())

    def test_preview_runs_off_gui_thread_and_displays_letter_and_number(self):
        import threading
        gui_thread = threading.get_ident()
        result = ExcelPreview(Path("input.xlsx"), "One", ("One",), {(0, 26): "Header", (5, 26): "Sample"}, 27)

        def read(*args, **kwargs):
            self.assertNotEqual(threading.get_ident(), gui_thread)
            time.sleep(0.05)
            return result

        self.picker.set_path("input.xlsx")
        with patch("tools.excel_preview.read_excel_preview", side_effect=read):
            self.selector.open_preview()
            self.wait_for_preview(self.page)
        self.assertIs(self.selector.model.preview, result)
        self.assertEqual(self.selector.model.headerData(26, Qt.Orientation.Horizontal), "AA (27)")
        self.assertEqual(self.selector.model.index(5, 26).data(), "Sample")
        self.assertEqual(self.selector.model.rowCount(), 6)

    def test_all_column_selection_pages_read_the_corresponding_input(self):
        from phraseloom.qt_page import PhraseLoomPage
        from tools.content_sync.qt_page import MasterToTargetPage, TargetToMasterPage
        from tools.excel_utilities_pages import ColumnToolsPage, UntranslatedStatsPage
        from tools.file_grouping.qt_page import FileGroupingPage
        from tools.french_nbsp_restorer.qt_page import FrenchNbspPage
        from tools.workflow.qt_page import WorkflowPage
        from tools.workflow.gui_options import WorkflowOptionsStore

        ExcelPreviewTests.setUp(self)
        workflow = WorkflowPage(options_store=WorkflowOptionsStore(self.root / "options.json"))
        phrase, nbsp, grouping = PhraseLoomPage(), FrenchNbspPage(), FileGroupingPage()
        columns, stats = ColumnToolsPage(), UntranslatedStatsPage()
        forward, reverse = MasterToTargetPage(), TargetToMasterPage()
        for page in (workflow, phrase, nbsp, grouping, columns, stats, forward, reverse):
            self.addCleanup(page.close)
        cases = [
            (workflow, workflow.input_preview, workflow.input_picker, workflow.sheet, False, "First"),
            (workflow, workflow.history_preview, workflow.history_picker, workflow.history_sheet, False, "First"),
            (phrase, phrase.input_preview, phrase.input_picker, None, False, "First"),
            (nbsp, nbsp.input_preview, nbsp.input_picker, nbsp.sheet, False, "First"),
            (columns, columns.input_preview, columns.input_picker, columns.sheet, True, "First"),
            (stats, stats.input_preview, stats.input_picker, stats.sheet, True, "First"),
            (grouping, grouping.input_preview, grouping.source_picker, None, True, "Other"),
        ]
        for page in (forward, reverse):
            cases.extend([
                (page, page.master_preview, page.master_picker, page.master_sheet, False, "First"),
                (page, page.target_preview, page.target_picker, page.target_sheet, True, "First"),
            ])
        for page, selector, picker, sheet, folder, expected_sheet in cases:
            with self.subTest(page=type(page).__name__, selector=selector.title):
                # Master must be outside the small-table directory for folder cases.
                if selector is getattr(page, "target_preview", None):
                    page.master_picker.set_path(str(self.root.parent / "master.xlsx"))
                picker.set_path(str(self.root if folder else self.path))
                if isinstance(sheet, QComboBox):
                    sheet.addItem("First")
                    sheet.setCurrentText("First")
                elif sheet is not None:
                    sheet.setText("First")
                selector.open_preview()
                self.wait_for_preview(page)
                self.assertIsNotNone(selector.model.preview, selector.status.text())
                self.assertEqual(selector.model.preview.sheet, expected_sheet)
                self.assertEqual(selector.model.preview.path, self.path)
                for choice, combo in zip(selector.choices, selector.column_choices):
                    if choice.field.isHidden():
                        continue
                    value = "标题" if expected_sheet == "Other" and choice.kind == "header" else (
                        "Source" if choice.kind == "header" else "2" if choice.kind == "index" else "B")
                    combo.setCurrentIndex(combo.findData(value))
                selector.apply_button.click()
                self.assertFalse(selector.dialog.isVisible())
                for choice, combo in zip(selector.choices, selector.column_choices):
                    if not choice.field.isHidden():
                        self.assertEqual(choice.field.text(), combo.currentData())
                selector.dialog.reject()
        # File-grouping previews can inspect each sheet of the same sample.
        grouping.input_preview.open_preview()
        self.wait_for_preview(grouping)
        grouping.input_preview.sheet_choice.setCurrentText("First")
        self.wait_for_preview(grouping)
        self.assertEqual(grouping.input_preview.model.preview.sheet, "First")
        grouping.input_preview.dialog.reject()
        grouping.input_kind.setCurrentIndex(1)
        grouping.source_picker.set_path(str(self.path))
        grouping.input_preview.open_preview()
        self.wait_for_preview(grouping)
        self.assertFalse(grouping.input_preview.folder)
        self.assertEqual(grouping.input_preview.model.preview.path, self.path)

    def test_cancel_does_not_change_columns_and_optional_columns_can_be_cleared(self):
        source, context = QLineEdit("A", self.page), QLineEdit("D", self.page)
        selector = ExcelColumnSelector(self.page, self.picker, self.sheet,
                                    choices=(ColumnChoice("Source", source),
                                             ColumnChoice("Context", context, optional=True)))
        self.page.run_in_background = Mock()
        self.picker.set_path("input.xlsx")
        preview = ExcelPreview(Path("input.xlsx"), "One", ("One", "Two"),
                               {(0, 0): "Key", (0, 1): "Source", (0, 2): "Context"}, 3)
        selector.open_preview()
        self.page.run_in_background.call_args.kwargs["on_success"](preview)
        selector.column_choices[0].setCurrentIndex(selector.column_choices[0].findData("B"))
        selector.dialog.reject()
        self.assertEqual(source.text(), "A")
        self.assertEqual(context.text(), "D")
        selector.open_preview()
        self.page.run_in_background.call_args.kwargs["on_success"](preview)
        selector.column_choices[0].setCurrentIndex(selector.column_choices[0].findData("B"))
        selector.column_choices[1].setCurrentIndex(0)
        changed = Mock()
        self.sheet.currentTextChanged.connect(changed)
        selector.apply_button.click()
        self.assertEqual(source.text(), "B")
        self.assertEqual(context.text(), "")
        changed.assert_not_called()

    def test_folder_selection_preserves_each_files_active_sheet_setting(self):
        self.page.run_in_background = Mock()
        sheet = QLineEdit("", self.page)
        source = QLineEdit("B", self.page)
        selector = ExcelColumnSelector(self.page, self.picker, sheet, folder=True,
                                    choices=(ColumnChoice("Source", source),))
        self.picker.set_path("folder")
        selector.open_preview()
        preview = ExcelPreview(Path("folder/one.xlsx"), "One", ("One", "Two"), {(0, 1): "Source"}, 2)
        self.page.run_in_background.call_args.kwargs["on_success"](preview)
        selector.apply_button.click()
        self.assertEqual(sheet.text(), "")
        selector.open_preview()
        self.page.run_in_background.call_args.kwargs["on_success"](preview)
        selector.sheet_choice.setCurrentText("Two")
        pending = self.page.run_in_background.call_args.kwargs
        self.assertEqual(pending["args"], (str(preview.path), "Two"))
        self.assertFalse(pending["kwargs"]["folder"])
        pending["on_success"](ExcelPreview(preview.path, "Two", preview.sheet_names, preview.cells, 2))
        selector.apply_button.click()
        self.assertEqual(sheet.text(), "Two")

    def test_manual_file_and_folder_choosers_open_column_selection_automatically(self):
        from tools.workflow.qt_page import WorkflowPage
        from tools.workflow.gui_options import WorkflowOptionsStore
        from tools.french_nbsp_restorer.qt_page import FrenchNbspPage
        from phraseloom.qt_page import PhraseLoomPage
        from tools.excel_utilities_pages import ColumnToolsPage, UntranslatedStatsPage
        from tools.file_grouping.qt_page import FileGroupingPage
        ExcelPreviewTests.setUp(self)
        workflow = WorkflowPage(options_store=WorkflowOptionsStore(self.root / "options.json"))
        nbsp, phrase = FrenchNbspPage(), PhraseLoomPage()
        columns, stats, grouping = ColumnToolsPage(), UntranslatedStatsPage(), FileGroupingPage()
        for page in (workflow, nbsp, phrase, columns, stats, grouping):
            self.addCleanup(page.close)
        for page, module, choose in (
            (workflow, "tools.workflow.qt_page", workflow.choose_input_file),
            (nbsp, "tools.french_nbsp_restorer.qt_page", nbsp.choose_input),
            (phrase, "phraseloom.qt_page", phrase.choose_input),
        ):
            with self.subTest(page=type(page).__name__), \
                 patch(f"{module}._choose_excel", return_value=str(self.path)), \
                 patch("openpyxl.load_workbook", side_effect=AssertionError("synchronous workbook load")):
                choose()
                self.wait_for_preview(page)
                self.assertTrue(page.input_preview.dialog.isVisible())
                self.assertIsNotNone(page.input_preview.model.preview)
                page.input_preview.dialog.reject()
        with patch("tools.workflow.qt_page._choose_excel", return_value=str(self.path)), \
             patch("openpyxl.load_workbook", side_effect=AssertionError("synchronous TB load")):
            workflow.choose_history_file()
            self.wait_for_preview(workflow)
        self.assertTrue(workflow.history_preview.dialog.isVisible())
        workflow.history_preview.dialog.reject()
        for page in (columns, stats):
            with patch("tools.excel_utilities_pages.QFileDialog.getExistingDirectory", return_value=str(self.root)):
                page.choose_folder(page.input_picker)
            self.wait_for_preview(page)
            self.assertTrue(page.input_preview.dialog.isVisible())
            page.input_preview.dialog.reject()
        with patch("tools.file_grouping.qt_page.QFileDialog.getExistingDirectory", return_value=str(self.root)):
            grouping.choose_source()
        self.wait_for_preview(grouping)
        self.assertTrue(grouping.input_preview.dialog.isVisible())

    def test_missing_previous_sheet_can_be_reselected_in_folder_preview(self):
        ExcelPreviewTests.setUp(self)
        sheet, source = QLineEdit("PreviousSheet", self.page), QLineEdit("B", self.page)
        selector = ExcelColumnSelector(self.page, self.picker, sheet, folder=True,
                                       choices=(ColumnChoice("Source", source),))
        self.picker.set_path(str(self.root))
        selector.open_preview()
        self.wait_for_preview(self.page)
        self.assertIsNotNone(selector.model.preview, selector.status.text())
        self.assertIn("PreviousSheet", selector.status.text())
        self.assertEqual(sheet.text(), "PreviousSheet")
        selector.sheet_choice.setCurrentText("First")
        self.wait_for_preview(self.page)
        selector.apply_button.click()
        self.assertEqual(sheet.text(), "First")
        self.assertEqual(source.text(), "B")

    def test_invalid_alias_config_keeps_manual_column_selection_available(self):
        from tools.header_aliases import HeaderAliasStore
        ExcelPreviewTests.setUp(self)
        aliases = self.root / "aliases.json"
        aliases.write_text("invalid json", encoding="utf-8")
        source = QLineEdit("", self.page)
        selector = ExcelColumnSelector(self.page, self.picker, self.sheet,
                                       choices=(ColumnChoice("Source", source, alias_role="source"),),
                                       header_alias_store=HeaderAliasStore(aliases))
        self.page.run_in_background = Mock()
        self.picker.set_path(str(self.path))
        selector.open_preview()
        preview = ExcelPreview(self.path, "One", ("One",), {(0, 0): "Key", (0, 1): "Source"}, 2)
        self.page.run_in_background.call_args.kwargs["on_success"](preview)
        self.assertIn("别名", selector.status.text())
        self.assertTrue(selector.apply_button.isEnabled())
        selector.column_choices[0].setCurrentIndex(selector.column_choices[0].findData("B"))
        selector.apply_button.click()
        self.assertEqual(source.text(), "B")

    def test_long_headers_are_display_limited_but_applied_in_full(self):
        from tools.file_grouping.grouping import scan_file_groups
        ExcelPreviewTests.setUp(self)
        long_header = "Source " + "x" * 1500
        workbook = Workbook()
        try:
            workbook.active.append([long_header, "Target"])
            workbook.active.append(["待翻译", None])
            workbook.save(self.path)
        finally:
            workbook.close()
        source, target = QLineEdit("", self.page), QLineEdit("Target", self.page)
        selector = ExcelColumnSelector(self.page, self.picker, first_sheet=True,
                                       choices=(ColumnChoice("Source", source, "header"),
                                                ColumnChoice("Target", target, "header")))
        self.picker.set_path(str(self.path))
        selector.open_preview()
        self.wait_for_preview(self.page)
        choice = selector.column_choices[0]
        index = choice.findData(long_header)
        self.assertGreater(index, 0)
        self.assertLess(len(choice.itemText(index)), 1100)
        self.assertEqual(len(selector.model.index(0, 0).data()), 1000)
        choice.setCurrentIndex(index)
        selector.apply_button.click()
        self.assertEqual(source.text(), long_header)
        result = scan_file_groups(self.root, source_header=source.text(), target_header=target.text())
        self.assertTrue(result.can_export, result.issues)


if __name__ == "__main__":
    unittest.main()
