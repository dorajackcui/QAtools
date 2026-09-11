from __future__ import annotations

import os
from pathlib import Path
import tempfile
import threading
import time
import unittest
from unittest.mock import Mock, patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from PySide6.QtCore import QCoreApplication, QEvent, QTimer
from PySide6.QtWidgets import QApplication
from openpyxl import Workbook, load_workbook

from tools.content_sync.master_to_target import FileResult, SyncSummary
from tools.content_sync.qt_page import ContentSyncPage, MasterToTargetPage


class ContentSyncGuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        self.page = MasterToTargetPage()
        self.addCleanup(self.dispose_page, self.page)

    @staticmethod
    def dispose_page(page):
        # Destroy Qt timers on the GUI thread before a later worker can trigger GC.
        page.close()
        page.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)

    def test_page_construction_does_not_load_workbook_or_com(self):
        with patch("openpyxl.load_workbook", side_effect=AssertionError("eager IO")), \
             patch("tools.content_sync.qt_page._master_sheets", side_effect=AssertionError("eager probe")):
            page = ContentSyncPage()
            self.dispose_page(page)
        self.assertEqual(self.page.run_button.parentWidget().objectName(), "pageActionBar")
        self.assertFalse(self.page.fill_blank_only.isChecked())
        self.assertFalse(self.page.allow_blank_write.isChecked())

    def test_direction_tabs_preserve_independent_inputs_results_and_logs(self):
        page = ContentSyncPage()
        self.addCleanup(self.dispose_page, page)
        self.assertEqual([page.tabs.tabText(i) for i in range(page.tabs.count())],
                         ["Master → 小表", "小表 → Master"])
        for index in (0, 1):
            page.tabs.setCurrentIndex(index)
            direction = page.tabs.currentWidget()
            direction.master_picker.set_path(f"master-{index}.xlsx")
            direction.result.setPlainText(f"result-{index}")
            direction.logs.start_run(f"log-{index}")
            direction.logs.finish_run("done")
        for index in (0, 1):
            page.tabs.setCurrentIndex(index)
            direction = page.tabs.currentWidget()
            self.assertEqual(direction.master_picker.path(), f"master-{index}.xlsx")
            self.assertEqual(direction.result.toPlainText(), f"result-{index}")
            direction.logs.open_logs()
            self.app.processEvents()
            self.assertIn(f"log-{index}", direction.logs.text.toPlainText())
            self.assertNotIn(f"log-{1 - index}", direction.logs.text.toPlainText())
            direction.logs.close()

    def test_run_captures_widget_values_before_dispatch(self):
        self.page.master_picker.set_path("master.xlsx")
        self.page.target_picker.set_path("targets")
        self.page.column_count.setValue(2)
        self.page.fill_blank_only.setChecked(True)
        self.page.run_in_background = Mock()
        self.page.run_sync()
        kwargs = self.page.run_in_background.call_args.kwargs["kwargs"]
        self.page.column_count.setValue(3)
        self.page.target_sheet.setText("changed afterwards")
        self.assertEqual(kwargs["column_count"], 2)
        self.assertIsNone(kwargs["target_sheet"])
        self.assertTrue(kwargs["fill_blank_only"])
        self.assertFalse(self.page.run_button.isEnabled())
        self.assertFalse(self.page.content.isEnabled())

    def test_missing_input_and_overlapping_columns_do_not_start_worker(self):
        self.page.run_in_background = Mock()
        with patch("tools.content_sync.qt_page.show_error") as error:
            self.page.run_sync()
            error.assert_called_once()
        self.page.master_picker.set_path("master.xlsx")
        self.page.target_picker.set_path("targets")
        self.page.master_columns[0].setText("D")
        with patch("tools.content_sync.qt_page.show_error") as error:
            self.page.run_sync()
            error.assert_called_once()
        self.page.run_in_background.assert_not_called()

    def test_directory_selection_dispatches_snapshot_and_displays_count(self):
        from tools.content_sync.preflight import DirectoryCheck
        self.page.master_picker.set_path("master.xlsx")
        self.page.run_in_background = Mock()
        with patch("tools.content_sync.qt_page.QFileDialog.getExistingDirectory", return_value="targets"):
            self.page.choose_targets()
        call = self.page.run_in_background.call_args.kwargs
        self.assertEqual(call["args"], ("targets",))
        self.assertEqual(call["kwargs"], {"master": "master.xlsx", "output": "", "reverse": False})
        self.assertFalse(self.page.run_button.isEnabled())
        with patch("tools.content_sync.qt_page.show_info") as info:
            call["on_success"](DirectoryCheck(32, 1, tuple(str(i) for i in range(20)), (), ()))
        self.assertIn("32", info.call_args.args[2])
        self.assertIn("20 / 32", info.call_args.args[2])
        self.assertTrue(self.page.run_button.isEnabled())

    def test_master_warning_preserves_sheet_selection_and_selection_errors_preserve_run_logs(self):
        from tools.content_sync.preflight import MasterCheck
        from tools.excel_metadata import WorkbookSheetChoices
        self.page.run_in_background = Mock()
        with patch("tools.content_sync.qt_page.QFileDialog.getOpenFileName", return_value=("master.xlsx", "")):
            self.page.choose_master()
        callback = self.page.run_in_background.call_args.kwargs["on_success"]
        with patch("tools.content_sync.qt_page.show_warning") as warning:
            callback(MasterCheck(WorkbookSheetChoices(("Data",), "Data"), ("Master 被占用",)))
        self.assertIn("占用", warning.call_args.args[2])
        self.assertEqual(self.page.master_sheet.currentText(), "Data")
        self.page.result.setPlainText("previous result")
        with patch("tools.content_sync.qt_page.show_warning"):
            self.page._selection_failed("scan failed")
        self.assertEqual(self.page.result.toPlainText(), "previous result")
        self.assertTrue(self.page.run_button.isEnabled())

    def test_directory_scan_is_background_and_gui_remains_responsive(self):
        from tools.content_sync.preflight import DirectoryCheck
        release, started = threading.Event(), threading.Event()
        owner = threading.get_ident()
        threads = []
        def scan(*args, **kwargs):
            threads.append(threading.get_ident())
            started.set()
            if not release.wait(5):
                raise RuntimeError("GUI did not release scan")
            return DirectoryCheck(1, 0, ("a.xlsx",), (), ())
        with patch("tools.content_sync.qt_page.QFileDialog.getExistingDirectory", return_value="targets"), \
             patch("tools.content_sync.qt_page._target_check", scan), \
             patch("tools.content_sync.qt_page.show_info") as info, \
             patch("tools.content_sync.qt_page.show_warning") as warning:
            try:
                self.page.choose_targets()
                deadline = time.monotonic() + 10
                while self.page.has_running_tasks() and time.monotonic() < deadline:
                    self.app.processEvents()
                    if started.is_set() and not release.is_set():
                        self.assertTrue(self.page.has_running_tasks())
                        release.set()
                    time.sleep(0.005)
            finally:
                release.set()
                self.page._thread_pool.waitForDone(10000)
                self.app.processEvents()
            info.assert_called_once()
            warning.assert_not_called()
        self.assertTrue(threads and all(thread != owner for thread in threads))
        self.assertFalse(self.page.has_running_tasks())

    def test_partial_failure_is_visible_and_controls_are_restored(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            summary = SyncSummary(root, 1, [],
                                  [FileResult("broken.xlsx", error="locked")])
            self.page._set_busy(True)
            with patch("tools.content_sync.qt_page.show_warning") as warning:
                self.page._finish(summary)
                warning.assert_called_once()
            self.assertIn("失败: 1", self.page.result.toPlainText())
            self.assertTrue(self.page.run_button.isEnabled())

    def test_real_background_run_generates_output_and_updates_gui(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            targets = root / "targets"
            targets.mkdir()
            master = root / "master.xlsx"
            for path, rows in (
                (master, [["id", "key", "match", "value"], [1, "k", "s", "  nan  "]]),
                (targets / "a.xlsx", [["key", "match", "value"], ["k", "s", "old"]]),
            ):
                wb = Workbook()
                for row in rows:
                    wb.active.append(row)
                wb.save(path)
                wb.close()
            self.page.master_picker.set_path(str(master))
            self.page.target_picker.set_path(str(targets))
            with patch("tools.content_sync.qt_page.show_error") as error:
                self.page.run_sync()
                deadline = time.monotonic() + 10
                while self.page.has_running_tasks() and time.monotonic() < deadline:
                    self.app.processEvents()
                    time.sleep(0.005)
                self.assertFalse(self.page.has_running_tasks(), "worker did not finish")
                error.assert_not_called()
            self.assertTrue(self.page.run_button.isEnabled())
            self.assertIn("实际更新单元格: 1", self.page.result.toPlainText())
            self.assertEqual(self.page.output_picker.path(), "")
            wb = load_workbook(targets / "a.xlsx")
            try:
                self.assertEqual(wb.active["C2"].value, "  nan  ")
            finally:
                wb.close()

    def test_parallel_files_leave_gui_timer_and_live_logs_responsive(self):
        from tools.content_sync import master_to_target as sync
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            targets = root / "targets"
            targets.mkdir()
            master = root / "master.xlsx"
            for path, row in ((master, [1, "k", "s", "new"]),
                              (targets / "a.xlsx", ["k", "s", "old"]),
                              (targets / "b.xlsx", ["k", "s", "old"])):
                book = Workbook()
                try:
                    book.active.append(["header"])
                    book.active.append(row)
                    book.save(path)
                finally:
                    book.close()
            release = threading.Event()
            started = [threading.Event(), threading.Event()]
            original = sync._sync_file
            ticks = []

            def process(source, *args, **kwargs):
                started[0 if source.stem == "a" else 1].set()
                if not release.wait(5):
                    raise RuntimeError("GUI timer did not respond")
                return original(source, *args, **kwargs)

            def heartbeat():
                if all(event.is_set() for event in started):
                    ticks.append(self.page.has_running_tasks())
                    self.page.logs.open_logs()
                    self.page.logs.flush()
                    release.set()

            self.page.master_picker.set_path(str(master))
            self.page.target_picker.set_path(str(targets))
            timer = QTimer(self.page)
            timer.setInterval(10)
            timer.timeout.connect(heartbeat)
            with patch.object(sync, "_sync_file", process), \
                 patch("tools.content_sync.qt_page.show_error") as error, \
                 patch("tools.content_sync.qt_page.show_warning") as warning:
                try:
                    timer.start()
                    self.page.run_sync()
                    deadline = time.monotonic() + 10
                    while self.page.has_running_tasks() and time.monotonic() < deadline:
                        self.app.processEvents()
                        time.sleep(0.005)
                finally:
                    release.set()
                    timer.stop()
                    self.page._thread_pool.waitForDone(10000)
                    self.app.processEvents()
                error.assert_not_called()
                warning.assert_not_called()
            self.assertIn(True, ticks)
            self.assertFalse(self.page.has_running_tasks())
            self.page.logs.flush()
            self.assertIn("更新 1 格", self.page.logs.text.toPlainText())
            self.assertIn("实际更新单元格: 2", self.page.result.toPlainText())


if __name__ == "__main__":
    unittest.main()
