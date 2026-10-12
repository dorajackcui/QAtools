from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from tools.file_grouping.merged import export_merged_groups, scan_merged_table


class MergedGroupingTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.input = self.root / "merged.xlsx"
        self.output = self.root / "output"

    def book(self, rows, *, headers=("File", "MsgStr", "fr", "Note"), extra=(), write_only=False):
        book = Workbook(write_only=write_only)
        try:
            sheet = book.create_sheet("Data") if write_only else book.active
            sheet.title = "Data"
            sheet.append(headers)
            for row in rows:
                sheet.append(row)
            for title, other_rows in extra:
                other = book.create_sheet(title)
                for row in other_rows:
                    other.append(row)
            book.save(self.input)
        finally:
            book.close()
        return self.input

    def read(self, path):
        book = load_workbook(path, read_only=True)
        try:
            return {sheet.title: [tuple(cell.value for cell in row) for row in sheet.iter_rows()]
                    for sheet in book.worksheets}
        finally:
            book.close()

    def test_groups_use_pending_sources_but_exports_keep_all_rows_columns_and_order(self):
        date = datetime(2026, 10, 12)
        rows = [["a.xlsx", "共同", None, "a pending", date],
                ["d.xlsx", "共同", "translated", "d translated", None],
                ["b.xlsx", "共同", None, "b pending", None],
                ["a.xlsx", "已完成", "done", "keep translation", None],
                ["b.xlsx", "桥梁", None, "bridge", None],
                ["c.xlsx", "桥梁", None, "c pending", None],
                ["c.xlsx", None, "metadata", "keep empty source", None],
                ["e.xlsx", "独立", None, "other pending", None],
                ["a.xlsx", None, None, "keep file metadata", None]]
        self.book(rows, headers=("File", "MsgStr", "fr", "Note", "Date"))
        book = load_workbook(self.input)
        try:
            book.active["A1"].font = Font(bold=True)
            book.active["B2"].fill = PatternFill("solid", fgColor="FFFF00")
            book.active["C2"].fill = PatternFill("solid", fgColor="EEEEEE")
            book.save(self.input)
        finally:
            book.close()
        before = self.input.read_bytes()
        result = scan_merged_table(self.input)
        self.assertTrue(result.can_export)
        self.assertEqual(result.row_count, len(rows))
        self.assertEqual([(len(group.files), group.unique_chars) for group in result.groups], [(3, 4), (2, 2)])
        self.assertEqual(sum(file.pending_rows for file in result.files), 5)
        summary = export_merged_groups(result, self.output)
        self.assertEqual(summary.file_count, 2)
        for name, filenames in (("文件群1", {"a.xlsx", "b.xlsx", "c.xlsx"}), ("其他", {"d.xlsx", "e.xlsx"})):
            exported = self.read(self.output / f"{name}.xlsx")["Data"]
            self.assertEqual(exported[0], ("File", "MsgStr", "fr", "Note", "Date"))
            # Streamed XLSX can omit trailing empty cells; compare padded row contents.
            actual = [tuple(row) + (None,) * (5 - len(row)) for row in exported[1:]]
            self.assertEqual(actual, [tuple(row) for row in rows if row[0] in filenames])
        book = load_workbook(self.output / "文件群1.xlsx")
        try:
            self.assertTrue(book.active["A1"].font.bold)
            self.assertEqual(book.active["B2"].fill.fgColor.rgb, "00FFFF00")
            self.assertEqual(book.active["C2"].fill.fgColor.rgb, "00EEEEEE")
            self.assertEqual(book.active["E2"].value, date)
        finally:
            book.close()
        self.assertEqual(self.input.read_bytes(), before)

    def test_case_slashes_and_dot_segments_identify_one_virtual_file(self):
        self.book([["独立", " Folder\\A.xlsx ", None], ["独立", "folder/./a.xlsx", None],
                   [None, "folder//A.xlsx", "done"]], headers=("MsgStr", "File", "fr"))
        result = scan_merged_table(self.input, filename_column="b")
        self.assertTrue(result.can_export)
        self.assertEqual(len(result.files), 1)
        self.assertEqual(result.files[0].path.as_posix(), "Folder/A.xlsx")
        self.assertEqual(result.files[0].pending_rows, 2)
        export_merged_groups(result, self.output)
        self.assertEqual(len(self.read(self.output / "其他.xlsx")["Data"]), 4)

    def test_dimensionless_tables_reversed_columns_and_extra_data_columns(self):
        self.book([[None, "a.xlsx", "共同", "outside header"],
                   [None, "b.xlsx", "共同"], [False, "c.xlsx", "共同"]],
                  headers=("fr", "File", "MsgStr"), write_only=True)
        result = scan_merged_table(self.input, filename_column="B")
        self.assertFalse(result.issues)
        self.assertEqual([len(group.files) for group in result.groups], [2, 1])
        self.assertEqual(sum(file.pending_rows for file in result.files), 2)
        export_merged_groups(result, self.output)
        rows = self.read(self.output / "文件群1.xlsx")["Data"]
        self.assertEqual(rows[1], (None, "a.xlsx", "共同", "outside header"))

    def test_multiple_matching_sheets_share_file_identity(self):
        self.book([["a.xlsx", "共同", None]], extra=(
            ("OtherData", [["File", "MsgStr", "fr"], ["b.xlsx", "共同", None], ["a.xlsx", "old", "done"]]),
            ("Readme", [["说明"], ["irrelevant"]]),))
        result = scan_merged_table(self.input)
        self.assertEqual(len(result.groups[0].files), 2)
        export_merged_groups(result, self.output)
        exported = self.read(self.output / "文件群1.xlsx")
        self.assertEqual(set(exported), {"Data", "OtherData"})
        self.assertEqual(len(exported["OtherData"]), 3)

    def test_missing_filename_and_formulas_block_export(self):
        for rows in ([[None, "原文", None]], [["a.xlsx", "原文", None, "=A2"]],
                     [["=A3", "原文", None]], [["a.xlsx", "=B3", None]]):
            with self.subTest(rows=rows):
                self.book(rows)
                result = scan_merged_table(self.input)
                self.assertFalse(result.can_export)
                self.assertTrue(result.issues)
                with self.assertRaises(ValueError):
                    export_merged_groups(result, self.output)
                self.assertFalse(self.output.exists())

    def test_literal_formula_like_text_and_typed_source_values_are_preserved(self):
        self.book([["a.xlsx", "=literal", None], ["b.xlsx", "=literal", None],
                   ["a.xlsx", 0, None], ["b.xlsx", False, None], ["b.xlsx", "nan", "nan"]])
        book = load_workbook(self.input)
        try:
            book.active["B2"].data_type = "s"
            book.active["B3"].data_type = "s"
            book.save(self.input)
        finally:
            book.close()
        result = scan_merged_table(self.input)
        self.assertTrue(result.can_export)
        self.assertEqual(sum(file.pending_rows for file in result.files), 4)
        export_merged_groups(result, self.output)
        book = load_workbook(self.output / "文件群1.xlsx")
        try:
            self.assertEqual(book.active["B2"].value, "=literal")
            self.assertEqual(book.active["B2"].data_type, "s")
            self.assertEqual(book.active["B4"].value, 0)
            self.assertIs(book.active["B5"].value, False)
        finally:
            book.close()

    def test_changed_input_requires_rescan_before_output(self):
        self.book([["a.xlsx", "原文", None]])
        result = scan_merged_table(self.input)
        self.book([["a.xlsx", "改后", None]])
        with self.assertRaisesRegex(ValueError, "重新扫描"):
            export_merged_groups(result, self.output)
        self.assertFalse(self.output.exists())

    def test_existing_destination_created_during_export_is_not_overwritten(self):
        self.book([["a.xlsx", "原文", None]])
        result = scan_merged_table(self.input)
        from tools.file_grouping import merged
        original = merged._check_input
        calls = []

        def race(value):
            calls.append(None)
            if len(calls) == 3:
                (self.output / "其他.xlsx").write_bytes(b"keep existing")
            original(value)

        with patch.object(merged, "_check_input", side_effect=race):
            with self.assertRaisesRegex(RuntimeError, "未覆盖"):
                export_merged_groups(result, self.output)
        self.assertEqual((self.output / "其他.xlsx").read_bytes(), b"keep existing")
        self.assertFalse(list(self.output.glob("*.tmp")))

    def test_invalid_headers_filename_column_and_output_directory(self):
        self.book([["a.xlsx", "原文", None]])
        for column in ("", "A1", "XFE"):
            with self.subTest(column=column), self.assertRaises(ValueError):
                scan_merged_table(self.input, filename_column=column)
        for column in ("B", "C", "D", "E"):
            result = scan_merged_table(self.input, filename_column=column)
            self.assertFalse(result.can_export)
        result = scan_merged_table(self.input, source_header="wrong")
        self.assertFalse(result.can_export)
        result = scan_merged_table(self.input)
        with self.assertRaises(ValueError):
            export_merged_groups(result, self.root)
        self.output.mkdir()
        with self.assertRaises(ValueError):
            export_merged_groups(result, self.output)


if __name__ == "__main__":
    unittest.main()
