"""Ghost dimensions must not create work, lose data, or shift coordinates."""
from __future__ import annotations

from datetime import datetime
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.cell.rich_text import CellRichText
from openpyxl.comments import Comment
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from tools.excel_rows import discard_empty_cells, iter_stored_rows, stored_row
from tools.excel_batcher.excel_batcher import split_workbook, restore_batches
from tools.excel_merger.merge_active_sheets import merge_active_sheets
from tools.term_pair_checker.extract_terms_from_excel import load_history_tb_mapping

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}


def set_dimension(path, ref):
    with ZipFile(path) as archive:
        entries = [(item, archive.read(item)) for item in archive.infolist()]
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for item, data in entries:
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                root.find("s:dimension", NS).set("ref", ref)
                data = ET.tostring(root)
            archive.writestr(item, data)


def worksheet_xml(path):
    with ZipFile(path) as archive:
        return ET.fromstring(archive.read("xl/worksheets/sheet1.xml"))


class SparseReadTests(unittest.TestCase):
    def test_stale_dimensions_and_styles_do_not_hide_values_or_pad_rows(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            book = Workbook()
            try:
                sheet = book.active
                sheet["A1"] = "header"
                values = [("A2", 0), ("C2", False), ("D2", "=literal"),
                          ("E2", "=1+2"), ("F2", datetime(2026, 1, 2)),
                          ("G2", "#N/A"), ("H2", CellRichText([""])),
                          ("XFD100001", "far content")]
                for address, value in values:
                    sheet[address] = value
                sheet["D2"].data_type = "s"
                sheet["XFD1048576"].number_format = "@"
                book.save(path)
            finally:
                book.close()
            for ref in ("A1:XFD1048576", "A1:A1"):
                set_dimension(path, ref)
                for read_only in (False, True):
                    with self.subTest(ref=ref, read_only=read_only):
                        book = load_workbook(path, read_only=read_only)
                        try:
                            sheet = book.active
                            with patch.object(sheet, "iter_rows", side_effect=AssertionError("dense read")):
                                self.assertEqual(list(stored_row(sheet)), [1])
                                rows = dict(iter_stored_rows(sheet, min_row=2))
                            self.assertEqual(list(rows), [2, 100001])
                            self.assertEqual(rows[2][1].value, 0)
                            self.assertIs(rows[2][3].value, False)
                            self.assertEqual((rows[2][4].value, rows[2][4].data_type), ("=literal", "s"))
                            self.assertEqual(rows[2][5].data_type, "f")
                            self.assertEqual(rows[2][6].value, datetime(2026, 1, 2))
                            self.assertEqual(rows[2][7].data_type, "e")
                            self.assertEqual(rows[2][8].value, "")
                            self.assertEqual(rows[100001][16384].value, "far content")
                        finally:
                            book.close()
            # Early header reads close their ZIP member, so Windows can replace the file.
            path.rename(path.with_suffix(".renamed"))

    def test_discard_inert_cells_preserves_format_comments_links_and_merges(self):
        book = Workbook()
        try:
            sheet = book.active
            sheet["A1"] = "source"
            sheet["A2"] = 0
            sheet["B2"] = False
            sheet["D10"].number_format = "@"
            sheet["E11"].comment = Comment("keep", "QA")
            sheet["F12"].hyperlink = "https://example.com"
            sheet.merge_cells("G13:H13")
            sheet.cell(1048576, 16384)
            discard_empty_cells(sheet)
            self.assertNotIn((1048576, 16384), sheet._cells)
            self.assertEqual(sheet["D10"].number_format, "@")
            self.assertIsNotNone(sheet["E11"].comment)
            self.assertIsNotNone(sheet["F12"].hyperlink)
            self.assertIn("G13:H13", sheet.merged_cells)
            self.assertIs(sheet["B2"].value, False)
            self.assertEqual(sheet["A2"].value, 0)
        finally:
            book.close()

    def test_history_import_reads_past_large_gaps_without_using_dense_iteration(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "history.xlsx"
            book = Workbook()
            try:
                sheet = book.active
                sheet.append(["source", "target"])
                sheet.append(["Early", "First"])
                sheet["A100001"] = "Late"
                sheet["B100001"] = "Last"
                sheet["XFD1048576"].number_format = "@"
                book.save(path)
            finally:
                book.close()
            with patch.object(ReadOnlyWorksheet, "iter_rows", side_effect=AssertionError("dense read")):
                self.assertEqual(set(load_history_tb_mapping(path)), {"early", "late"})


class SparseBatchTests(unittest.TestCase):
    def test_split_restore_retains_positions_deletions_and_new_columns(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            book = Workbook()
            try:
                sheet = book.active
                sheet.append(["source", "target"])
                sheet.append(["Early", "remove me"])
                sheet["A100001"] = "Late"
                sheet["B100001"] = "old"
                sheet["XFD2"].number_format = "@"
                sheet["XFD1048576"].number_format = "@"
                book.save(source)
            finally:
                book.close()
            with patch.object(ReadOnlyWorksheet, "iter_rows", side_effect=AssertionError("dense read")):
                split = split_workbook(source, batch_size=50000)
            self.assertEqual(split.batch_count, 2)
            self.assertEqual(split.data_row_count, 100000)
            for index, batch in enumerate(split.batch_files):
                xml = worksheet_xml(batch)
                self.assertEqual(len(xml.findall("s:sheetData/s:row", NS)), 2)
                self.assertNotIn("XFD", xml.find("s:dimension", NS).get("ref"))
                book = load_workbook(batch)
                try:
                    sheet = book.active
                    if index == 0:
                        sheet["A2"] = None
                        sheet["B2"] = None
                    else:
                        sheet["B50001"] = "new"
                        sheet["D1"] = "review"
                        sheet["D50001"] = "pass"
                    book.save(batch)
                finally:
                    book.close()
                set_dimension(batch, "A1:XFD1048576")
            manifest = json.loads(split.manifest_path.read_text(encoding="utf-8"))
            manifest["original_max_column"] = 16384  # Legacy manifests can contain ghost widths.
            split.manifest_path.write_text(json.dumps(manifest), encoding="utf-8")
            restored_path = root / "restored.xlsx"
            restore_batches(split.output_dir, output_file=restored_path)
            xml = worksheet_xml(restored_path)
            self.assertEqual(len(xml.findall("s:sheetData/s:row", NS)), 4)
            book = load_workbook(restored_path)
            try:
                sheet = book.active
                self.assertEqual(sheet["A100001"].value, "Late")
                self.assertIsNone(sheet["A2"].value)
                self.assertIsNone(sheet["B2"].value)
                self.assertEqual(sheet["B100001"].value, "new")
                self.assertEqual(sheet["D100001"].value, "pass")
                self.assertEqual(sheet["D1"].value, "review")
                self.assertEqual(sheet["XFD1048576"].number_format, "@")
                self.assertLess(len(sheet._cells), 20)
            finally:
                book.close()

    def test_macro_batch_trimming_moves_stored_cells_without_dense_delete_rows(self):
        from tools.excel_batcher.excel_batcher import _trim_worksheet_to_batch
        book = Workbook()
        try:
            sheet = book.active
            sheet.append(["source", "target"])
            sheet["A100001"] = "Late"
            sheet.row_dimensions[100001].height = 32
            sheet["XFD1048576"].number_format = "@"
            with patch.object(sheet, "delete_rows", side_effect=AssertionError("dense delete")):
                _trim_worksheet_to_batch(sheet, header_rows=1,
                                         source_start_row=50002, source_end_row=100001)
            self.assertEqual(sheet["A50001"].value, "Late")
            self.assertEqual(sheet.row_dimensions[50001].height, 32)
            self.assertEqual(sheet.max_column, 2)
            self.assertLess(len(sheet._cells), 10)
        finally:
            book.close()


class GhostMergeTests(unittest.TestCase):
    def test_format_only_rows_and_columns_are_not_exported(self):
        with tempfile.TemporaryDirectory() as directory:
            folder = Path(directory) / "inputs"
            folder.mkdir()
            source = folder / "input.xlsx"
            book = Workbook()
            try:
                sheet = book.active
                sheet.append(["source", "target", "context"])
                sheet["A2"] = "Early"
                sheet["C2"] = "keep column position"
                sheet["XFD2"].number_format = "@"
                sheet["A100001"] = "Late"
                for row in range(3, 1003):
                    sheet.cell(row, 1).number_format = "@"
                sheet["XFD1048576"].number_format = "@"
                book.save(source)
            finally:
                book.close()
            output = Path(directory) / "merged.xlsx"
            merge_active_sheets(folder, output_path=output)
            book = load_workbook(output)
            try:
                self.assertEqual(book.active.max_row, 3)
                self.assertEqual(book.active.max_column, 4)
                self.assertEqual(book.active["B3"].value, "Late")
                self.assertEqual(book.active["D2"].value, "keep column position")
                self.assertIsNone(book.active["C2"].value)
            finally:
                book.close()


class SparseRevisionTests(unittest.TestCase):
    def test_revisions_in_far_report_rows_do_not_expand_or_shift_data(self):
        from tools.workflow.workflow_runner import run_workflow
        from tools.workflow.revision_applier import apply_workflow_revisions
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsx"
            report = Path(directory) / "report.xlsx"
            output = Path(directory) / "revised.xlsx"
            book = Workbook()
            try:
                book.active.title = "Data"
                book.active.append(["source", "target"])
                book.active["A2"] = "Early message"
                book.active["A100001"] = "Distant message"
                book.save(source)
            finally:
                book.close()
            run_workflow(input_file=source, output_file=report, source_column="A", target_column="B")
            book = load_workbook(report)
            try:
                review = book["问题处理"]
                review["D2"] = "First translation"
                review["D3"] = "Last translation"
                review.move_range("A3:F3", rows=100000)
                review["XFD1048576"].number_format = "@"
                before = set(review._cells)
                with patch("tools.workflow.revision_applier.load_workbook_for_editing", return_value=book):
                    result = apply_workflow_revisions(report, output)
                self.assertEqual(result.revised_count, 2)
                self.assertEqual(set(review._cells), before)
            finally:
                book.close()
            book = load_workbook(output)
            try:
                self.assertEqual(book["Data"]["B2"].value, "First translation")
                self.assertEqual(book["Data"]["B100001"].value, "Last translation")
                self.assertLess(len(book["Data"]._cells), 10)
            finally:
                book.close()
