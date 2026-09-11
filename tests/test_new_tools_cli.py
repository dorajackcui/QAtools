from __future__ import annotations

from contextlib import redirect_stderr, redirect_stdout
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from qatools.cli import main
from tools.excel_file_ops import BatchSummary, OperationResult


class NewToolsCliTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)

    def invoke(self, *arguments):
        output, errors = io.StringIO(), io.StringIO()
        with redirect_stdout(output), redirect_stderr(errors):
            code = main(list(map(str, arguments)))
        return code, output.getvalue(), errors.getvalue()

    def book(self, path, rows):
        path.parent.mkdir(parents=True, exist_ok=True)
        book = Workbook()
        try:
            book.active.title = "Data"
            for row in rows:
                book.active.append(row)
            book.save(path)
        finally:
            book.close()
        return path

    def values(self, path, row=2):
        book = load_workbook(path)
        try:
            return [cell.value for cell in book.active[row]]
        finally:
            book.close()

    def test_forward_multi_column_defaults_and_text_fidelity(self):
        master = self.book(self.root / "master.xlsx", [["id", "k", "s", "t", "t2"],
            [1, " k ", " S ", "  nan\n ", None], [2, "n", "N", 0, False]])
        target = self.book(self.root / "small" / "a.xlsx", [["k", "s", "t", "t2"],
            ["k", "S", "old", "keep"], ["n", "N", "old", "old"]])
        original = master.read_bytes()
        code, output, logs = self.invoke("content-sync", "master-to-target", master, target.parent,
                                        "--column-count", 2)
        self.assertEqual(code, 0, logs)
        self.assertEqual(self.values(target), ["k", "S", "  nan\n ", "keep"])
        self.assertEqual(self.values(target, 3), ["n", "N", "0", "False"])
        self.assertEqual(master.read_bytes(), original)
        self.assertIn("实际更新单元格: 3", output)
        self.assertIn("a.xlsx", logs)
        self.assertEqual(list(self.root.rglob("*.json")), [])

    def test_reverse_deterministic_order_and_copy_output(self):
        master = self.book(self.root / "master.xlsx", [["id", "k", "s", "t"], [1, "k", "S", "old"]])
        small = self.root / "small"
        self.book(small / "a.xlsx", [["k", "s", "t"], ["k", "S", "first"]])
        self.book(small / "z.xlsx", [["k", "s", "t"], ["k", "S", "None"], ["k", "S", None]])
        before = {p: p.read_bytes() for p in self.root.rglob("*.xlsx")}
        code, _, errors = self.invoke("content-sync", "target-to-master", master, small,
                                     "-o", self.root / "out", "--quiet")
        self.assertEqual(code, 0, errors)
        self.assertEqual(self.values(self.root / "out" / master.name)[3], "None")
        self.assertTrue(all(p.read_bytes() == value for p, value in before.items()))

    def test_sync_custom_mapping_and_flags_are_forwarded(self):
        from tools.content_sync.master_to_target import ColumnMapping
        result = BatchSummary("sync", self.root)
        with patch("tools.content_sync.master_to_target.sync_master_to_targets", return_value=result) as operation:
            code, _, _ = self.invoke("content-sync", "master-to-target", "m.xlsx", "small",
                "--master-key-column", "F", "--master-source-column", "G", "--master-content-column", "H",
                "--target-key-column", "a", "--target-source-column", "b", "--target-content-column", "e",
                "--master-sheet", "Master", "--target-sheet", "Translation", "--master-header-rows", 2,
                "--target-header-rows", 0, "--column-count", 3, "--fill-blank-only", "--allow-blank-write",
                "--compatibility-resave", "--quiet")
        self.assertEqual(code, 0)
        kwargs = operation.call_args.kwargs
        self.assertEqual(kwargs["master_columns"], ColumnMapping("F", "G", "H"))
        self.assertEqual(kwargs["target_columns"], ColumnMapping("A", "B", "E"))
        self.assertEqual((kwargs["master_sheet"], kwargs["target_sheet"]), ("Master", "Translation"))
        self.assertEqual((kwargs["master_header_rows"], kwargs["target_header_rows"], kwargs["column_count"]), (2, 0, 3))
        self.assertTrue(all(kwargs[key] for key in ("fill_blank_only", "allow_blank_write", "compatibility_resave")))
        self.assertIsNone(kwargs["output_dir"])
        self.assertIsNone(kwargs["log_callback"])

    def test_deep_replace_default_backup_is_byte_identical(self):
        source = self.book(self.root / "source" / "a.xlsx", [["new"]])
        target = self.book(self.root / "target" / "sub" / "a.xlsx", [["old"]])
        before = target.read_bytes()
        code, output, logs = self.invoke("deep-replace", source.parent, target.parent.parent, "--quiet")
        self.assertEqual(code, 0, logs)
        self.assertEqual(source.read_bytes(), target.read_bytes())
        backups = list((self.root / "target" / ".qatools-backups").rglob("a.xlsx"))
        self.assertEqual(len(backups), 1)
        self.assertEqual(backups[0].read_bytes(), before)
        self.assertIn("backup_dir:", output)

    def test_sync_worker_defaults_override_and_invalid_values(self):
        from tools.content_sync.cli import build_parser
        for direction, default in (("master-to-target", 2), ("target-to-master", 1)):
            parser = build_parser()
            args = [direction, "master.xlsx", "targets"]
            self.assertEqual(parser.parse_args(args).workers, default)
            self.assertEqual(parser.parse_args(args + ["--workers", "4"]).workers, 4)
            for value in ("0", "5", "1.5"):
                with redirect_stderr(io.StringIO()), self.assertRaises(SystemExit) as error:
                    parser.parse_args(args + ["--workers", value])
                self.assertEqual(error.exception.code, 2)
        result = BatchSummary("sync", self.root)
        with patch("tools.content_sync.target_to_master.sync_targets_to_master", return_value=result) as operation:
            code, _, _ = self.invoke("content-sync", "target-to-master", "m.xlsx", "small", "--workers", 2)
        self.assertEqual(code, 0)
        self.assertEqual(operation.call_args.kwargs["workers"], 2)

    def test_statistics_mode_and_reserved_output(self):
        target = self.book(self.root / "small" / "a.xlsx", [["id", "s", "t"],
                          [1, "don't re-use", "nan"], [2, "hello", "None"]])
        original = target.read_bytes()
        for _ in range(2):
            code, output, errors = self.invoke("untranslated-stats", target.parent, "--mode", "english_words", "--quiet")
            self.assertEqual(code, 0, errors)
            self.assertIn("excel_report:", output)
            self.assertEqual(self.values(target.parent / "未翻译统计.xlsx"), ["a.xlsx", 2, 1, 3, 2])
        self.assertEqual(target.read_bytes(), original)

    def test_column_actions_and_compatibility_forward_to_business(self):
        cases = [
            (["columns", "clear", "input", "--column", "AA", "--sheet", "Data", "--header-rows", "3"],
             "tools.column_tools.processor.operate_columns", {"action": "clear", "column": "AA", "sheet": "Data", "header_rows": 3}),
            (["columns", "insert", "input", "--inserted-header", "=literal"],
             "tools.column_tools.processor.operate_columns", {"action": "insert", "column": "C", "inserted_header": "=literal"}),
            (["columns", "delete", "input"], "tools.column_tools.processor.operate_columns", {"action": "delete", "column": "C"}),
            (["compatibility", "input", "-o", "output"], "tools.excel_compatibility.processor.resave_workbooks", {"output_dir": "output"}),
        ]
        for arguments, name, expected in cases:
            with self.subTest(arguments=arguments), patch(name, return_value=BatchSummary("native", self.root)) as operation:
                self.assertEqual(self.invoke(*arguments)[0], 0)
                self.assertEqual(operation.call_args.kwargs["folder_path"], "input")
                for key, value in expected.items():
                    self.assertEqual(operation.call_args.kwargs[key], value)

    def test_partial_failure_all_failure_and_master_save_failure_exit_codes(self):
        master = self.book(self.root / "master.xlsx", [["id", "k", "s", "t"], [1, "k", "S", "new"]])
        target = self.book(self.root / "small" / "good.xlsx", [["k", "s", "t"], ["k", "S", "old"]])
        (target.parent / "bad.xlsx").write_bytes(b"broken")
        self.assertEqual(self.invoke("content-sync", "master-to-target", master, target.parent)[0], 3)
        target.unlink()
        self.assertEqual(self.invoke("content-sync", "master-to-target", master, target.parent)[0], 1)
        self.book(target, [["k", "s", "t"], ["k", "S", "translated"]])
        before = master.read_bytes()
        with patch("tools.content_sync.master_to_target._save_output", side_effect=PermissionError("locked")):
            self.assertEqual(self.invoke("content-sync", "target-to-master", master, target.parent)[0], 1)
        self.assertEqual(master.read_bytes(), before)

    def test_skips_warnings_and_startup_errors_are_nonzero(self):
        for summary in (BatchSummary("x", self.root, [OperationResult("x", "skipped")]),
                        BatchSummary("x", self.root, warnings=["Excel quit failed"])):
            with patch("tools.excel_compatibility.processor.resave_workbooks", return_value=summary):
                self.assertEqual(self.invoke("compatibility", "input")[0], 3)
        with patch("tools.excel_compatibility.processor.resave_workbooks", side_effect=RuntimeError("Excel unavailable")):
            code, _, errors = self.invoke("compatibility", "input")
        self.assertEqual(code, 1)
        self.assertIn("Excel unavailable", errors)
        self.assertNotIn("Traceback", errors)

    def test_invalid_arguments_do_not_prompt_or_create_output(self):
        cases = [("content-sync",), ("columns", "clear"), ("compatibility",),
                 ("deep-replace", "source"), ("untranslated-stats",),
                 ("columns", "clear", "input", "--column", "XFE"),
                 ("content-sync", "target-to-master", "m", "s", "--column-count", "2"),
                 ("untranslated-stats", "input", "--header-rows", "-1")]
        with patch("builtins.input", side_effect=AssertionError("interactive prompt")):
            for args in cases:
                with self.subTest(args=args):
                    self.assertEqual(self.invoke(*args)[0], 2)
        self.assertEqual(list(self.root.iterdir()), [])

    def test_help_imports_no_optional_runtime_and_never_starts_excel(self):
        code = '''
import importlib.abc, sys
class Blocker(importlib.abc.MetaPathFinder):
    def find_spec(self, fullname, path=None, target=None):
        if fullname.split('.')[0] in {'PySide6', 'openpyxl', 'pythoncom', 'win32com', 'pandas'}:
            raise AssertionError('Eager runtime import: ' + fullname)
sys.meta_path.insert(0, Blocker())
from qatools.cli import main
commands = [['content-sync'], ['content-sync', 'master-to-target'], ['content-sync', 'target-to-master'],
            ['columns'], ['columns', 'clear'], ['columns', 'insert'], ['columns', 'delete'],
            ['compatibility'], ['deep-replace'], ['untranslated-stats']]
for command in commands:
    assert main(command + ['--help']) == 0
'''
        result = subprocess.run([sys.executable, "-c", code], cwd=Path(__file__).resolve().parents[1],
                                capture_output=True, text=True)
        self.assertEqual(result.returncode, 0, result.stderr)


if __name__ == "__main__":
    unittest.main()
