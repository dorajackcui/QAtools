from __future__ import annotations

import io
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path

from openpyxl import Workbook, load_workbook

from tools.workflow.cli import main
from tools.workflow.revision_applier import apply_workflow_revisions
from tools.workflow.workflow_runner import run_workflow, write_empty_target_problems


OPTIONAL_CHECKS = (
    "term_pair", "tag", "line_break", "source_consistency", "target_consistency",
    "substring_consistency", "number", "url", "chinese_target", "target_text",
)


class WorkflowEmptyTargetTests(unittest.TestCase):
    def test_empty_check_keeps_sparse_input_sparse_and_skips_formatted_tail(self):
        book = Workbook()
        try:
            worksheet = book.active
            worksheet["A3"] = "source"
            worksheet["A1000000"].number_format = "@"
            cells_before = set(worksheet._cells)
            name, count = write_empty_target_problems(book, worksheet, "A", "B", 2)
            self.assertEqual(count, 1)
            self.assertEqual(set(worksheet._cells), cells_before)
            self.assertEqual(book[name]["A2"].value, 3)
        finally:
            book.close()

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.input = Path(temporary.name) / "input.xlsx"

    def fixture(self, rows, *, data_name="Data"):
        book = Workbook()
        try:
            book.active.title = data_name
            book.active.append(["target", "note", "source"])
            for source, target in rows:
                book.active.append([target, "keep", source])
            book.save(self.input)
        finally:
            book.close()

    def run_checks(self, **options):
        flags = {f"run_{name}_check": False for name in OPTIONAL_CHECKS}
        flags.update(options)
        return run_workflow(input_file=self.input, source_column="C", target_column="A", **flags)

    def test_blank_values_scope_and_existing_sheet_are_preserved(self):
        self.fixture([
            ("before start", None), ("empty cell", None), ("empty string", ""),
            ("whitespace", " \t\n\u00a0\u3000"), (None, None), (" \t", None),
            ("zero target", 0), ("false target", False), ("literal nan", "nan"),
            (0, None), (False, None), ("nan", None), ("formula", '=IF(1=1,"","x")'),
        ], data_name="Target 为空")
        book = load_workbook(self.input)
        try:
            book.create_sheet("Other").append([None, None, "must not inspect"])
            book.active = 0
            book.save(self.input)
        finally:
            book.close()
        original = self.input.read_bytes()
        result = self.run_checks(start_row=3)
        book = load_workbook(result.output_path)
        try:
            rows = list(book["问题处理"].values)[1:]
            self.assertEqual([row[0] for row in rows], [3, 4, 5, 11, 12, 13])
            self.assertTrue(all(row[5] == "Target 为空" for row in rows))
            self.assertEqual(rows[2][2], " \t\n\u00a0\u3000")
            self.assertEqual(rows[0][4], "【Target 为空】译文为空或仅含空白")
            self.assertEqual(book["问题处理"]["A2"].hyperlink.location, "'Target 为空'!A3")
            self.assertEqual(list(book["质量检查汇总"].values),
                             [("检查项", "问题行数"), ("Target 为空", 6)])
            self.assertEqual(book["Target 为空"]["C3"].value, "empty cell")
            self.assertNotIn("Target 为空1", book.sheetnames)
        finally:
            book.close()
        self.assertEqual(self.input.read_bytes(), original)

    def test_merges_with_other_issues_and_revisions_fill_empty_targets(self):
        self.fixture([("Same {name}", None), ("Same {name}", "Translated {name}")])
        result = self.run_checks(run_source_consistency_check=True, run_tag_check=True)
        book = load_workbook(result.output_path)
        try:
            rows = list(book["问题处理"].values)[1:]
            self.assertEqual([row[0] for row in rows], [2, 3])
            self.assertEqual(set(rows[0][5].split("；")),
                             {"Target 为空", "同 Source 不同 Target", "Tag 检查"})
            book["问题处理"]["D2"] = "New {name}"
            book.save(result.output_path)
        finally:
            book.close()
        revision = apply_workflow_revisions(result.output_path)
        self.assertEqual(revision.revised_count, 1)
        book = load_workbook(revision.output_path)
        try:
            self.assertEqual(book.sheetnames, ["Data"])
            self.assertEqual(book["Data"]["A2"].value, "New {name}")
            self.assertEqual(book["Data"]["B2"].value, "keep")
        finally:
            book.close()

    def test_explicit_cli_check_still_reports_empty_target(self):
        self.fixture([("plain source", None)])
        output = self.input.with_name("report.xlsx")
        with redirect_stdout(io.StringIO()):
            code = main([str(self.input), "-c", "C", "-t", "A", "--check", "url", "-o", str(output)])
        self.assertEqual(code, 0)
        book = load_workbook(output)
        try:
            self.assertEqual(book["问题处理"]["F2"].value, "Target 为空")
            self.assertEqual(list(book["质量检查汇总"].values),
                             [("检查项", "问题行数"), ("Target 为空", 1), ("URL 一致性", 0)])
        finally:
            book.close()


if __name__ == "__main__":
    unittest.main()
