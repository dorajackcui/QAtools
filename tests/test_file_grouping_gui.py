import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

from openpyxl import Workbook
from PySide6.QtCore import QCoreApplication, QEvent
from PySide6.QtGui import QCloseEvent
from PySide6.QtTest import QTest
from PySide6.QtWidgets import QApplication, QMessageBox

from toolshub_gui import ToolshubApp
from tools.header_aliases import HeaderAliasStore


class FileGroupingGuiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "source"
        self.source.mkdir()
        self.errors = self.enterContext(patch.object(QMessageBox, "critical"))
        self.callback_errors = self.enterContext(patch("sys.excepthook"))
        self.window = ToolshubApp(show_window=False, initial_tool="file_grouping",
                                  header_alias_store=HeaderAliasStore(self.root / "aliases.json"))
        self.page = self.window.tool_frames["file_grouping"]
        self.page.source_picker.set_path(str(self.source))
        self.addCleanup(self.dispose)

    def dispose(self):
        self.page.details_dialog.reject()
        self.window.close()
        self.window.deleteLater()
        QCoreApplication.sendPostedEvents(None, QEvent.Type.DeferredDelete)

    def book(self, relative, rows, headers=("MsgStr", "fr")):
        path = self.source / relative
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        try:
            workbook.active.append(headers)
            for row in rows:
                workbook.active.append(row)
            workbook.save(path)
        finally:
            workbook.close()
        return path

    def click(self, button):
        self.assertTrue(button.isEnabled())
        button.click()
        deadline = time.monotonic() + 15
        while self.page.has_running_tasks() and time.monotonic() < deadline:
            QTest.qWait(10)
        self.assertFalse(self.page.has_running_tasks())
        self.app.processEvents()
        self.callback_errors.assert_not_called()

    def test_real_scan_details_navigation_and_whole_file_export(self):
        self.book("one/a.xlsx", [("共同", None)])
        self.book("two/a.xlsx", [("共同", None)])
        self.book("b.xlsx", [("独立", None)])
        self.click(self.page.scan_button)
        self.assertEqual(self.page.counts.text(), "3 个文件，2 个分组")
        self.assertEqual(self.page.model.index(0, 2).data(), "2")
        self.page.table.clicked.emit(self.page.model.index(0, 0))
        self.app.processEvents()
        self.assertTrue(self.page.details_dialog.isVisible())
        self.assertEqual(self.page.details_text.toPlainText(), "one/a.xlsx\ntwo/a.xlsx")
        self.page.details_dialog.reject()
        self.window.select_tool("file_collector")
        self.window.select_tool("file_grouping")
        self.assertIs(self.window.current_tool_frame, self.page)
        with patch("tools.file_grouping.qt_page.QFileDialog.getExistingDirectory", return_value=str(self.root)):
            self.click(self.page.export_button)
        output = self.page.exported.output_dir
        self.assertEqual((output / "文件群1/one/a.xlsx").read_bytes(), (self.source / "one/a.xlsx").read_bytes())
        self.assertEqual((output / "其他/b.xlsx").read_bytes(), (self.source / "b.xlsx").read_bytes())
        self.assertIn("已导出 3 个文件", self.page.status.text())
        self.errors.assert_not_called()

    def test_every_input_invalidates_old_scan_and_export(self):
        self.book("a.xlsx", [("原文", None)])
        for field, value in ((self.page.source_header, "Other"), (self.page.target_header, "translation"),
                             (self.page.source_picker.line_edit, str(self.root))):
            self.page.source_picker.set_path(str(self.source))
            self.page.source_header.setText("MsgStr")
            self.page.target_header.setText("fr")
            self.click(self.page.scan_button)
            self.assertTrue(self.page.export_button.isEnabled())
            field.setText(value)
            self.assertIsNone(self.page.result)
            self.assertEqual(self.page.model.rowCount(), 0)
            self.assertFalse(self.page.export_button.isEnabled())
        self.errors.assert_not_called()

    def test_scan_issues_block_export_and_can_be_viewed_and_corrected(self):
        path = self.book("bad.xlsx", [("原文", None)], headers=("wrong", "fr"))
        self.click(self.page.scan_button)
        self.assertFalse(self.page.export_button.isEnabled())
        self.assertIn("扫描不完整", self.page.status.text())
        self.page.status.linkActivated.emit("issues")
        self.app.processEvents()
        self.assertIn("bad.xlsx", self.page.details_text.toPlainText())
        self.page.details_dialog.reject()
        path.unlink()
        self.book("good.xlsx", [("原文", None)])
        self.click(self.page.scan_button)
        self.assertTrue(self.page.export_button.isEnabled())
        self.errors.assert_not_called()

    def test_changed_sources_require_rescan_and_cancelled_export_keeps_result(self):
        self.book("a.xlsx", [("原文", None)])
        self.click(self.page.scan_button)
        result = self.page.result
        with patch("tools.file_grouping.qt_page.QFileDialog.getExistingDirectory", return_value=""):
            self.click(self.page.export_button)
        self.assertIs(self.page.result, result)
        self.book("a.xlsx", [("已变", None)])
        with patch("tools.file_grouping.qt_page.QFileDialog.getExistingDirectory", return_value=str(self.root)):
            self.click(self.page.export_button)
        self.errors.assert_called_once()
        self.assertIn("重新扫描", self.page.status.text())
        self.assertFalse(list(self.root.glob("source_文件分组_*")))
        self.errors.reset_mock()
        self.click(self.page.scan_button)
        self.assertTrue(self.page.export_button.isEnabled())
        self.errors.assert_not_called()

    def test_hidden_background_task_blocks_window_exit(self):
        from unittest.mock import Mock
        self.page._workers.add(Mock())
        try:
            self.window.select_tool("file_collector")
            event = QCloseEvent()
            with patch("toolshub_gui.show_warning") as warning:
                self.window.closeEvent(event)
            self.assertFalse(event.isAccepted())
            self.assertIn("文件分组", warning.call_args.args[2])
        finally:
            self.page._workers.clear()

    def test_merged_input_mode_scans_and_exports_all_rows_per_group(self):
        path = self.book("merged.xlsx", [["a.xlsx", "共同", None, "pending"],
                        ["b.xlsx", "共同", None, "pending"], ["a.xlsx", "已译", "done", "keep"],
                        ["c.xlsx", "独立", None, "other"]], headers=("File", "MsgStr", "fr", "Note"))
        self.page.input_kind.setCurrentIndex(1)
        self.assertEqual(self.page.source_picker.path(), "")
        self.assertFalse(self.page.filename_column.isHidden())
        self.page.source_picker.set_path(str(path))
        self.click(self.page.scan_button)
        self.assertEqual(self.page.counts.text(), "3 个文件，2 个分组")
        self.assertTrue(self.page.export_button.isEnabled())
        with patch("tools.file_grouping.qt_page.QFileDialog.getExistingDirectory", return_value=str(self.root)):
            self.click(self.page.export_button)
        output = self.page.exported.output_dir
        from openpyxl import load_workbook
        book = load_workbook(output / "文件群1.xlsx", read_only=True)
        try:
            rows = list(book.active.iter_rows(values_only=True))
            self.assertEqual(len(rows), 4)
            self.assertEqual(rows[-1][2], "done")
        finally:
            book.close()
        self.assertIn("已导出 2 个工作簿", self.page.status.text())
        self.page.filename_column.setText("B")
        self.assertIsNone(self.page.result)
        self.assertFalse(self.page.export_button.isEnabled())
        self.page.input_kind.setCurrentIndex(0)
        self.assertTrue(self.page.filename_column.isHidden())
        self.assertEqual(self.page.source_picker.path(), "")
        self.errors.assert_not_called()

    def test_merged_mode_uses_excel_picker_and_opens_preview_on_selection(self):
        self.page.input_kind.setCurrentIndex(1)
        with patch("tools.file_grouping.qt_page.QFileDialog.getOpenFileName", return_value=("merged.xlsx", "")):
            self.page.choose_source()
        self.assertEqual(self.page.source_picker.path(), "merged.xlsx")
        self.assertIsNone(self.page.result)
        deadline = time.monotonic() + 5
        while self.page.has_running_tasks() and time.monotonic() < deadline:
            QTest.qWait(10)
        self.assertFalse(self.page.has_running_tasks())
        self.assertTrue(self.page.input_preview.dialog.isVisible())


if __name__ == "__main__":
    unittest.main()
