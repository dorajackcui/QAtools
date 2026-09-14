from __future__ import annotations

import importlib
import itertools
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook

from tools.excel_output import PROBLEM_BASE_HEADERS, validate_report_output_path
from tools.report_text import CHECK_DESCRIPTION_LIMIT, format_consistency_variants
from tools.term_pair_checker.extract_terms_from_excel import process_excel as check_terms
from tools.workflow.review_sheet import collect_review_rows, WORKFLOW_METADATA_SHEET_NAME
from tools.workflow.revision_applier import apply_workflow_revisions
from tools.workflow.workflow_runner import run_workflow


CHECK_FLAGS = (
    "run_term_pair_check", "run_tag_check", "run_line_break_check",
    "run_source_consistency_check", "run_target_consistency_check",
    "run_substring_consistency_check", "run_number_check", "run_url_check",
    "run_chinese_target_check", "run_target_text_check",
)


class QaReportIntegrityTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.input_path = self.root / "input.xlsx"

    def fixture(self, rows, *, headers=("source", "target"), before=()):
        book = Workbook()
        try:
            book.active.title = "Data"
            book.active.append(headers)
            for row in rows:
                book.active.append(row)
            for name in before:
                book.create_sheet(name, 0)["A1"] = "Keep existing content"
            book.active = book.sheetnames.index("Data")
            book.save(self.input_path)
        finally:
            book.close()

    def run_checks(self, *enabled, **kwargs):
        return run_workflow(
            input_file=self.input_path,
            source_column=kwargs.pop("source_column", "A"), target_column="B",
            **{flag: flag in enabled for flag in CHECK_FLAGS}, **kwargs,
        )

    def test_term_checks_preserve_existing_columns_and_revision_source_coordinates(self):
        for helper_header in ("术语QA问题", "User annotation"):
            with self.subTest(header=helper_header):
                self.fixture(
                    [("=1+2", "【Cat】", "Keep note", "【猫咪】"),
                     ("=1+3", "【Dog】", "Keep note 2", "【猫咪】")],
                    headers=("Formula", "target", helper_header, "source"),
                )
                original_bytes = self.input_path.read_bytes()
                with ZipFile(self.input_path) as archive:
                    original_sheet = archive.read("xl/worksheets/sheet1.xml")
                standalone = check_terms(self.input_path, "D", "B")[3]
                workflow = self.run_checks("run_term_pair_check", source_column="D")
                for output in (standalone, workflow.output_path):
                    with ZipFile(output) as archive:
                        self.assertEqual(archive.read("xl/worksheets/sheet1.xml"), original_sheet)
                self.assertEqual(self.input_path.read_bytes(), original_bytes)
                book = load_workbook(workflow.output_path)
                try:
                    book["问题处理"]["D2"] = "【Cat】"
                    # Even an older report's obsolete cleanup flag must not delete columns.
                    book[WORKFLOW_METADATA_SHEET_NAME].append(["remove_term_helper", "1"])
                    book.save(workflow.output_path)
                finally:
                    book.close()
                revision = apply_workflow_revisions(workflow.output_path)
                self.assertEqual(revision.conflict_rows, ())
                self.assertEqual(revision.revised_count, 1)
                book = load_workbook(revision.output_path)
                try:
                    self.assertEqual(book["Data"]["B3"].value, "【Cat】")
                    self.assertEqual(book["Data"]["C3"].value, "Keep note 2")
                    self.assertEqual(book["Data"]["D3"].value, "【猫咪】")
                    self.assertEqual(book["Data"]["A3"].value, "=1+3")
                finally:
                    book.close()

    def test_disabled_tag_preserves_existing_summary_even_when_it_is_the_data_sheet(self):
        for use_summary_as_data in (False, True):
            with self.subTest(summary_is_data=use_summary_as_data):
                self.fixture([("text", "中文")], before=() if use_summary_as_data else ("检查汇总",))
                if use_summary_as_data:
                    book = load_workbook(self.input_path)
                    try:
                        book["Data"].title = "检查汇总"
                        book.save(self.input_path)
                    finally:
                        book.close()
                result = self.run_checks("run_chinese_target_check")
                book = load_workbook(result.output_path)
                try:
                    self.assertIn("检查汇总", book.sheetnames)
                    if use_summary_as_data:
                        self.assertEqual(book["检查汇总"]["B2"].value, "中文")
                    else:
                        self.assertEqual(book["检查汇总"]["A1"].value, "Keep existing content")
                finally:
                    book.close()
                revision = apply_workflow_revisions(result.output_path)
                self.assertTrue(revision.output_path.exists())

    def test_rebuilding_previous_results_does_not_switch_the_input_sheet(self):
        for before in (("检查汇总",), ("术语汇总", "检查汇总")):
            with self.subTest(before=before):
                self.fixture([("{a}", "{b}"), ("A\nB", "AB")], before=before)
                result = self.run_checks("run_term_pair_check", "run_tag_check", "run_line_break_check")
                self.assertEqual(result.worksheet_title, "Data")
                self.assertEqual(result.line_break_problem_count, 1)
                book = load_workbook(result.output_path)
                try:
                    review = book["问题处理"]
                    self.assertEqual(review["A3"].value, 3)
                    self.assertIn("换行数量检查", review["F3"].value)
                    review["D3"] = "A\nB"
                    book.save(result.output_path)
                finally:
                    book.close()
                self.assertEqual(apply_workflow_revisions(result.output_path).revised_count, 1)

    def test_all_file_checkers_reject_overwriting_input_before_writing(self):
        self.fixture([("{a}", "中文")])
        original = self.input_path.read_bytes()
        modules = (
            "term_pair_checker.extract_terms_from_excel",
            "tag_placeholder_checker.check_tags_and_placeholders",
            "line_break_checker.check_line_breaks",
            "source_consistency_checker.check_source_consistency",
            "target_consistency_checker.check_target_consistency",
            "content_fidelity_checker.check_content_fidelity",
            "chinese_target_checker.check_chinese_target",
            "target_text_checker.check_target_text",
        )
        processors = [run_workflow] + [
            importlib.import_module(f"tools.{module}").process_excel for module in modules
        ]
        for processor in processors:
            with self.subTest(processor=processor.__module__):
                with self.assertRaisesRegex(ValueError, "输出路径不能与输入文件相同"):
                    processor(input_file=self.input_path, source_column="A", target_column="B",
                              output_file=self.root / "." / "input.xlsx")
                self.assertEqual(self.input_path.read_bytes(), original)

    def test_output_hardlink_to_input_is_also_rejected(self):
        self.fixture([("text", "中文")])
        alias = self.root / "alias.xlsx"
        try:
            alias.hardlink_to(self.input_path)
        except OSError as error:
            self.skipTest(f"Hard links unavailable: {error}")
        with self.assertRaisesRegex(ValueError, "输出路径不能与输入文件相同"):
            validate_report_output_path(self.input_path, alias)

    def test_long_variants_leave_room_for_tag_details_and_keep_complete_snapshots(self):
        for reverse in (False, True):
            with self.subTest(reverse=reverse):
                rows = [("Same {a}", "a" * 18000), ("Same {a}", "b" * 18000)]
                if reverse:
                    rows = [(target, source) for source, target in rows]
                self.fixture(rows)
                flag = "run_target_consistency_check" if reverse else "run_source_consistency_check"
                result = self.run_checks(flag, "run_tag_check")
                book = load_workbook(result.output_path)
                try:
                    review_rows = list(book["问题处理"].values)[1:]
                    for index, row in enumerate(review_rows):
                        self.assertEqual(row[1:3], rows[index])
                        self.assertIn("已截断", row[4])
                        self.assertIn("【Tag 检查】", row[4])
                        self.assertLess(len(row[4]), 1000)
                finally:
                    book.close()

    def test_many_variants_keep_full_counts_and_reference_rows(self):
        for reverse in (False, True):
            with self.subTest(reverse=reverse):
                rows = [("Same", f"Variant {index}") for index in range(15)]
                if reverse:
                    rows = [(target, source) for source, target in rows]
                self.fixture(rows)
                module = "target_consistency_checker.check_target_consistency" if reverse else "source_consistency_checker.check_source_consistency"
                result = importlib.import_module(f"tools.{module}").process_excel(self.input_path, "A", "B")
                book = load_workbook(result.output_path)
                try:
                    report = book.worksheets[-1]
                    self.assertEqual(report.max_row, 16)
                    self.assertEqual(report["E2"].value, 15)
                    self.assertEqual(report["F2"].value, "、".join(map(str, range(2, 17))))
                    self.assertIn("另 5 种未展示", report["D2"].value)
                    self.assertNotIn("Variant 10", report["D2"].value)
                finally:
                    book.close()

    def test_description_budget_is_per_check_even_with_many_issues_in_one_check(self):
        book = Workbook()
        try:
            checks = []
            for index in range(10):
                sheet = book.create_sheet(f"check{index}")
                sheet.append(PROBLEM_BASE_HEADERS)
                for issue in range(4):
                    sheet.append([2, "source", "target", str(issue) + "x" * 20000])
                checks.append((f"Check {index}", sheet.title))
            row = collect_review_rows(book, checks)[0]
            self.assertLessEqual(len(row[4]), 10 * CHECK_DESCRIPTION_LIMIT + 9)
            for label, _ in checks:
                self.assertIn(f"【{label}】", row[4])
            self.assertEqual(row[4].count("已截断"), 10)
        finally:
            book.close()

    def test_variant_preview_limits_can_be_adjusted_without_changing_count(self):
        text = format_consistency_variants(["x" * 200, "y" * 200, "z"], "译法",
                                           max_variant_chars=20, max_variants=2)
        self.assertTrue(text.startswith("3 种译法："))
        self.assertEqual(text.count("已截断"), 2)
        self.assertIn("另 1 种未展示", text)
        self.assertNotIn("x" * 20, text)

    def test_all_pairs_and_all_checks_merge_the_same_rows_and_counts_as_single_checks(self):
        self.fixture([
            ("【阿童木】", "【Astro Boy】"), ("【阿童木】", "【Atom】"),
            ("Value {n}\n100 https://example.com/a", "Value {m}  200 https://example.com/b 中文.. "),
            ("保存更改", "Save changes"), ("是否保存更改？", "Save modifications?"),
            ("Alpha", "Shared"), ("Beta", "Shared"), ("Same", "First"), ("Same", "Second"),
        ])

        def rows_and_counts(enabled):
            result = self.run_checks(*enabled)
            book = load_workbook(result.output_path)
            try:
                rows = {row[0]: set(row[5].split("；")) for row in list(book["问题处理"].values)[1:]}
                for name, count in list(book["质量检查汇总"].values)[1:]:
                    self.assertEqual(count, sum(name in checks for checks in rows.values()))
                return rows
            finally:
                book.close()

        singles = {flag: rows_and_counts([flag]) for flag in CHECK_FLAGS}
        for enabled in itertools.chain(itertools.combinations(CHECK_FLAGS, 2), [CHECK_FLAGS]):
            with self.subTest(enabled=enabled):
                expected = {}
                for flag in enabled:
                    for row, checks in singles[flag].items():
                        expected.setdefault(row, set()).update(checks)
                self.assertEqual(rows_and_counts(enabled), expected)


if __name__ == "__main__":
    unittest.main()
