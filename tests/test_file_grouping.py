from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools.file_grouping.grouping import export_file_groups, scan_file_groups


class FileGroupingTests(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.source = self.root / "source"
        self.source.mkdir()
        self.output = self.root / "output"

    def book(self, name, rows, *, headers=("MsgStr", "fr"), extra_sheets=(), write_only=False):
        path = self.source / name
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook(write_only=write_only)
        try:
            worksheet = workbook.create_sheet("Sheet1") if write_only else workbook.active
            worksheet.append(headers)
            for row in rows:
                worksheet.append(row)
            for title, other_rows in extra_sheets:
                sheet = workbook.create_sheet(title)
                for row in other_rows:
                    sheet.append(row)
            workbook.save(path)
        finally:
            workbook.close()
        return path

    def scan(self, **kwargs):
        return scan_file_groups(self.source, **kwargs)

    def test_transitive_groups_independent_files_and_deduplicated_counts(self):
        self.book("nested/a.xlsx", [("共享一", None), ("共享一", " "), ("独有", None)])
        self.book("b.xlsx", [("共享一", None), ("共享二", None)])
        self.book("c.xlsx", [("共享二", None)])
        self.book("d.xlsx", [("共享一", "translated"), ("独立", None)])
        self.book("e.xlsx", [("abc", None)])
        result = self.scan()
        self.assertTrue(result.can_export)
        self.assertEqual([(group.name, len(group.files), group.unique_chars) for group in result.groups],
                         [("文件群 1", 3, 8), ("其他", 2, 2)])
        self.assertEqual(set(result.groups[0].files), {Path("nested/a.xlsx"), Path("b.xlsx"), Path("c.xlsx")})
        self.assertEqual(sum(file.pending_rows for file in result.files), 8)
        # Check the central guarantee independently of the connected-component implementation.
        sources = {file.path: file.sources for file in result.files}
        collections = [set().union(*(sources[path] for path in group.files)) for group in result.groups]
        self.assertFalse(collections[0] & collections[1])
        self.assertFalse(sources[Path("d.xlsx")] & sources[Path("e.xlsx")])

    def test_exact_source_comparison_and_strict_empty_targets(self):
        self.book("a.xlsx", [("你好", None), ("trim", " "), ("", None), (" ", None),
                              (0, None), (False, None), ("zero", 0), ("bool", False),
                              ("nan text", "nan"), ("formula target", "=A2")])
        self.book("b.xlsx", [(" 你好", None), ("你好 ", None), ("TRIM", None),
                              ("zero", None), ("bool", None), ("nan text", None), ("formula target", None)])
        result = self.scan()
        self.assertEqual(len(result.groups), 1)
        self.assertEqual(result.groups[0].name, "其他")
        self.assertEqual(result.files[0].pending_rows, 4)
        self.assertEqual(result.files[0].sources, {"你好", "trim", "0", "False"})
        self.assertEqual(result.groups[0].unique_chars, 6)

    def test_every_matching_sheet_contributes_and_unrelated_sheets_are_ignored(self):
        self.book("a.xlsx", [("唯一", None)], headers=(" text ", "translation"), extra_sheets=(
            ("extra", [["text", "translation"], ["跨表", None]]),
            ("readme", [["说明"], ["not translation data"]]),))
        self.book("b.xlsm", [("跨表", None)], headers=("text", "translation"))
        result = self.scan(source_header="text", target_header="translation")
        self.assertEqual(len(result.groups[0].files), 2)
        self.assertEqual(result.groups[0].unique_chars, 4)
        self.assertFalse(result.issues)

    def test_dimensionless_sparse_rows_are_scanned_and_exported(self):
        for index, headers in enumerate((
            ("MsgStr", "fr"),
            ("Ignored", "MsgStr", "Note", "fr", "Trailing"),
            ("fr", "Note", "MsgStr", "Trailing"),
        )):
            with self.subTest(headers=headers):
                def row(source, target):
                    values = [None] * len(headers)
                    values[headers.index("MsgStr")] = source
                    values[headers.index("fr")] = target
                    if "Note" in headers:
                        values[headers.index("Note")] = "context"
                    return values

                rows = [row("共同原文", None), row("共同原文", " \t"),
                        row("已翻译", "translated"), row(None, "译文"),
                        row(None, None), row(" ", None), row(0, None),
                        row(False, None), row("公式译文", "=A2")]
                paths = [self.book(f"case_{index}/a.xlsx", rows, headers=headers, write_only=True),
                         self.book(f"case_{index}/b.xlsx", [row("共同原文", None)], headers=headers, write_only=True)]
                original = {path.name: path.read_bytes() for path in paths}
                workbook = load_workbook(paths[0], read_only=True)
                try:
                    self.assertIsNone(workbook.active.max_row)
                    self.assertIsNone(workbook.active.max_column)
                finally:
                    workbook.close()
                result = scan_file_groups(paths[0].parent)
                self.assertFalse(result.issues)
                self.assertTrue(result.can_export)
                self.assertEqual([(group.name, len(group.files), group.unique_chars) for group in result.groups],
                                 [("文件群 1", 2, 4)])
                self.assertEqual([file.pending_rows for file in result.files], [4, 1])
                self.assertEqual(result.files[0].sources, {"共同原文", "0", "False"})
                self.assertEqual(result.files[1].sources, {"共同原文"})
                output = self.root / f"export_{index}"
                summary = export_file_groups(result, output)
                self.assertEqual(summary.file_count, 2)
                for path in paths:
                    self.assertEqual(path.read_bytes(), original[path.name])
                    self.assertEqual((output / "文件群1" / path.name).read_bytes(), original[path.name])

    def test_dimensionless_pending_formulas_still_report_the_original_row(self):
        self.book("formula.xlsx", [[None, "=B3", None, None]],
                  headers=("Ignored", "MsgStr", "Note", "fr"), write_only=True)
        result = self.scan()
        self.assertFalse(result.can_export)
        self.assertEqual(len(result.issues), 1)
        self.assertIn("第 2 行", result.issues[0].message)
        self.assertIn("公式", result.issues[0].message)

    def test_missing_duplicate_headers_and_pending_source_formulas_block_export(self):
        self.book("good.xlsx", [("文本", None)])
        self.book("missing.xlsx", [("a", None)], headers=("MsgStr", "wrong"))
        self.book("duplicate.xlsx", [("a", None)], headers=("MsgStr", "MsgStr", "fr"))
        self.book("formula.xlsx", [("=A3", None)])
        result = self.scan()
        self.assertEqual(len(result.files), 1)
        self.assertEqual(len(result.issues), 3)
        self.assertFalse(result.can_export)
        with self.assertRaisesRegex(ValueError, "扫描不完整"):
            export_file_groups(result, self.output)
        self.assertFalse(self.output.exists())

    def test_corrupt_unsupported_and_no_pending_files_are_not_misreported(self):
        self.book("translated.xlsx", [("你好", "Bonjour")])
        (self.source / "broken.xlsx").write_bytes(b"broken")
        (self.source / "legacy.xls").write_bytes(b"legacy")
        result = self.scan()
        self.assertEqual(result.files[0].pending_rows, 0)
        self.assertEqual(result.groups[0].unique_chars, 0)
        self.assertEqual(len(result.issues), 2)
        self.assertFalse(result.can_export)

    def test_lock_files_backup_trees_and_generated_directories_are_excluded(self):
        self.book("good.xlsx", [("你好", None)])
        self.book("~$lock.xlsx", [("你好", None)])
        self.book(".qatools-backups/backup.xlsx", [("你好", None)])
        result = self.scan()
        export_file_groups(result, self.output)
        # A generated subtree is ignored when its parent is scanned in this session.
        root_result = scan_file_groups(self.root)
        self.assertEqual([file.path for file in root_result.files], [Path("source/good.xlsx")])

    def test_export_preserves_whole_file_bytes_relative_paths_and_source_files(self):
        a = self.book("one/same.xlsx", [("重复", None)])
        b = self.book("two/same.xlsx", [("重复", None), ("不同", None)])
        c = self.book("alone.xlsm", [("独立", None)])
        before = {path: (path.read_bytes(), path.stat().st_mtime_ns) for path in (a, b, c)}
        result = self.scan()
        progress = []
        summary = export_file_groups(result, self.output, progress_callback=lambda *item: progress.append(item))
        self.assertEqual(summary.file_count, 3)
        self.assertEqual(progress[-1], (3, 3))
        for group in result.groups:
            for relative in group.files:
                source = self.source / relative
                copied = self.output / group.name.replace(" ", "") / relative
                self.assertEqual(copied.read_bytes(), before[source][0])
                self.assertEqual(source.read_bytes(), before[source][0])
                self.assertEqual(source.stat().st_mtime_ns, before[source][1])
                self.assertEqual(copied.stat().st_mtime_ns, before[source][1])
        self.assertEqual(len(list(self.output.rglob("*.*"))), 3)

    def test_changes_new_files_and_deleted_files_require_new_scan_before_any_output(self):
        path = self.book("a.xlsx", [("文本", None)])
        result = self.scan()
        self.book("b.xlsx", [("文本", None)])
        with self.assertRaisesRegex(ValueError, "重新扫描"):
            export_file_groups(result, self.output)
        self.assertFalse(self.output.exists())
        result = self.scan()
        path.unlink()
        with self.assertRaisesRegex(ValueError, "重新扫描"):
            export_file_groups(result, self.output)
        self.assertFalse(self.output.exists())
        result = self.scan()
        self.book("b.xlsx", [("改后", None)])
        with self.assertRaisesRegex(ValueError, "重新扫描"):
            export_file_groups(result, self.output)
        self.assertFalse(self.output.exists())

    def test_content_hash_is_checked_even_if_metadata_appears_unchanged(self):
        path = self.book("a.xlsx", [("文本", None)])
        result = self.scan()
        path.write_bytes(b"different content")
        with patch("tools.file_grouping.grouping._fingerprint", return_value=result.files[0].fingerprint):
            with self.assertRaisesRegex(ValueError, "重新扫描"):
                export_file_groups(result, self.output)
        self.assertFalse(self.output.exists())

    def test_input_and_output_overlap_and_existing_outputs_are_rejected(self):
        self.book("a.xlsx", [("文本", None)])
        result = self.scan()
        for path in (self.source, self.source / "output", self.root, ""):
            with self.subTest(output=path), self.assertRaises(ValueError):
                export_file_groups(result, path)
        self.output.mkdir()
        existing = self.output / "preserved.txt"
        existing.write_text("keep")
        with self.assertRaises(ValueError):
            export_file_groups(result, self.output)
        self.assertEqual(existing.read_text(), "keep")

    def test_file_created_during_copy_is_not_overwritten_and_temporary_is_cleaned(self):
        self.book("a.xlsx", [("文本", None)])
        result = self.scan()
        target = self.output / "其他" / "a.xlsx"
        import shutil
        original = shutil.copyfile

        def race(source, temporary):
            original(source, temporary)
            target.write_bytes(b"created by someone else")

        with patch("tools.file_grouping.grouping.shutil.copyfile", side_effect=race):
            with self.assertRaisesRegex(RuntimeError, "未覆盖"):
                export_file_groups(result, self.output)
        self.assertEqual(target.read_bytes(), b"created by someone else")
        self.assertFalse(list(self.output.rglob("*.tmp")))

    def test_new_file_during_export_prevents_complete_result(self):
        self.book("a.xlsx", [("文本", None)])
        result = self.scan()
        with self.assertRaisesRegex(RuntimeError, "来源文件列表已变化"):
            export_file_groups(result, self.output,
                               progress_callback=lambda *_: self.book("new.xlsx", [("文本", None)]))
        self.assertTrue((self.output / "其他/a.xlsx").is_file())

    def test_invalid_configuration_and_empty_directories(self):
        for kwargs in ({"source_header": ""}, {"target_header": ""}, {"target_header": "MsgStr"}):
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                self.scan(**kwargs)
        with self.assertRaisesRegex(ValueError, "没有支持"):
            self.scan()
        with self.assertRaisesRegex(ValueError, "选择来源"):
            scan_file_groups("")


if __name__ == "__main__":
    unittest.main()
