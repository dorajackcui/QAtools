from __future__ import annotations

from contextlib import contextmanager
from concurrent.futures import ThreadPoolExecutor
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock, patch

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from tools.content_sync.target_to_master import sync_targets_to_master
from tools.deep_replace.replacer import replace_files
from tools.untranslated_stats.stats import count_units, untranslated_stats
from tools.excel_com import edit_copy, excel_session
from tools.excel_compatibility.processor import resave_workbooks
from tools.column_tools.processor import operate_columns


def workbook(path, rows, title="Data"):
    path.parent.mkdir(parents=True, exist_ok=True)
    book = Workbook()
    try:
        book.active.title = title
        for row in rows:
            book.active.append(row)
        book.active["C2"].fill = PatternFill("solid", fgColor="FFFF00")
        book.create_sheet("Unchanged")["A1"] = "retained"
        book.save(path)
    finally:
        book.close()


def values(path, addresses):
    book = load_workbook(path)
    try:
        return [book.active[address].value for address in addresses]
    finally:
        book.close()


class UtilityFixtures(unittest.TestCase):
    def setUp(self):
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.targets = self.root / "targets"
        self.targets.mkdir()

    def tearDown(self):
        # None of the six migrated tools should leave an auxiliary report/log file.
        self.assertEqual(list(self.root.rglob("*.json")), [])
        self.assertEqual(list(self.root.rglob("*.log")), [])


class ReverseSyncTests(UtilityFixtures):
    def setUp(self):
        super().setUp()
        self.master = self.root / "master.xlsx"
        workbook(self.master, [["id", "key", "source", "translation"],
                               [1, "a", "Alpha", "old"], [2, "b", "Beta", None],
                               [3, "a", "Alpha", "duplicate"]])

    def run_sync(self, **kwargs):
        kwargs.setdefault("output_dir", self.root / "updated")
        return sync_targets_to_master(self.master, self.targets, **kwargs)

    def test_default_output_updates_original_master_and_preserves_small_tables(self):
        path = self.targets / "a.xlsx"
        workbook(path, [["k", "s", "t"], ["a", "Alpha", "new"]])
        original = path.read_bytes()
        result = sync_targets_to_master(self.master, self.targets)
        self.assertEqual(Path(result.details["master_output"]), self.master)
        self.assertEqual(values(self.master, ["D2", "D4"]), ["new", "new"])
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual(set(self.root.iterdir()), {self.master, self.targets})

    def test_default_master_save_failure_keeps_original(self):
        workbook(self.targets / "a.xlsx", [["k", "s", "t"], ["a", "Alpha", "new"]])
        original = self.master.read_bytes()
        with patch("tools.content_sync.master_to_target.os.replace", side_effect=PermissionError("locked")):
            result = sync_targets_to_master(self.master, self.targets)
        self.assertEqual(result.failed_files, 1)
        self.assertEqual(self.master.read_bytes(), original)

    def test_order_blank_duplicates_unknown_rows_and_source_preservation(self):
        workbook(self.targets / "a.xlsx", [["k", "s", "t"], ["a", "Alpha", "first"], ["b", "Beta", "nan"]])
        workbook(self.targets / "z.xlsx", [["k", "s", "t"], [" a ", " Alpha ", "  last\n "],
                                            ["a", "Alpha", None], ["new", "new", "absent"]])
        original = self.master.read_bytes()
        logs = []
        summary = self.run_sync(log_callback=logs.append)
        self.assertEqual(values(Path(summary.details["master_output"]), ["D2", "D3", "D4"]),
                         ["  last\n ", "nan", "  last\n "])
        self.assertEqual(self.master.read_bytes(), original)
        self.assertEqual(summary.details["updated_cells"], 3)
        self.assertEqual(len(summary.details["conflicts"]), 1)
        self.assertEqual(summary.details["unmatched_candidates"][0]["key"], "new")
        self.assertEqual(summary.tool, "target-to-master")
        self.assertEqual([p.name for p in summary.output_dir.iterdir()], [self.master.name])
        self.assertTrue(any("跨文件重复身份" in line for line in logs))
        self.assertTrue(any("Master 中未找到身份" in line for line in logs))

    def test_fill_only_and_allow_blank_are_independent(self):
        workbook(self.targets / "a.xlsx", [["k", "s", "t"], ["a", "Alpha", "new"], ["b", "Beta", False]])
        summary = self.run_sync(fill_blank_only=True)
        self.assertEqual(values(Path(summary.details["master_output"]), ["D2", "D3", "D4"]),
                         ["old", "False", "duplicate"])
        workbook(self.targets / "z.xlsx", [["k", "s", "t"], ["a", "Alpha", None], ["b", "Beta", 0]])
        summary = self.run_sync(output_dir=self.root / "second", allow_blank_write=True)
        self.assertEqual(values(Path(summary.details["master_output"]), ["D2", "D3", "D4"]), [None, "0", None])

    def test_literal_missing_markers_and_formula_like_text(self):
        tokens = ["None", "nan", "NA", "NULL", "=literal", "#N/A"]
        workbook(self.master, [["id", "k", "s", "t"]] + [[i, str(i), "S", "old"] for i in range(len(tokens))])
        path = self.targets / "a.xlsx"
        workbook(path, [["k", "s", "t"]] + [[str(i), "S", token] for i, token in enumerate(tokens)])
        book = load_workbook(path)
        try:
            for row in book.active.iter_rows(min_row=2):
                row[2].data_type = "s"
            book.save(path)
        finally:
            book.close()
        result = self.run_sync()
        book = load_workbook(result.details["master_output"])
        try:
            self.assertEqual([book.active.cell(i + 2, 4).value for i in range(len(tokens))], tokens)
            self.assertTrue(all(book.active.cell(i + 2, 4).data_type == "s" for i in range(len(tokens))))
        finally:
            book.close()

    def test_bad_file_continues_and_all_bad_does_not_emit_master(self):
        (self.targets / "bad.xlsx").write_text("invalid")
        summary = self.run_sync()
        self.assertIsNone(summary.details["master_output"])
        self.assertEqual(summary.failed_files, 2)
        workbook(self.targets / "good.xlsx", [["k", "s", "t"], ["a", "Alpha", "new"]])
        summary = self.run_sync(output_dir=self.root / "second")
        self.assertEqual(summary.failed_files, 1)
        self.assertEqual(values(Path(summary.details["master_output"]), ["D2"]), ["new"])

    def test_missing_formula_cache_does_not_partially_merge_file(self):
        workbook(self.targets / "a.xlsx", [["k", "s", "t"], ["a", "Alpha", "partial"], ["b", "Beta", "=1+1"]])
        workbook(self.targets / "b.xlsx", [["k", "s", "t"], ["b", "Beta", "ok"]])
        summary = self.run_sync()
        self.assertEqual(values(Path(summary.details["master_output"]), ["D2", "D3"]), ["old", "ok"])
        self.assertEqual(summary.failed_files, 1)

    def test_existing_output_rejected_and_nested_previous_output_excluded(self):
        workbook(self.targets / "a.xlsx", [["k", "s", "t"], ["a", "Alpha", "new"]])
        output = self.targets / "generated"
        self.run_sync(output_dir=output)
        with self.assertRaises(ValueError):
            self.run_sync(output_dir=output)
        result = self.run_sync(output_dir=self.root / "next")
        self.assertEqual(len([r for r in result.results if r.status == "read"]), 1)

    def test_master_save_failure_has_no_partial_output(self):
        workbook(self.targets / "a.xlsx", [["k", "s", "t"], ["a", "Alpha", "new"]])
        original = self.master.read_bytes()
        with patch("tools.content_sync.master_to_target._save_output", side_effect=PermissionError("locked")):
            result = self.run_sync()
        self.assertEqual(result.failed_files, 1)
        self.assertEqual(result.details["updated_cells"], 0)
        self.assertIsNone(result.details["master_output"])
        self.assertFalse((result.output_dir / self.master.name).exists())
        self.assertEqual(self.master.read_bytes(), original)


class StatsTests(UtilityFixtures):
    def test_default_saves_beside_inputs_and_repeated_run_excludes_own_result(self):
        path = self.targets / "a.xlsx"
        workbook(path, [["id", "s", "t"], [1, "你好", None]])
        original = path.read_bytes()
        first = untranslated_stats(self.targets)
        second = untranslated_stats(self.targets)
        self.assertEqual(first.output_dir, self.targets)
        self.assertEqual(second.details["totals"], [2, 1, 2, 1])
        self.assertEqual(len(second.results), 1)
        self.assertEqual(path.read_bytes(), original)
        self.assertEqual({p.name for p in self.targets.iterdir()}, {"a.xlsx", "未翻译统计.xlsx"})
        copied = untranslated_stats(self.targets, output_dir=self.root / "statistics")
        self.assertEqual(copied.details["totals"], second.details["totals"])

    def test_counts_and_excel_report_with_failure_details(self):
        workbook(self.targets / "a.xlsx", [["id", "source", "target"], [1, "你好 world", None],
                                            [2, "世界", " NaN "], [3, "完成", "None"],
                                            [4, 0, False], [5, "   ", None]])
        (self.targets / "broken.xlsx").write_text("bad")
        (self.targets / "old.xls").write_bytes(b"old")
        logs = []
        summary = untranslated_stats(self.targets, log_callback=logs.append)
        self.assertTrue(any("[失败] broken.xlsx" in line for line in logs))
        self.assertTrue(any("未翻译量: 4" in line for line in logs))
        self.assertEqual([line for line in logs if "a.xlsx" in line],
                         ["[已读取] a.xlsx · 未翻译量: 4/6；未翻译行: 2/4"])
        self.assertEqual(summary.details["totals"], [4, 2, 6, 4])
        self.assertEqual(summary.failed_files, 1)
        self.assertEqual(summary.skipped_files, 1)
        book = load_workbook(summary.details["excel_report"])
        try:
            self.assertEqual(book.sheetnames, ["未翻译统计", "未统计文件"])
            self.assertEqual(list(book.active.values)[-1], ("总计", 4, 2, 6, 4))
            self.assertEqual(list(book.active.values)[0], ("文件名", "未翻译字数", "未翻译行数", "总字数", "总行数"))
        finally:
            book.close()

    def test_english_units_follow_original_rule(self):
        self.assertEqual(count_units("don't re-use ABC 123 中文", "english_words"), 3)
        workbook(self.targets / "a.xlsx", [["id", "s", "t"], [1, "don't re-use", None], [2, "one two", "yes"]])
        summary = untranslated_stats(self.targets, mode="english_words")
        self.assertEqual(summary.details["totals"], [2, 1, 4, 2])

    def test_sheet_header_and_column_validation(self):
        workbook(self.targets / "a.xlsx", [["k", "s", "t"], ["k", "你好", None]])
        summary = untranslated_stats(self.targets, sheet="missing")
        self.assertEqual(summary.failed_files, 1)
        with self.assertRaises(ValueError):
            untranslated_stats(self.targets, source_column="C", target_column="C")
        with self.assertRaises(ValueError):
            untranslated_stats(self.targets, header_rows=-1)


class DeepReplaceTests(UtilityFixtures):
    def setUp(self):
        super().setUp()
        self.sources = self.root / "sources"
        self.sources.mkdir()

    def file(self, root, name, content):
        path = root / name
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_bytes(content)
        return path

    def test_case_insensitive_copy_preserves_tree_and_unmatched_targets(self):
        self.file(self.sources, "A.XLSX", b"new")
        self.file(self.sources, "missing.xlsb", b"unmatched")
        target = self.file(self.targets, "sub/a.xlsx", b"old")
        self.file(self.targets, "keep.xls", b"keep")
        logs = []
        summary = replace_files(self.sources, self.targets, output_dir=self.root / "replaced", log_callback=logs.append)
        self.assertTrue(any("[已替换]" in line for line in logs))
        self.assertEqual((summary.output_dir / "sub/a.xlsx").read_bytes(), b"new")
        self.assertEqual((summary.output_dir / "keep.xls").read_bytes(), b"keep")
        self.assertEqual(target.read_bytes(), b"old")
        self.assertEqual(summary.details["replaced_files"], 1)
        self.assertEqual(summary.details["unmatched_sources"], ["missing.xlsb"])

    def test_duplicate_sources_or_targets_are_skipped(self):
        for root, name in ((self.sources, "one/a.xlsx"), (self.sources, "two/a.xlsx"),
                           (self.targets, "a.xlsx"), (self.sources, "b.xlsx"),
                           (self.targets, "one/b.xlsx"), (self.targets, "two/b.xlsx")):
            self.file(root, name, name.encode())
        summary = replace_files(self.sources, self.targets)
        self.assertEqual(summary.skipped_files, 3)
        self.assertEqual(summary.details["replaced_files"], 0)
        self.assertEqual((summary.output_dir / "a.xlsx").read_bytes(), b"a.xlsx")

    def test_inplace_backup_and_atomic_failure_preserve_original(self):
        self.file(self.sources, "a.xlsx", b"new")
        target = self.file(self.targets, "sub/a.xlsx", b"old")
        summary = replace_files(self.sources, self.targets)
        self.assertEqual(target.read_bytes(), b"new")
        self.assertEqual(Path(summary.results[0].details["backup"]).read_bytes(), b"old")
        second = replace_files(self.sources, self.targets)
        self.assertEqual(len(second.results), 1, "backup must not become a duplicate target")
        self.assertNotEqual(second.details["backup_dir"], summary.details["backup_dir"])
        self.assertEqual(Path(summary.results[0].details["backup"]).read_bytes(), b"old")
        with patch("tools.excel_file_ops.os.replace", side_effect=PermissionError("locked")):
            # Report also uses os.replace; patch only the copy operation below instead.
            from tools.excel_file_ops import atomic_copy
            with self.assertRaises(PermissionError):
                atomic_copy(self.sources / "a.xlsx", target)
        self.assertEqual(target.read_bytes(), b"new")
        self.assertEqual(len(list(target.parent.iterdir())), 1)

    def test_backup_failure_prevents_inplace_replacement(self):
        self.file(self.sources, "a.xlsx", b"new")
        target = self.file(self.targets, "a.xlsx", b"old")
        with patch("tools.deep_replace.replacer.atomic_copy", side_effect=PermissionError("backup locked")) as copy:
            summary = replace_files(self.sources, self.targets)
        copy.assert_called_once()
        self.assertEqual(summary.failed_files, 1)
        self.assertEqual(target.read_bytes(), b"old")

    def test_overlapping_directories_rejected(self):
        with self.assertRaises(ValueError):
            replace_files(self.targets, self.targets / "child")


class ExcelComTests(UtilityFixtures):
    def test_session_isolated_and_cleaned_up_on_exception(self):
        com, client = Mock(), Mock()
        app = client.DispatchEx.return_value
        with patch("tools.excel_com._load_com", return_value=(com, client)):
            with self.assertRaisesRegex(ValueError, "failure"):
                with excel_session() as excel:
                    self.assertIs(excel, app)
                    self.assertFalse(excel.Visible)
                    self.assertEqual(excel.AutomationSecurity, 3)
                    raise ValueError("failure")
        client.DispatchEx.assert_called_once_with("Excel.Application")
        app.Quit.assert_called_once()
        com.CoUninitialize.assert_called_once()

    def test_startup_failure_uninitializes_com(self):
        com, client = Mock(), Mock()
        client.DispatchEx.side_effect = RuntimeError("missing")
        with patch("tools.excel_com._load_com", return_value=(com, client)):
            with self.assertRaisesRegex(RuntimeError, "无法启动"):
                with excel_session():
                    self.fail("should not start")
        com.CoUninitialize.assert_called_once()

    def test_unsupported_platform_and_missing_optional_dependency_are_actionable(self):
        from tools.excel_com import _load_com
        with patch("tools.excel_com.sys.platform", "linux"):
            with self.assertRaisesRegex(RuntimeError, "Windows"):
                _load_com()
        with patch("tools.excel_com.sys.platform", "win32"), patch.dict("sys.modules", {"pythoncom": None}):
            with self.assertRaisesRegex(RuntimeError, "excel-com"):
                _load_com()

    def test_quit_failure_is_reported_and_com_is_still_uninitialized(self):
        com, client = Mock(), Mock()
        client.DispatchEx.return_value.Quit.side_effect = RuntimeError("quit failed")
        with patch("tools.excel_com._load_com", return_value=(com, client)):
            with self.assertRaisesRegex(RuntimeError, "quit failed"):
                with excel_session():
                    pass
        com.CoUninitialize.assert_called_once()

    def test_edit_failure_closes_workbook_and_preserves_source_and_output(self):
        source, destination = self.targets / "a.xlsx", self.root / "result.xlsx"
        source.write_bytes(b"original")
        destination.write_bytes(b"existing")
        application = Mock()
        book = application.Workbooks.Open.return_value
        book.ReadOnly = False
        book.Save.side_effect = RuntimeError("save failed")
        with self.assertRaisesRegex(RuntimeError, "save failed"):
            edit_copy(application, source, destination)
        book.Close.assert_called_once_with(SaveChanges=False)
        self.assertEqual(source.read_bytes(), b"original")
        self.assertEqual(destination.read_bytes(), b"existing")
        self.assertFalse(list(self.root.glob(".qatools-excel-*")))

    def test_preflight_no_excel_does_not_create_output(self):
        (self.targets / "a.xlsx").write_bytes(b"data")
        output = self.root / "out"
        with patch("tools.excel_compatibility.processor.excel_session", side_effect=RuntimeError("not installed")):
            with self.assertRaises(RuntimeError):
                resave_workbooks(self.targets, output_dir=output)
        self.assertFalse(output.exists())

    def test_clear_uses_actual_last_row_and_skips_header_only_range(self):
        source = self.targets / "a.xlsx"
        source.write_bytes(b"original")
        application = Mock()
        book = application.Workbooks.Open.return_value
        book.ReadOnly = False
        sheet = book.ActiveSheet
        sheet.Columns.Count = 16384
        sheet.UsedRange.Row = 10
        sheet.UsedRange.Rows.Count = 5
        edit_copy(application, source, self.root / "out.xlsx", action="clear", column=3)
        sheet.Cells.assert_any_call(2, 3)
        sheet.Cells.assert_any_call(14, 3)
        sheet.Range.return_value.ClearContents.assert_called_once()
        sheet.reset_mock()
        sheet.UsedRange.Row = 1
        sheet.UsedRange.Rows.Count = 1
        edit_copy(application, source, self.root / "header.xlsx", action="clear", column=3)
        sheet.Range.assert_not_called()

    def test_batch_continues_after_failed_workbook(self):
        for name in ("a.xlsx", "b.xlsm", "unsupported.xlsb"):
            (self.targets / name).write_bytes(b"data")
        @contextmanager
        def session():
            yield Mock()
        with patch("tools.excel_compatibility.processor.excel_session", session), \
             patch("tools.excel_compatibility.processor.edit_copy", side_effect=[RuntimeError("locked"), None]):
            logs = []
            summary = operate_columns(self.targets, column="AA", action="insert", log_callback=logs.append)
        self.assertTrue(any("[失败] a.xlsx — locked" in line for line in logs))
        self.assertTrue(any("[已更新] b.xlsm" in line for line in logs))
        self.assertEqual((summary.failed_files, summary.succeeded_files, summary.skipped_files), (1, 1, 1))
        self.assertEqual(summary.details["column"], 27)


@unittest.skipUnless(os.environ.get("QATOOLS_TEST_EXCEL_COM") == "1", "opt-in real Microsoft Excel smoke test")
class RealExcelTests(UtilityFixtures):
    def test_default_native_actions_update_original_workbook(self):
        source = self.targets / "a.xlsx"
        for action in ("clear", "insert", "delete", "resave"):
            workbook(source, [["id", "s", "t"], [1, "你好", "old"]])
            with ThreadPoolExecutor(max_workers=1) as executor:
                future = (executor.submit(resave_workbooks, self.targets) if action == "resave"
                          else executor.submit(operate_columns, self.targets, column="C", action=action))
                result = future.result(timeout=60)
            self.assertEqual(result.failed_files, 0, str(result.results))
            self.assertEqual(Path(result.results[0].output), source)
            self.assertEqual(set(self.root.iterdir()), {self.targets})
            self.assertEqual(values(source, ["C2", "D2"]),
                             [None, "old"] if action == "insert" else ["old", None] if action == "resave" else [None, None])

    def test_legacy_and_macro_enabled_formats_and_literal_insert_header(self):
        source = self.root / "fixture.xlsx"
        workbook(source, [["id", "source", "translation"], [1, "你好", "old"]])
        with excel_session() as excel:
            for suffix, format_id in ((".xls", 56), (".xlsm", 52)):
                original = self.targets / f"fixture{suffix}"
                destination = self.root / "output" / original.name
                book = excel.Workbooks.Open(str(source), UpdateLinks=0, ReadOnly=False)
                try:
                    book.SaveAs(str(original), FileFormat=format_id)
                finally:
                    book.Close(SaveChanges=False)
                    book = None
                original_bytes = original.read_bytes()
                edit_copy(excel, original, destination, action="insert", column=3, inserted_header="=literal")
                book = excel.Workbooks.Open(str(destination), UpdateLinks=0, ReadOnly=True)
                try:
                    self.assertEqual(book.ActiveSheet.Cells(1, 3).Value, "=literal")
                    self.assertFalse(book.ActiveSheet.Cells(1, 3).HasFormula)
                    self.assertEqual(book.ActiveSheet.Cells(2, 4).Value, "old")
                finally:
                    book.Close(SaveChanges=False)
                    book = None
                self.assertEqual(original.read_bytes(), original_bytes)

    def test_forward_optional_resave_keeps_transferred_markers_as_text(self):
        from tools.content_sync.master_to_target import sync_master_to_targets
        master = self.root / "master.xlsx"
        workbook(master, [["id", "k", "s", "t"], [1, "a", "A", "None"], [2, "b", "B", "nan"]])
        workbook(self.targets / "a.xlsx", [["k", "s", "t"], ["a", "A", "old"], ["b", "B", "old"]])
        with ThreadPoolExecutor(max_workers=1) as executor:
            result = executor.submit(sync_master_to_targets, master, self.targets, compatibility_resave=True).result(timeout=60)
        self.assertEqual(result.failed_files, 0, str(result.files))
        self.assertFalse(result.warnings)
        self.assertEqual(values(result.output_dir / "a.xlsx", ["C2", "C3"]), ["None", "nan"])

    def test_clear_insert_delete_resave_using_independent_excel_sessions(self):
        source = self.targets / "a.xlsx"
        workbook(source, [["id", "source", "translation"], [1, "你好", "old"], [2, "Hi", "=1+1"]])
        original = source.read_bytes()
        for action in ("clear", "insert", "delete", "resave"):
            with self.subTest(action=action):
                output = self.root / action
                with ThreadPoolExecutor(max_workers=1) as executor:
                    task = (executor.submit(resave_workbooks, self.targets, output_dir=output) if action == "resave"
                            else executor.submit(operate_columns, self.targets, column="C", action=action, output_dir=output))
                    summary = task.result(timeout=60)
                self.assertEqual(summary.failed_files, 0, summary.describe() + str(summary.results))
                self.assertFalse(summary.warnings)
                book = load_workbook(output / "a.xlsx")
                try:
                    sheet = book.active
                    self.assertEqual(book["Unchanged"]["A1"].value, "retained")
                    if action == "clear":
                        self.assertEqual(sheet["C1"].value, "translation")
                        self.assertIsNone(sheet["C2"].value)
                        self.assertIsNone(sheet["C3"].value)
                        self.assertEqual(sheet["C2"].fill.fgColor.rgb[-6:], "FFFF00")
                    elif action == "insert":
                        self.assertEqual(sheet["C1"].value, "Translation")
                        self.assertEqual(sheet["D2"].value, "old")
                    elif action == "delete":
                        self.assertIsNone(sheet["C1"].value)
                        self.assertIsNone(sheet["C2"].value)
                    else:
                        self.assertEqual(sheet["C2"].value, "old")
                        self.assertEqual(sheet["C3"].value, "=1+1")
                finally:
                    book.close()
        self.assertEqual(source.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()
