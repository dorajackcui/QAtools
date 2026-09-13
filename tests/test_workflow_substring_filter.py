from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from tools.term_pair_checker import extract_terms_from_excel as term_checker
from tools.workflow.workflow_runner import run_workflow


class WorkflowSubstringFilterTests(unittest.TestCase):
    def test_reuses_active_tb_and_batch_terms_but_not_blank_translations_or_stale_sheets(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "input.xlsx"
            tb = root / "tb.xlsx"
            book = Workbook()
            try:
                book.active.title = "Data"
                for row in [
                    ("source", "target"),
                    ("阿童木", "Astro Boy"), ("阿童木登场", "Wrong"),
                    ("光合作", "Cooperation"), ("光合作用", "Photosynthesis"),
                    ("确认设置", "Confirm settings"), ("请确认设置", "Wrong"),
                    ("【新术语】", "【New term】"), ("新术语提示", "Wrong"),
                    ("无译术语", "Reference"), ("无译术语提示", "Wrong"),
                ]:
                    book.active.append(row)
                stale = book.create_sheet("术语表")
                stale.append(["source术语", "target术语"])
                stale.append(["确认设置", "Confirm settings"])
                book.save(source)
            finally:
                book.close()
            book = Workbook()
            try:
                for row in [
                    ("source", "target"), ("阿童木", "Astro Boy"),
                    ("光合作用", "Photosynthesis"), ("无译术语", None),
                    ("未涉及有效词", "Unused"),
                ]:
                    book.active.append(row)
                book.save(tb)
            finally:
                book.close()

            for enabled, expected in ((True, [7, 11]), (False, [3, 5, 7, 9, 11])):
                with self.subTest(term_check=enabled), patch.object(
                    term_checker, "load_history_tb_mapping", wraps=term_checker.load_history_tb_mapping,
                ) as load_tb:
                    result = run_workflow(
                        input_file=source, source_column="A", target_column="B", sheet="Data",
                        output_file=root / f"result-{enabled}.xlsx",
                        run_term_pair_check=enabled, term_history_tb_file=tb,
                        run_tag_check=False, run_line_break_check=False,
                        run_source_consistency_check=False, run_number_check=False,
                        run_url_check=False, run_chinese_target_check=False,
                        run_target_text_check=False, run_substring_consistency_check=True,
                    )
                    self.assertEqual(load_tb.call_count, int(enabled))
                    self.assertEqual(result.substring_consistency_problem_rows, len(expected))
                    report = load_workbook(result.output_path)
                    try:
                        rows = [row[0] for row in report["问题处理"].iter_rows(
                            min_row=2, max_col=6, values_only=True,
                        ) if row[5] and "子串译文一致性" in row[5]]
                        self.assertEqual(rows, expected)
                    finally:
                        report.close()

    def test_runner_rejects_invalid_enabled_threshold_before_file_access(self):
        with self.assertRaisesRegex(ValueError, "最小有效字符数"):
            run_workflow(
                input_file="missing.xlsx", source_column="A", target_column="B",
                run_substring_consistency_check=True, substring_min_cjk_chars=0,
            )


if __name__ == "__main__":
    unittest.main()
