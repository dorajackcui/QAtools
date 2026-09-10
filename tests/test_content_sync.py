from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill

from tools.content_sync.master_to_target import (
    ColumnMapping, content_text, is_blank, sync_master_to_targets,
)


class ContentSyncTests(unittest.TestCase):
    def test_equal_formula_or_error_is_written_as_text_in_both_directions(self):
        from tools.content_sync.target_to_master import sync_targets_to_master

        for reverse in (False, True):
            for value, kind in (("=1+1", "f"), ("#N/A", "e")):
                for fill_only in (False, True):
                    with self.subTest(reverse=reverse, value=value, fill_only=fill_only):
                        target = self.targets / "a.xlsx"
                        self.write(self.master, [["id", "k", "s", "t"], [1, "k", "s", value]])
                        self.write(target, [["k", "s", "t"], ["k", "s", value]])
                        source, address = (target, "C2") if reverse else (self.master, "D2")
                        book = load_workbook(source)
                        try:
                            book.active[address].data_type = "s"
                            book.save(source)
                        finally:
                            book.close()
                        operation = sync_targets_to_master if reverse else sync_master_to_targets
                        result = operation(self.master, self.targets, fill_blank_only=fill_only)
                        destination, address = (self.master, "D2") if reverse else (target, "C2")
                        book = load_workbook(destination)
                        try:
                            self.assertEqual(book.active[address].value, value)
                            self.assertEqual(book.active[address].data_type, kind if fill_only else "s")
                        finally:
                            book.close()
                        updates = result.details["updated_cells"] if reverse else result.updated_cells
                        self.assertEqual(updates, 0 if fill_only else 1)

    def test_successful_file_has_one_compact_log_record(self):
        self.write(self.master, [["id", "key", "source", "target"], [1, "k", "s", "new"]])
        self.write(self.targets / "a.xlsx", [["key", "source", "target"], ["k", "s", "old"]])
        logs = []
        result = sync_master_to_targets(self.master, self.targets, log_callback=logs.append)
        file_logs = [line for line in logs if "a.xlsx" in line]
        self.assertEqual(file_logs, ["[已更新] a.xlsx · 匹配 1 行，更新 1 格"])
        self.assertEqual(result.updated_cells, 1)
        self.assertTrue(logs[-1].startswith("完成：成功 1，失败 0，跳过 0"))

    def setUp(self):
        directory = tempfile.TemporaryDirectory()
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.master = self.root / "master.xlsx"
        self.targets = self.root / "targets"
        self.targets.mkdir()
        self.output = self.root / "synced"

    def write(self, path, rows, *, sheet="Data", active_second=False):
        path.parent.mkdir(parents=True, exist_ok=True)
        workbook = Workbook()
        if active_second:
            workbook.active.title = "Ignored"
            workbook.active.append(["do not change"])
            worksheet = workbook.create_sheet(sheet)
            workbook.active = 1
        else:
            worksheet = workbook.active
            worksheet.title = sheet
        for row in rows:
            worksheet.append(row)
        workbook.save(path)
        workbook.close()

    @contextmanager
    def book(self, path):
        workbook = load_workbook(path)
        try:
            yield workbook
        finally:
            workbook.close()

    def sync(self, **kwargs):
        return sync_master_to_targets(self.master, self.targets, output_dir=self.output, **kwargs)

    def basic(self, value="new", target="old"):
        self.write(self.master, [["id", "key", "source", "target"], [1, "k", "s", value]])
        self.write(self.targets / "a.xlsx", [["key", "source", "target"], ["k", "s", target]])

    def test_blank_output_updates_targets_in_place_and_unchanged_file_is_untouched(self):
        self.basic(value="  nan\n ")
        original_master = self.master.read_bytes()
        result = sync_master_to_targets(self.master, self.targets, output_dir="  ")
        self.assertEqual(result.output_dir, self.targets)
        self.assertEqual(result.updated_cells, 1)
        self.assertEqual(self.master.read_bytes(), original_master)
        with self.book(self.targets / "a.xlsx") as book:
            self.assertEqual(book.active["C2"].value, "  nan\n ")
        original = (self.targets / "a.xlsx").read_bytes()
        with patch("tools.content_sync.master_to_target.os.replace", side_effect=AssertionError("unchanged file rewritten")):
            result = sync_master_to_targets(self.master, self.targets)
        self.assertEqual(result.updated_cells, 0)
        self.assertEqual((self.targets / "a.xlsx").read_bytes(), original)
        self.assertFalse(self.output.exists())

    def test_inplace_replace_failure_preserves_original_bytes(self):
        self.basic()
        target = self.targets / "a.xlsx"
        original = target.read_bytes()
        with patch("tools.content_sync.master_to_target.os.replace", side_effect=PermissionError("locked")):
            result = sync_master_to_targets(self.master, self.targets)
        self.assertEqual(result.failed_files, 1)
        self.assertEqual(result.updated_cells, 0)
        self.assertEqual(target.read_bytes(), original)
        self.assertEqual(list(self.targets.iterdir()), [target])

    def test_optional_resave_preflight_and_failure_preserve_synced_output(self):
        self.basic()
        with patch("tools.excel_com.excel_session", side_effect=RuntimeError("missing Excel")):
            with self.assertRaisesRegex(RuntimeError, "missing Excel"):
                self.sync(compatibility_resave=True)
        self.assertFalse(self.output.exists())
        @contextmanager
        def session():
            yield object()
        self.write(self.targets / "unchanged.xlsx", [["key", "source", "target"], ["k", "s", "new"]])
        with patch("tools.excel_com.excel_session", session), \
             patch("tools.excel_com.edit_copy", side_effect=RuntimeError("resave failed")) as resave:
            summary = self.sync(compatibility_resave=True)
        resave.assert_called_once()
        self.assertEqual(summary.updated_cells, 1)
        self.assertEqual(summary.failed_files, 1)
        self.assertEqual(summary.succeeded_files, 1)
        self.assertIn("resave failed", summary.files[0].postprocess_error)
        with self.book(self.output / "a.xlsx") as book:
            self.assertEqual(book.active["C2"].value, "new")

    def test_single_preserves_inputs_nested_structure_styles_and_other_sheets(self):
        self.basic()
        path = self.targets / "nested" / "a.xlsx"
        self.write(path, [["key", "source", "target"], ["k", "s", "old"]], active_second=True)
        with self.book(path) as wb:
            wb.active["C2"].fill = PatternFill("solid", fgColor="FF0000")
            wb.active["D2"] = "=1+2"
            wb["Ignored"].sheet_state = "hidden"
            wb.save(path)
        original = {p: p.read_bytes() for p in (self.master, self.targets / "a.xlsx", path)}
        summary = self.sync()
        self.assertEqual(summary.updated_cells, 2)
        self.assertEqual(summary.succeeded_files, 2)
        with self.book(self.output / "nested" / "a.xlsx") as wb:
            self.assertEqual(wb.active["C2"].value, "new")
            self.assertEqual(wb.active["C2"].fill.fgColor.rgb, "00FF0000")
            self.assertEqual(wb.active["D2"].value, "=1+2")
            self.assertEqual(wb["Ignored"].sheet_state, "hidden")
        for p, value in original.items():
            self.assertEqual(p.read_bytes(), value)

    def test_value_contract_round_trip(self):
        values = [None, "", "   ", "None", "nan", "NA", "NULL", 0, False,
                  "00123", "  <b>hello</b>\nworld  ", "=literal", "#N/A", datetime(2026, 9, 10)]
        self.write(self.master, [["key", "match", "value"]] +
                   [[str(i), "s", value] for i, value in enumerate(values)])
        # Explicitly store formula-like text as text in the source workbook.
        with self.book(self.master) as wb:
            wb.active.cell(13, 3).data_type = "s"
            wb.active.cell(14, 3).data_type = "s"
            wb.save(self.master)
        self.write(self.targets / "a.xlsx", [["key", "match", "value"]] +
                   [[str(i), "s", "keep"] for i in range(len(values))])
        self.sync(master_columns=ColumnMapping("A", "B", "C"))
        with self.book(self.output / "a.xlsx") as wb:
            for i, value in enumerate(values, 2):
                cell = wb.active.cell(i, 3)
                self.assertEqual(cell.value, "keep" if is_blank(value) else content_text(value))
                if not is_blank(value):
                    self.assertEqual(cell.data_type, "s")

    def test_blank_write_and_fill_blank_are_independent_per_cell(self):
        for fill, allow in ((False, False), (False, True), (True, False), (True, True)):
            with self.subTest(fill=fill, allow=allow):
                self.output = self.root / f"out_{fill}_{allow}"
                self.write(self.master, [["key", "match", "v1", "v2", "v3"],
                                         ["k", "s", None, "new", "   "]])
                self.write(self.targets / "a.xlsx", [["key", "match", "v1", "v2", "v3"],
                                                      ["k", "s", "old", None, "old"],
                                                      ["k", "s", None, "keep", None]])
                self.sync(master_columns=ColumnMapping("A", "B", "C"), column_count=3,
                          fill_blank_only=fill, allow_blank_write=allow)
                with self.book(self.output / "a.xlsx") as wb:
                    self.assertEqual(wb.active["C2"].value, None if allow and not fill else "old")
                    self.assertEqual(wb.active["D2"].value, "new")
                    self.assertEqual(wb.active["D3"].value, "keep" if fill else "new")
                    self.assertEqual(wb.active["E2"].value, "   " if allow and not fill else "old")
                    self.assertEqual(wb.active["E3"].value, "   " if allow else None)

    def test_exact_tuple_identity_duplicates_and_unmatched_rows(self):
        self.write(self.master, [["key", "match", "value"],
                                 ["a|b", "c", "first"], ["a", "b|c", "different"],
                                 [" a|b ", " c ", "last"], ["", "c", "skip"]])
        self.write(self.targets / "a.xlsx", [["key", "match", "value"],
                                              ["a|b", "c", "old"], ["a", "b|c", None],
                                              ["a|b", "C", "keep"], [None, "c", "keep"],
                                              ["a|b", "c", "old"]])
        logs = []
        summary = self.sync(master_columns=ColumnMapping("A", "B", "C"), log_callback=logs.append)
        with self.book(self.output / "a.xlsx") as wb:
            self.assertEqual([wb.active.cell(i, 3).value for i in range(2, 7)],
                             ["last", "different", "keep", "keep", "last"])
        self.assertEqual(summary.duplicate_keys[0]["previous_row"], 2)
        self.assertEqual(summary.duplicate_keys[0]["selected_row"], 4)
        self.assertTrue(summary.duplicate_keys[0]["conflicting"])
        self.assertEqual(summary.files[0].unmatched_rows, [4])
        self.assertEqual(summary.files[0].invalid_rows, 1)
        self.assertTrue(any("Master 重复身份" in line for line in logs))
        self.assertEqual([p.name for p in self.output.iterdir()], ["a.xlsx"])

    def test_sparse_reordered_columns_and_selected_sheets_and_headers(self):
        self.write(self.master, [["intro"], ["v1", "v2", "key", "source"],
                                 ["one", "two", "k", "s"]], sheet="Master", active_second=True)
        self.write(self.targets / "a.xlsx", [["v1", "v2", "source", "key"],
                                              [None, None, "s", "k"]], sheet="Target", active_second=True)
        summary = self.sync(master_columns=ColumnMapping("C", "D", "A"),
                            target_columns=ColumnMapping("D", "C", "A"), column_count=2,
                            master_sheet="Master", target_sheet="Target", master_header_rows=2)
        self.assertEqual(summary.updated_cells, 2)
        with self.book(self.output / "a.xlsx") as wb:
            self.assertEqual(tuple(wb["Target"].values)[1], ("one", "two", "s", "k"))

    def test_zero_headers_and_missing_source_content_column(self):
        self.write(self.master, [["k", "s", "one"]])
        self.write(self.targets / "a.xlsx", [["k", "s", None, "keep"]])
        self.sync(master_columns=ColumnMapping("A", "B", "C"), column_count=2,
                  master_header_rows=0, target_header_rows=0)
        with self.book(self.output / "a.xlsx") as wb:
            self.assertEqual(wb.active["C1"].value, "one")
            self.assertEqual(wb.active["D1"].value, "keep")

    def test_unchanged_workbook_is_copied_byte_for_byte(self):
        self.basic(target="new")
        data = (self.targets / "a.xlsx").read_bytes()
        summary = self.sync()
        self.assertEqual(summary.updated_cells, 0)
        self.assertEqual(summary.files[0].status, "unchanged")
        self.assertEqual((self.output / "a.xlsx").read_bytes(), data)

    def test_enumeration_ignores_locks_master_and_handles_formats(self):
        self.basic()
        self.master = self.targets / "master.xlsx"
        self.write(self.master, [["id", "key", "match", "value"], [1, "k", "s", "new"]])
        (self.targets / "a.xlsx").rename(self.targets / "a.XLSX")
        (self.targets / "~$temp.xlsx").write_bytes(b"not excel")
        (self.targets / "old.xls").write_bytes(b"legacy")
        progress = []
        summary = self.sync(progress_callback=lambda done, total: progress.append((done, total)))
        self.assertEqual(summary.succeeded_files, 1)
        self.assertEqual(summary.skipped_files, 1)
        self.assertEqual(progress, [(1, 1)])
        self.assertFalse((self.output / "master.xlsx").exists())

    def test_bad_file_and_missing_sheet_are_failures_without_input_writes(self):
        self.basic()
        (self.targets / "broken.xlsx").write_bytes(b"not excel")
        self.write(self.targets / "missing.xlsx", [["key", "source", "target"]], sheet="Other")
        summary = self.sync(target_sheet="Data")
        self.assertEqual(summary.succeeded_files, 1)
        self.assertEqual(summary.failed_files, 2)
        self.assertFalse((self.output / "broken.xlsx").exists())
        self.assertFalse((self.output / "missing.xlsx").exists())
        self.assertTrue(all(f.error for f in summary.files if f.status == "failed"))

    def test_save_failure_does_not_count_updates_or_leave_temporary_file(self):
        self.basic()
        with patch("openpyxl.workbook.workbook.Workbook.save", side_effect=PermissionError("locked")):
            summary = self.sync()
        self.assertEqual(summary.updated_cells, 0)
        self.assertEqual(summary.failed_files, 1)
        self.assertEqual(list(self.output.iterdir()), [])
        with self.book(self.targets / "a.xlsx") as wb:
            self.assertEqual(wb.active["C2"].value, "old")

    def test_rejects_existing_overlapping_outputs_before_writing(self):
        self.basic()
        for output in (self.targets, self.root, self.master):
            with self.subTest(output=output), self.assertRaises(ValueError):
                sync_master_to_targets(self.master, self.targets, output_dir=output)
        self.output.mkdir()
        with self.assertRaisesRegex(ValueError, "已存在"):
            self.sync()
        self.assertEqual(list(self.output.iterdir()), [])

    def test_new_output_inside_input_is_not_reprocessed(self):
        self.basic()
        self.output = self.targets / "result"
        summary = self.sync()
        self.assertEqual(len(summary.files), 1)
        self.assertTrue((self.output / "a.xlsx").exists())

    def test_invalid_columns_and_header_settings_fail_before_creating_output(self):
        self.basic()
        cases = [dict(column_count=0), dict(master_header_rows=-1), dict(target_header_rows=True),
                 dict(master_columns=ColumnMapping("B", "C", "C")),
                 dict(target_columns=ColumnMapping("A", "A", "C")),
                 dict(master_columns=ColumnMapping("1", "B", "C")),
                 dict(target_columns=ColumnMapping("A", "B", "XFD"), column_count=2)]
        for kwargs in cases:
            with self.subTest(kwargs=kwargs), self.assertRaises(ValueError):
                self.sync(**kwargs)
            self.assertFalse(self.output.exists())

    def test_previous_nested_output_is_excluded_on_later_run(self):
        self.basic()
        self.output = self.targets / "first_result"
        self.sync()
        self.output = self.targets / "second_result"
        summary = self.sync()
        self.assertEqual([item.source for item in summary.files], ["a.xlsx"])
        self.assertFalse((self.output / "first_result").exists())

    def test_missing_formula_cache_does_not_clear_target(self):
        self.basic(value="=1+2")
        with self.assertRaisesRegex(ValueError, "公式缺少缓存"):
            self.sync(allow_blank_write=True)
        self.assertFalse(self.output.exists())

    def test_formula_cached_value_is_transferred_as_text(self):
        self.basic(value="=1+2")
        with ZipFile(self.master) as archive:
            entries = {name: archive.read(name) for name in archive.namelist()}
        sheet = ET.fromstring(entries["xl/worksheets/sheet1.xml"])
        ns = "{http://schemas.openxmlformats.org/spreadsheetml/2006/main}"
        sheet.find(f".//{ns}c[@r='D2']/{ns}v").text = "3"
        entries["xl/worksheets/sheet1.xml"] = ET.tostring(sheet)
        with ZipFile(self.master, "w") as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        self.sync()
        with self.book(self.output / "a.xlsx") as wb:
            self.assertEqual(wb.active["C2"].value, "3")
            self.assertEqual(wb.active["C2"].data_type, "s")

    def test_invalid_identity_skips_even_when_content_formula_has_no_cache(self):
        self.basic()
        with self.book(self.master) as wb:
            wb.active.append([2, None, "s", "=1+2"])
            wb.save(self.master)
        self.assertEqual(self.sync().updated_cells, 1)

    def test_target_formula_is_occupied_in_fill_blank_mode(self):
        self.basic(target="=1+2")
        summary = self.sync(fill_blank_only=True)
        self.assertEqual(summary.updated_cells, 0)
        with self.book(self.output / "a.xlsx") as wb:
            self.assertEqual(wb.active["C2"].value, "=1+2")

    def test_formatted_empty_tail_does_not_expand_target_scan(self):
        self.basic()
        with self.book(self.targets / "a.xlsx") as wb:
            wb.active["XFD1048576"].fill = PatternFill("solid", fgColor="FF0000")
            wb.save(self.targets / "a.xlsx")
        with patch("openpyxl.worksheet.worksheet.Worksheet.iter_rows", side_effect=AssertionError("dense scan")):
            summary = self.sync()
        self.assertEqual(summary.updated_cells, 1)

    def test_xlsm_macro_payload_is_preserved(self):
        self.basic()
        source = self.targets / "a.xlsx"
        macro = self.targets / "a.xlsm"
        source.rename(macro)
        with ZipFile(macro, "a") as archive:
            archive.writestr("xl/vbaProject.bin", b"test-macro-payload")
        self.sync()
        with ZipFile(self.output / "a.xlsm") as archive:
            self.assertEqual(archive.read("xl/vbaProject.bin"), b"test-macro-payload")


if __name__ == "__main__":
    unittest.main()
