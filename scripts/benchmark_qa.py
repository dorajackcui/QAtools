"""Synthetic QA performance diagnostics; each measurement uses a fresh process.

Core timings exclude fixture construction and disk I/O, and include temporary
problem-sheet construction. Workflow timings include load, checks, merge, save.
Peak RSS includes the interpreter and fixture. No user workbooks are accessed.
"""
from __future__ import annotations

import argparse
from contextlib import ExitStack
import cProfile
from dataclasses import fields, is_dataclass
import gc
import importlib
import json
from pathlib import Path
import platform
import pstats
import subprocess
import sys
import tempfile
from time import perf_counter, process_time
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts.benchmark_content_sync import Timings, peak_rss_mib

MODULES = {
    "term": "tools.term_pair_checker.extract_terms_from_excel",
    "source": "tools.source_consistency_checker.check_source_consistency",
    "target": "tools.target_consistency_checker.check_target_consistency",
    "substring": "tools.substring_consistency_checker.check_substring_consistency",
    "tag": "tools.tag_placeholder_checker.check_tags_and_placeholders",
    "line": "tools.line_break_checker.check_line_breaks",
    "number": "tools.content_fidelity_checker.check_content_fidelity",
    "url": "tools.content_fidelity_checker.check_content_fidelity",
    "chinese": "tools.chinese_target_checker.check_chinese_target",
    "text": "tools.target_text_checker.check_target_text",
}
TEXT_RULES = ("abnormal-punctuation", "consecutive-spaces", "leading-trailing-spaces",
              "mixed-width", "paired-symbols")
CHECKS = (*MODULES, "empty", *("text." + rule for rule in TEXT_RULES))


def fixture(count, scenario, error_percent=10, vocabulary=100):
    from openpyxl import Workbook
    book = Workbook()
    ws = book.active
    ws.title = "Data"
    ws.append(["source", "target"])
    for index in range(count):
        group = index // 4
        term = group % vocabulary
        source = f'Open [Term{term}] menu {group} and save the changes.'
        target = f'Ouvrir [Terme{term}] menu {group} et enregistrer les modifications.'
        if group % 3 == 0:
            source = f'<b>"{source}"</b> {{name}}\nhttps://example.test/{group}'
            target = f'<b>"{target}"</b> {{name}}\nhttps://example.test/{group}'
        if index % 100 < error_percent:
            defect = (index // 100) % 5
            if defect == 0:
                target = ""
            elif defect == 1:
                target = f'  Erreur..  中文（texte) {group + 1}'
            elif defect == 2:
                target = target.replace("Terme", "Incorrect").replace("</b>", "")
            elif defect == 3:
                target = "Reusable translation"
            else:
                target += ' "unfinished'
        if scenario == "source-group":
            source, target = "One source", f"Variant {index % 2}"
        elif scenario == "target-group":
            source, target = f"Source {index % 2}", "One target"
        elif scenario == "substring-dense":
            source, target = "中" * (index + 3), "译" * (index + 3)
        elif scenario == "long":
            source, target = source * 20, target * 20
        elif scenario == "sparse":
            if index:
                continue
            source, target = "Start", "Début"
        ws.append([source, target])
    if scenario == "sparse":
        ws.cell(count + 1, 1, "End")
        ws.cell(count + 1, 2, "Fin")
    return book


def instrument_tables(stack, timings):
    names = set(MODULES.values()) | {"tools.term_pair_checker.workbook_output"}
    for name in names:
        module = importlib.import_module(name)
        if hasattr(module, "write_output_table"):
            stack.enter_context(patch.object(module, "write_output_table",
                                            timings.wrap(module.write_output_table, "tables")))


def core(book, folder, job):
    import tools.workflow.workflow_runner as wf
    key = job["check"]
    if key == "empty":
        return wf.write_empty_target_problems(book, book["Data"], "A", "B", 2)
    base = key.split(".")[0]
    module = importlib.import_module(MODULES[base])
    options = dict(workbook=book, output_path=folder / "unused.xlsx", source_column="A",
                   target_column="B", sheet="Data", start_row=2, format_output=False)
    if base == "term":
        dummy = folder / "input.xlsx"
        dummy.touch()  # process_workbook only validates this path's existence.
        options["input_file"] = dummy
    if base in ("number", "url"):
        options["rules"] = (base,)
    if "." in key:
        options["rules"] = (key.split(".")[1],)
    return module.process_workbook(**options)


def workflow(book, folder, job, stack, timings):
    import tools.workflow.workflow_runner as wf
    from openpyxl.workbook.workbook import Workbook
    source = folder / "input.xlsx"
    book.save(source)
    book.close()
    # Fixture save is deliberately outside the timed run.
    for name in vars(wf):
        if name.startswith("run_") and name.endswith("_workbook"):
            stack.enter_context(patch.object(wf, name, timings.wrap(getattr(wf, name), name)))
    for name in ("load_workbook_for_editing", "write_empty_target_problems", "write_review_sheet"):
        stack.enter_context(patch.object(wf, name, timings.wrap(getattr(wf, name), name)))
    stack.enter_context(patch.object(Workbook, "save", timings.wrap(Workbook.save, "save")))
    enabled_all = job["check"] == "all"
    return lambda: wf.run_workflow(
        input_file=source, source_column="A", target_column="B", output_file=folder / "report.xlsx",
        run_target_consistency_check=enabled_all, run_substring_consistency_check=enabled_all,
    )


def worker(job):
    import openpyxl
    import tools.term_matching as matching
    with tempfile.TemporaryDirectory(prefix="qatools-qa-bench-") as directory:
        folder = Path(directory)
        book = fixture(job["rows"], job["scenario"], job["errors"], job.get("vocabulary", 100))
        fixture_cells = len(book["Data"]._cells)
        timings = Timings()
        profiler = cProfile.Profile() if job.get("profile") else None
        try:
            with ExitStack() as stack:
                instrument_tables(stack, timings)
                if job.get("trace_term_copies"):
                    # The baseline term module copies its mapping with dict().
                    # The optimized overlay no longer uses that constructor;
                    # keep the diagnostic available for baseline comparisons.
                    timings.seconds["term_mapping_copies"] = 0.0
                    term_module = importlib.import_module(MODULES["term"])
                    stack.enter_context(patch.object(term_module, "dict",
                        timings.wrap(dict, "term_mapping_copies"), create=True))
                action = (workflow(book, folder, job, stack, timings) if job["mode"] == "workflow"
                          else lambda: core(book, folder, job))
                gc.collect()
                if profiler:
                    profiler.enable()
                start = perf_counter()
                cpu_start = process_time()
                result = action()
                cpu_elapsed = process_time() - cpu_start
                elapsed = perf_counter() - start
                if profiler:
                    profiler.disable()
            output = dict(job, seconds=elapsed, cpu_seconds=cpu_elapsed,
                          phases=dict(timings.seconds), peak_rss_mib=peak_rss_mib(),
                          fixture_cells=fixture_cells,
                          final_cells=sum(len(s._cells) for s in book) if job["mode"] == "core" else None,
                          python=platform.python_version(), openpyxl=openpyxl.__version__,
                          aho_available=matching.ahocorasick is not None)
            if is_dataclass(result):
                output["counts"] = {field.name: getattr(result, field.name) for field in fields(result)
                                    if type(getattr(result, field.name)) is int}
            elif job["check"] == "empty":
                output["counts"] = {"problem_rows": result[1]}
            elif job["mode"] == "core" and job["check"] == "term":
                output["counts"] = {"term_count": result[-2], "problem_count": result[-1]}
            if profiler:
                stats = pstats.Stats(profiler)
                output["hotspots"] = [dict(file=Path(key[0]).name, line=key[1], function=key[2],
                                            calls=value[1], self_seconds=value[2], cumulative=value[3])
                                      for key, value in sorted(stats.stats.items(), key=lambda pair: pair[1][3], reverse=True)[:40]]
            return output
        finally:
            book.close()


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--worker", help=argparse.SUPPRESS)
    parser.add_argument("--mode", choices=("core", "workflow"), default="core")
    parser.add_argument("--check", action="append")
    parser.add_argument("--rows", type=int, nargs="+", default=[10000, 50000])
    parser.add_argument("--repeat", type=int, default=3)
    parser.add_argument("--scenario", choices=("mixed", "long", "source-group", "target-group", "substring-dense", "sparse"), default="mixed")
    parser.add_argument("--errors", type=int, default=10)
    parser.add_argument("--vocabulary", type=int, default=100)
    parser.add_argument("--profile", action="store_true")
    parser.add_argument("--trace-term-copies", action="store_true",
                        help="Measure row-transaction dictionary copying (adds timing overhead)")
    parser.add_argument("--timeout", type=int, default=120)
    parser.add_argument("--output", type=Path, default=ROOT / "outputs/qa-performance/core.json")
    args = parser.parse_args()
    if args.worker:
        print(json.dumps(worker(json.loads(args.worker)), ensure_ascii=True))
        return
    if min(args.rows) < 1 or args.repeat < 1 or args.vocabulary < 1 or args.timeout < 1 or not 0 <= args.errors <= 100:
        parser.error("rows, repeat, vocabulary and timeout must be positive; errors must be 0..100")
    checks = args.check or (list(CHECKS) if args.mode == "core" else ["default", "all"])
    allowed = CHECKS if args.mode == "core" else ("default", "all")
    if any(check not in allowed for check in checks):
        parser.error("unknown check for selected mode")
    output = {"platform": platform.platform(), "results": []}
    args.output.parent.mkdir(parents=True, exist_ok=True)
    for rows in args.rows:
        for check in checks:
            for iteration in range(args.repeat):
                job = dict(mode=args.mode, check=check, rows=rows, scenario=args.scenario, errors=args.errors,
                           vocabulary=args.vocabulary, profile=args.profile,
                           trace_term_copies=args.trace_term_copies, iteration=iteration)
                try:
                    result = subprocess.run([sys.executable, __file__, "--worker", json.dumps(job)],
                                            cwd=ROOT, capture_output=True, text=True, timeout=args.timeout)
                    if result.returncode:
                        entry = dict(job, error=result.stderr[-3000:])
                    else:
                        entry = json.loads(result.stdout)
                except subprocess.TimeoutExpired:
                    entry = dict(job, error=f"timeout after {args.timeout}s")
                output["results"].append(entry)
                args.output.write_text(json.dumps(output, indent=2, ensure_ascii=False), encoding="utf-8")
                print(json.dumps(entry, ensure_ascii=True), flush=True)
    return int(any("error" in entry for entry in output["results"]))


if __name__ == "__main__":
    raise SystemExit(main())
